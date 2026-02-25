# -*- coding: utf-8 -*-
import os
import sys
from pathlib import Path
from openpyxl import load_workbook

# Add src to path
sys.path.append(str(Path(__file__).parent / "src"))

from dsf_general_info import get_gulfcam_config
from dsf_general_prefill import DSFGeneralPrefiller
from smart_general_filler import SmartGeneralFiller

def verify_r3():
    template = Path("templates/DSF Normal standard.xlsx")
    output_manual = Path("output/test_r3_manual.xlsx")
    output_smart = Path("output/test_r3_smart.xlsx")
    mapping = Path("config/dsf_prefill_mapping.json")
    
    info = get_gulfcam_config()
    
    # --- Test Manual Filler ---
    print("Testing DSFGeneralPrefiller...")
    prefiller = DSFGeneralPrefiller(template, mapping)
    prefiller.load()
    prefiller.fill(info)
    prefiller.save(output_manual)
    
    wb_m = load_workbook(output_manual)
    ws_r3_m = None
    for sn in wb_m.sheetnames:
        if "R3" in sn.upper():
            ws_r3_m = wb_m[sn]
            break
            
    if ws_r3_m:
        # Actionnaires start at row 44 in mapping
        val = ws_r3_m["A44"].value
        print(f"Manual Filler R3 A44: {val}")
        if val and val == info.actionnaires[0].nom:
            print("ERROR: Manual filler still filling actionnaires in R3!")
        else:
            print("SUCCESS: Manual filler skipped actionnaires in R3.")
    
    # --- Test Smart Filler ---
    print("\nTesting SmartGeneralFiller...")
    smart = SmartGeneralFiller(template)
    smart.load()
    smart.fill(info)
    smart.save(output_smart)
    
    wb_s = load_workbook(output_smart)
    ws_r3_s = None
    for sn in wb_s.sheetnames:
        if "R3" in sn.upper():
            ws_r3_s = wb_s[sn]
            break
            
    if ws_r3_s:
        # Actionnaires start at row 44 in smart filler
        val = ws_r3_s["A44"].value
        print(f"Smart Filler R3 A44: {val}")
        if val and val == info.actionnaires[0].nom:
            print("ERROR: Smart filler still filling actionnaires in R3!")
        else:
            print("SUCCESS: Smart filler skipped actionnaires in R3.")
            
    # --- Verify Note 13 (should still be filled) ---
    ws_n13_s = None
    for sn in wb_s.sheetnames:
        if "NOTE 13" in sn.upper():
            ws_n13_s = wb_s[sn]
            break
            
    if ws_n13_s:
        val = ws_n13_s["A11"].value
        print(f"Smart Filler Note 13 A11: {val}")
        if val and val == info.actionnaires[0].nom:
            print("SUCCESS: Smart filler still fills Note 13.")
        else:
            print("WARNING: Smart filler skipped Note 13 too?")

if __name__ == "__main__":
    verify_r3()
