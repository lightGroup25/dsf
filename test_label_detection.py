"""
Test rapide de la détection dynamique des colonnes de labels
"""

import openpyxl
from pathlib import Path

TEMPLATE_FILE = Path(r"c:\Users\Emmaneul Ambadiang\Desktop\DSF\templates\DSF Normal standard.xlsx")

# Copier la fonction is_label_column du mapper
def is_label_column(ws, col_letter, data_start_row=15, sample_rows=20):
    """Déterminer si une colonne contient des LABELS"""
    text_count = 0
    number_count = 0
    empty_count = 0
    long_text_count = 0
    
    for row_idx in range(data_start_row, min(data_start_row + sample_rows, ws.max_row + 1)):
        try:
            cell_value = ws[f'{col_letter}{row_idx}'].value
        except:
            continue
        
        if cell_value is None or cell_value == '':
            empty_count += 1
        elif isinstance(cell_value, str):
            text_count += 1
            if len(cell_value.strip()) > 10:
                long_text_count += 1
        elif isinstance(cell_value, (int, float)):
            number_count += 1
    
    total_analyzed = text_count + number_count + empty_count
    
    if total_analyzed < 5:
        return False
    
    if long_text_count > (total_analyzed * 0.5):
        return True
    
    if text_count > (total_analyzed * 0.7):
        return True
    
    return False

def main():
    print("="*80)
    print("TEST DÉTECTION DYNAMIQUE DES COLONNES LABELS")
    print("="*80)
    
    wb = openpyxl.load_workbook(TEMPLATE_FILE)
    
    # Tester sur plusieurs sheets différentes
    test_sheets = ['BILAN PAYSAGE', 'NOTE 3A', 'NOTE 3B', 'NOTE 28']
    
    for sheet_name in test_sheets:
        if sheet_name not in wb.sheetnames:
            continue
        
        ws = wb[sheet_name]
        
        print(f"\n{'='*80}")
        print(f"Sheet: {sheet_name}")
        print('='*80)
        
        label_cols = []
        data_cols = []
        
        for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J']:
            is_label = is_label_column(ws, col, data_start_row=15, sample_rows=20)
            
            if is_label:
                label_cols.append(col)
            else:
                data_cols.append(col)
        
        print(f"\nCOLONNES LABELS détectées: {', '.join(label_cols) if label_cols else 'Aucune'}")
        print(f"COLONNES DATA disponibles: {', '.join(data_cols) if data_cols else 'Aucune'}")
        
        # Vérifier quelques cellules pour validation
        if label_cols:
            print(f"\nExemples de contenu (colonnes labels):")
            for col in label_cols[:2]:  # Montrer 2 colonnes max
                sample_values = []
                for row_idx in range(15, 20):
                    val = ws[f'{col}{row_idx}'].value
                    if val and isinstance(val, str):
                        sample_values.append(f"  {col}{row_idx}: {val[:50]}")
                
                if sample_values:
                    print(f"  Colonne {col}:")
                    for sv in sample_values[:3]:
                        print(sv)

if __name__ == "__main__":
    main()
