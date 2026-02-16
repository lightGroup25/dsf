"""
VRAI MAPPER - Remplit CELLULE PAR CELLULE en respectant:
1. Les libellés des colonnes (pour savoir quoi remplir)
2. Les libellés des lignes (pour identifier le compte)
3. Les données disponibles dans la balance
4. NE remplit QUE les colonnes appropriées
"""

import openpyxl
from difflib import SequenceMatcher
import json
from pathlib import Path

BALANCE_FILE = r"input/BALANCE AU 31 12 2024 VERSION DU 26 05 25.xlsx"
DSF_TEMPLATE = r"templates/DSF Normal standard.xlsx"
OUTPUT_FILE = r"output/DSF_FINAL_CORRECT.xlsx"

def load_balance():
    """Charger les comptes du balance"""
    wb = openpyxl.load_workbook(BALANCE_FILE)
    ws = wb.active
    
    accounts = {}
    for row_idx in range(1, ws.max_row + 1):
        account_num = ws[f'A{row_idx}'].value
        account_label = ws[f'D{row_idx}'].value
        opening = ws[f'N{row_idx}'].value  # Col 14
        closing = ws[f'AA{row_idx}'].value  # Col 27
        
        if account_num and account_label:
            try:
                int(str(account_num).split('.')[0])
                accounts[str(account_num).strip()] = {
                    'label': str(account_label).strip(),
                    'opening': opening,
                    'closing': closing
                }
            except:
                pass
    
    return accounts

def identify_column_types(ws, sheet_name):
    """Identifier le type de chaque colonne (ouverture, mouvements, clôture) basé sur les libellés"""
    column_types = {}
    
    # Scanner les en-têtes (rows 5-15 généralement)
    for row_idx in range(1, 20):
        for col_letter in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L']:
            cell_val = ws[f'{col_letter}{row_idx}'].value
            if cell_val:
                val_upper = str(cell_val).upper()
                
                # Identifier les colonnes spéciales
                if 'OUVERTURE' in val_upper or 'BRUTE A L' in val_upper:
                    column_types[col_letter] = 'OPENING'
                elif 'CLOTURE' in val_upper or 'BRUT A LA' in val_upper:
                    column_types[col_letter] = 'CLOSING'
                elif any(x in val_upper for x in ['ACQUISITION', 'APPORT', 'CREATION']):
                    column_types[col_letter] = 'MOVEMENT'
                elif any(x in val_upper for x in ['AMORT', 'DEPREC']):
                    column_types[col_letter] = 'AMORT'  # Ne pas remplir
    
    return column_types

def process_sheet_properly(wb, sheet_name, accounts):
    """Traiter chaque feuille correctement"""
    try:
        ws = wb[sheet_name]
    except:
        return 0
    
    # Identifier les types de colonnes
    col_types = identify_column_types(ws, sheet_name)
    
    if not col_types:
        return 0  # Pas de structure identifiée
    
    print(f"\n{sheet_name}:")
    print(f"  Column types found: {col_types}")
    
    cells_filled = 0
    
    # Scanner les lignes de donnéess
    for row_idx in range(10, min(200, ws.max_row + 1)):
        col_a = ws[f'A{row_idx}'].value
        col_b = ws[f'B{row_idx}'].value
        
        if not col_a and not col_b:
            continue
        
        row_label = str(col_b if col_b else col_a).strip()
        
        # Trouver le compte correspondant dans la balance
        best_score = 0
        best_account = None
        
        for acc_data in accounts.values():
            score = SequenceMatcher(None, row_label.lower(), acc_data['label'].lower()).ratio()
            if score > best_score:
                best_score = score
                best_account = acc_data
        
        # Remplir SEULEMENT les colonnes appropriées si match >= 0.4
        if best_score >= 0.4 and best_account:
            for col_letter, col_type in col_types.items():
                try:
                    cell = ws[f'{col_letter}{row_idx}']
                    
                    # Remplir seulement les colonnes d'ouverture et fermeture
                    if col_type == 'OPENING' and best_account['opening'] and not cell.value:
                        cell.value = best_account['opening']
                        cells_filled += 1
                    
                    elif col_type == 'CLOSING' and best_account['closing'] and not cell.value:
                        cell.value = best_account['closing']
                        cells_filled += 1
                    
                    # IGNORER les colonnes AMORT (orangées) et MOVEMENT (détail)
                
                except:
                    pass
    
    return cells_filled

def main():
    print("="*70)
    print("VRAI MAPPER - CELLULE PAR CELLULE AVEC LIBELLÉS RÉELS")
    print("="*70)
    
    accounts = load_balance()
    print(f"✓ Loaded {len(accounts)} accounts from balance")
    
    wb = openpyxl.load_workbook(DSF_TEMPLATE)
    print(f"✓ Template has {len(wb.sheetnames)} sheets")
    
    total_filled = 0
    sheets_done = 0
    
    for sheet_name in wb.sheetnames:
        filled = process_sheet_properly(wb, sheet_name, accounts)
        if filled > 0:
            total_filled += filled
            sheets_done += 1
            print(f"    → {filled} cells filled")
    
    print(f"\nSaving to {OUTPUT_FILE}...")
    Path(OUTPUT_FILE).parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_FILE)
    
    print("\n" + "="*70)
    print(f"✓ DONE: {total_filled} cells filled in {sheets_done} sheets")
    print(f"✓ Output: {OUTPUT_FILE}")
    print("="*70)

if __name__ == "__main__":
    main()
