import openpyxl

dsf_file = r'templates/DSF Normal standard.xlsx'
wb = openpyxl.load_workbook(dsf_file)

print('Available sheets:')
for i, sheet in enumerate(wb.sheetnames):
    print(f'{i}: {sheet}')
    
# Check first sheet of each
for sheet_name in wb.sheetnames[:3]:
    ws = wb[sheet_name]
    print(f'\n{sheet_name} - first 10 rows:')
    for row_idx in range(1, min(11, ws.max_row + 1)):
        col_a = ws[f'A{row_idx}'].value
        col_b = ws[f'B{row_idx}'].value
        if col_a or col_b:
            print(f'  Row {row_idx}: {col_a} | {col_b}')
