#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import openpyxl

wb = openpyxl.load_workbook(r'templates\DSF Normal standard.xlsx', data_only=False)
ws = wb['NOTE 3A']

print("="*130)
print("ANALYSE - FORMULES ET POURCENTAGES - NOTE 3A")
print("="*130)

# Vérifier les lignes 8-10
print("\nHEADERS (Rows 8-10):")
for row_num in [8, 9, 10]:
    print(f"\nRow {row_num}:")
    for col_num in [1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13]:
        cell = ws.cell(row_num, col_num)
        val = cell.value
        if val:
            print(f"  Col {col_num}: {str(val)[:70]}")

# Maintenant vérifier les formules dans les données
print("\n" + "="*130)
print("VERIFICATION - Y A-T-IL DES FORMULES DANS LES CELLULES DATA?")
print("="*130)

# Vérifier les premières lignes de données (row 11+)
print("\nCellules des premières lignes (Rows 11-15):")
for row_num in range(11, 16):
    print(f"\nRow {row_num}:")
    has_data = False
    for col_num in range(1, 14):
        cell = ws.cell(row_num, col_num)
        if cell.value:
            has_data = True
            # Vérifier si c'est une formule
            if isinstance(cell.value, str) and cell.value.startswith('='):
                print(f"  Col {col_num}: FORMULE = {str(cell.value)[:80]}")
            else:
                print(f"  Col {col_num}: {type(cell.value).__name__} = {str(cell.value)[:60]}")
    
    if not has_data:
        print("  [VIDE - pas de donnees]")

# Vérifier aussi NOTE 3B qui pourrait avoir des pourcentages/formules
if "NOTE 3B" in wb.sheetnames:
    print("\n" + "="*130)
    print("VERIFIER AUSSI NOTE 3B (DEPRECIATIONS)")
    print("="*130)
    
    ws3b = wb['NOTE 3B']
    print("\nHEADERS NOTE 3B (Row 8):")
    for col_num in range(1, 14):
        val = ws3b.cell(8, col_num).value
        if val:
            print(f"  Col {col_num}: {str(val)[:70]}")
    
    print("\nDonnees NOTE 3B (Row 11):")
    for col_num in range(1, 14):
        cell = ws3b.cell(11, col_num)
        if cell.value:
            if isinstance(cell.value, str) and cell.value.startswith('='):
                print(f"  Col {col_num}: FORMULE = {str(cell.value)[:80]}")
            else:
                print(f"  Col {col_num}: {str(cell.value)[:60]}")

wb.close()

print("\n" + "="*130)
print("QUESTION: Existe-t-il des colonnes avec taux d'amortissement (%)?")
print("="*130)
