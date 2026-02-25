#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Module pour remplir les cellules composites du BILAN PAYSAGE
qui référencent des Note Annexes  (P3-B1)

Les cellules L12:L16, L13, etc. doivent être remplies en cherchant les valeurs
dans les NOTE 7, NOTE 3C, etc. basé sur les références en colonne J.
"""
from __future__ import annotations

import logging
from pathlib import Path
from typing import Optional, Dict, Tuple
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

logger = logging.getLogger(__name__)


def fill_bilan_composite_cells(wb: Workbook) -> int:
    """
    Remplit les cellules composites du BILAN PAYSAGE qui référencent les Notes.
    
    Structure:
      - Colonne I: libellé de l'élément (Capital, Primes, Ecarts, etc.)
      - Colonne J: référence à la NOTE (13, 14, 3e, etc.)
      - Colonne L: à remplir avec la valeur de la NOTE
    
    Args:
        wb: Workbook ouvert
    
    Returns:
        Nombre de cellules remplies
    """
    try:
        ws = wb["BILAN PAYSAGE"]
    except KeyError:
        logger.warning("Feuille BILAN PAYSAGE introuvable")
        return 0
    
    filled_count = 0
    
    # Zones à remplir : (start_row, end_row, col_libellé, col_ref, col_valeur)
    # Ces zones contiennent les éléments qui ont une référence à une NOTE
    zones = [
        (12, 16, "I", "J", "L"),  # Capitaux propres
        # Ajouter d'autres zones si nécessaire
    ]
    
    for start_row, end_row, col_label, col_ref, col_value in zones:
        for row in range(start_row, end_row + 1):
            # Lire la référence à la NOTE (colonne J)
            ref_cell = ws[f"{col_ref}{row}"]
            ref_value = ref_cell.value
            
            if not ref_value:
                continue
            
            # Normaliser la référence (13 → "NOTE 13", 3e → "NOTE 3E", etc.)
            note_sheet_name = _resolve_note_reference(ref_value)
            
            if not note_sheet_name:
                continue
            
            # Chercher la NOTE dans le wb
            try:
                note_ws = wb[note_sheet_name]
            except KeyError:
                logger.debug(f"NOTE {note_sheet_name} introuvable pour {col_ref}{row}")
                continue
            
            # Chercher le libellé dans la NOTE
            label_cell = ws[f"{col_label}{row}"]
            label = label_cell.value
            
            if not label:
                continue
            
            # Chercher la ligne dans la NOTE qui correspond au libellé
            target_value = _find_value_in_note(note_ws, label)
            
            if target_value is not None:
                # Remplir la cellule
                value_cell = ws[f"{col_value}{row}"]
                if not value_cell.value:  # Ne pas écraser si déjà rempli
                    value_cell.value = target_value
                    filled_count += 1
                    logger.debug(f"  {col_value}{row} ← {target_value} (from {note_sheet_name})")
    
    return filled_count


def _resolve_note_reference(ref: str) -> Optional[str]:
    """
    Convertit une référence (13, 3e, etc.) en nom de feuille NOTE (NOTE 13, NOTE 3E, etc.)
    
    Args:
        ref: Référence brute (13, "3e", "3A", etc.)
    
    Returns:
        Nom de la feuille ou None si invalide
    """
    if not ref:
        return None
    
    ref_str = str(ref).strip()
    
    # Si déjà un nom de NOTE, retourner tel quel
    if ref_str.upper().startswith("NOTE"):
        return ref_str
    
    # Sinon, formatter : 13 → NOTE 13, 3e → NOTE 3E, 3A → NOTE 3A
    if ref_str.isdigit():
        return f"NOTE {ref_str}"
    else:
        # Lettre ajoutée : 3e → NOTE 3E
        return f"NOTE {ref_str.upper()}"


def _find_value_in_note(note_ws: Worksheet, label_pattern: str) -> Optional[float]:
    """
    Cherche dans la NOTE la ligne correspondant au libellé.
    
    Args:
        note_ws: Feuille NOTE
        label_pattern: Libellé à chercher
    
    Returns:
        La valeur trouvée ou None
    """
    if not note_ws or not label_pattern:
        return None
    
    label_norm = label_pattern.strip().upper()
    
    # Scanner la première colonne pour trouver le libellé (toutes les lignes)
    for row in range(1, note_ws.max_row + 1):
        cell = note_ws.cell(row, 1)  # Colonne A
        if cell.value:
            cell_norm = str(cell.value).strip().upper()
            if label_norm in cell_norm or cell_norm in label_norm:
                # Trouver la valeur (généralement colonne C ou D)
                for col in range(2, 6):
                    val_cell = note_ws.cell(row, col)
                    if val_cell.value and isinstance(val_cell.value, (int, float)):
                        return val_cell.value
    
    return None
