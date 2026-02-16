# -*- coding: utf-8 -*-
"""Inventory utility for DSF template input zones."""
from __future__ import annotations

import argparse
import json
import logging
from dataclasses import dataclass, asdict, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, List, Optional, Sequence

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

logger = logging.getLogger(__name__)


@dataclass
class ColumnDescriptor:
    column: int
    letter: str
    header_value: str
    header_cell: str
    column_type: str  # ex: "exercice_n", "exercice_n1"


@dataclass
class TemplateCell:
    sheet: str
    cell: str
    row: int
    column: int
    column_letter: str
    column_type: str
    label: Optional[str]
    section_hint: Optional[str]
    locked: bool
    has_formula: bool
    number_format: Optional[str]
    data_validation: List[Dict[str, str]] = field(default_factory=list)
    merged: bool = False
    merge_range: Optional[str] = None

    def to_dict(self) -> Dict[str, object]:
        payload = asdict(self)
        payload["locked"] = bool(self.locked)
        payload["has_formula"] = bool(self.has_formula)
        return payload


def _normalize_header_value(value: str) -> str:
    return " ".join(value.upper().split())


def _classify_header(value: str) -> Optional[str]:
    normalized = _normalize_header_value(value)
    compact = normalized.replace(" ", "")
    if "EXERCICE" not in normalized:
        return None
    if "N-2" in compact:
        return "exercice_n2"
    if "N-1" in compact:
        return "exercice_n1"
    if "N" in compact:
        return "exercice_n"
    return None


def _classify_note_header(value: str) -> Optional[str]:
    normalized = _normalize_header_value(value)
    compact = normalized.replace(" ", "")

    # Handle N-1 without requiring EXERCICE
    if "N-1" in compact or "N-1" in normalized:
        return "exercice_n1"

    # Catch common patterns: "N" alone or with labels (NET N, BRUT N, AMORT N)
    if "N" in compact:
        if "N-1" in compact:
            return "exercice_n1"
        return "exercice_n"

    return None


def detect_input_columns(ws: Worksheet, allow_note_mode: bool = False) -> List[ColumnDescriptor]:
    descriptors: Dict[int, ColumnDescriptor] = {}
    scan_rows = min(ws.max_row, 60 if not allow_note_mode else 200)
    max_col = min(ws.max_column, 80)
    for row in ws.iter_rows(min_row=1, max_row=scan_rows, max_col=max_col):
        for cell in row:
            value = cell.value
            if not isinstance(value, str):
                continue
            column_type = _classify_header(value)
            if not column_type:
                continue
            if cell.column not in descriptors:
                descriptors[cell.column] = ColumnDescriptor(
                    column=cell.column,
                    letter=get_column_letter(cell.column),
                    header_value=value.strip(),
                    header_cell=cell.coordinate,
                    column_type=column_type,
                )
    if descriptors or not allow_note_mode:
        return sorted(descriptors.values(), key=lambda desc: (desc.column, desc.column_type))

    # NOTE/SYNTHESE fallback: detect columns with N/N-1 headers even without EXERCICE
    note_descriptors: Dict[int, ColumnDescriptor] = {}
    for row in ws.iter_rows(min_row=1, max_row=scan_rows, max_col=max_col):
        for cell in row:
            value = cell.value
            if not isinstance(value, str):
                continue
            column_type = _classify_note_header(value)
            if not column_type:
                continue
            if cell.column not in note_descriptors:
                note_descriptors[cell.column] = ColumnDescriptor(
                    column=cell.column,
                    letter=get_column_letter(cell.column),
                    header_value=value.strip(),
                    header_cell=cell.coordinate,
                    column_type=column_type,
                )
    return sorted(note_descriptors.values(), key=lambda desc: (desc.column, desc.column_type))


def _is_numeric_format(number_format: Optional[str]) -> bool:
    if not number_format:
        return False
    fmt = number_format.replace(" ", "")
    if fmt.startswith("@"):  # text
        return False
    return any(ch in fmt for ch in ("0", "#", "%"))


def detect_note_value_columns(ws: Worksheet, validation_map: Dict[str, List[Dict[str, str]]]) -> List[ColumnDescriptor]:
    """Fallback detector for NOTE/SYNTHESE sheets without EXERCICE headers."""
    scores: Dict[int, int] = {}
    scan_rows = min(ws.max_row, 800)
    header_scan_rows = min(ws.max_row, 200)
    max_col = min(ws.max_column, 80)
    header_keywords = [
        "montant",
        "valeur",
        "solde",
        "brut",
        "net",
        "amort",
        "provision",
        "taux",
        "quant",
        "nombre",
        "effectif",
        "chiffre",
        "pourcentage",
        "variation",
    ]

    columns_from_headers = set()
    for row in ws.iter_rows(min_row=1, max_row=header_scan_rows, max_col=max_col):
        for cell in row:
            if not isinstance(cell.value, str):
                continue
            value = cell.value.strip().lower()
            if any(kw in value for kw in header_keywords):
                columns_from_headers.add(cell.column)

    for row_idx in range(1, scan_rows + 1):
        for col_idx in range(1, max_col + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            label = extract_label(ws, row_idx, col_idx)
            if not label:
                continue

            coord = cell.coordinate
            dv = validation_map.get(coord, [])
            has_numeric_validation = any(d.get("type") in ("decimal", "whole") for d in dv)
            is_numeric_like = isinstance(cell.value, (int, float)) or _is_numeric_format(cell.number_format)

            if has_numeric_validation or is_numeric_like:
                scores[col_idx] = scores.get(col_idx, 0) + 1

    # Threshold: at least 3 numeric-like cells with labels in the column
    descriptors: List[ColumnDescriptor] = []
    candidate_columns = set(scores.keys()) | columns_from_headers
    for col_idx in sorted(candidate_columns):
        score = scores.get(col_idx, 0)
        if score < 3 and col_idx not in columns_from_headers:
            continue
        col_letter = get_column_letter(col_idx)

        # Try to infer column type from headers in the column
        column_type = "exercice_n"
        for row in ws.iter_rows(min_row=1, max_row=header_scan_rows, max_col=max_col):
            cell = row[col_idx - 1]
            if isinstance(cell.value, str):
                detected = _classify_note_header(cell.value)
                if detected:
                    column_type = detected
                    break

        descriptors.append(
            ColumnDescriptor(
                column=col_idx,
                letter=col_letter,
                header_value="NOTE_VALUE",
                header_cell=f"{col_letter}1",
                column_type=column_type,
            )
        )
    return sorted(descriptors, key=lambda desc: (desc.column, desc.column_type))


def build_validation_map(ws: Worksheet) -> Dict[str, List[Dict[str, str]]]:
    mapping: Dict[str, List[Dict[str, str]]] = {}
    data_validations = getattr(ws, "data_validations", None)
    if not data_validations:
        return mapping

    for dv in data_validations.dataValidation:
        attrs = {
            "type": dv.type or "any",
            "operator": dv.operator or "",
            "formula1": dv.formula1 or "",
            "formula2": dv.formula2 or "",
            "allow_blank": str(dv.allowBlank),
            "show_error": str(dv.showErrorMessage),
        }
        for sqref in str(dv.sqref).split():
            min_col, min_row, max_col, max_row = range_boundaries(sqref)
            for row in range(min_row, max_row + 1):
                for col in range(min_col, max_col + 1):
                    coord = f"{get_column_letter(col)}{row}"
                    mapping.setdefault(coord, []).append(attrs)
    return mapping


def build_merge_lookup(ws: Worksheet) -> Dict[str, str]:
    lookup: Dict[str, str] = {}
    for merged in ws.merged_cells.ranges:
        for row in range(merged.min_row, merged.max_row + 1):
            for col in range(merged.min_col, merged.max_col + 1):
                coord = f"{get_column_letter(col)}{row}"
                lookup[coord] = str(merged)
    return lookup


def extract_label(ws: Worksheet, row: int, col: int) -> Optional[str]:
    for idx in range(col - 1, 0, -1):
        value = ws.cell(row=row, column=idx).value
        if value is None:
            continue
        if isinstance(value, str) and value.strip():
            return value.strip()
        if isinstance(value, (int, float)):
            return str(value)
    # Fallback: look upward in same column (header-based labels)
    for r_idx in range(row - 1, max(row - 25, 1), -1):
        value = ws.cell(row=r_idx, column=col).value
        if value is None:
            continue
        if isinstance(value, str) and value.strip():
            return value.strip()
    return None


def extract_section_hint(ws: Worksheet, row: int) -> Optional[str]:
    for r_idx in range(row, max(row - 25, 1), -1):
        for c_idx in (1, 2):
            value = ws.cell(row=r_idx, column=c_idx).value
            if isinstance(value, str):
                stripped = value.strip()
                if not stripped:
                    continue
                if stripped.isupper() and len(stripped) > 3:
                    return stripped
                if stripped.endswith(":"):
                    return stripped.rstrip(":")
    return None


def extract_sheet_inventory(
    ws: Worksheet,
    columns: Sequence[ColumnDescriptor],
    include_locked: bool = False,
    include_column_types: Optional[Sequence[str]] = None,
    allow_note_mode: bool = False,
) -> List[TemplateCell]:
    validation_map = build_validation_map(ws)
    merge_lookup = build_merge_lookup(ws)
    records: List[TemplateCell] = []

    if include_column_types:
        allowed_types = set(include_column_types)
    else:
        allowed_types = {descriptor.column_type for descriptor in columns}

    if not columns and allow_note_mode:
        columns = detect_note_value_columns(ws, validation_map)
        allowed_types = {descriptor.column_type for descriptor in columns}

    for descriptor in columns:
        if descriptor.column_type not in allowed_types:
            continue
        for row_idx in range(1, ws.max_row + 1):
            cell = ws.cell(row=row_idx, column=descriptor.column)
            locked = cell.protection.locked
            has_formula = bool(cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")))
            if locked and not include_locked:
                continue
            label = extract_label(ws, row_idx, descriptor.column)
            if not label:
                continue
            coord = cell.coordinate
            record = TemplateCell(
                sheet=ws.title,
                cell=coord,
                row=row_idx,
                column=descriptor.column,
                column_letter=descriptor.letter,
                column_type=descriptor.column_type,
                label=label,
                section_hint=extract_section_hint(ws, row_idx),
                locked=locked,
                has_formula=has_formula,
                number_format=cell.number_format,
                data_validation=validation_map.get(coord, []),
                merged=coord in merge_lookup,
                merge_range=merge_lookup.get(coord),
            )
            records.append(record)
    return records


def extract_template_inventory(
    template_path: Path,
    sheet_whitelist: Optional[Sequence[str]] = None,
    include_locked: bool = False,
    include_column_types: Optional[Sequence[str]] = None,
) -> Dict[str, object]:
    wb = load_workbook(template_path, data_only=False, keep_links=True, keep_vba=True)
    sheets = sheet_whitelist or wb.sheetnames
    payload: Dict[str, object] = {
        "template": str(template_path),
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "sheets": {},
    }

    for sheet_name in sheets:
        if sheet_name not in wb.sheetnames:
            logger.warning("Sheet %s absent du classeur", sheet_name)
            continue
        ws = wb[sheet_name]
        allow_note_mode = "NOTE" in sheet_name.upper() or "SYNTHESE" in sheet_name.upper() or "SYNTH" in sheet_name.upper()
        include_locked_sheet = include_locked or allow_note_mode
        descriptors = detect_input_columns(ws, allow_note_mode=allow_note_mode)
        if descriptors:
            records = extract_sheet_inventory(
                ws,
                descriptors,
                include_locked=include_locked_sheet,
                include_column_types=include_column_types,
                allow_note_mode=allow_note_mode,
            )
        else:
            records = extract_sheet_inventory(
                ws,
                descriptors,
                include_locked=include_locked_sheet,
                include_column_types=include_column_types,
                allow_note_mode=allow_note_mode,
            )
        payload["sheets"][sheet_name] = {
            "input_columns": [asdict(desc) for desc in descriptors],
            "fields": [record.to_dict() for record in records],
            "stats": {
                "rows": ws.max_row,
                "columns": ws.max_column,
                "input_fields": len(records),
            },
        }
        logger.info("Sheet %s: %s colonnes détectées | %s champs identifiés", sheet_name, len(descriptors), len(records))

    wb.close()
    return payload


def save_inventory(inventory: Dict[str, object], destination: Path) -> Path:
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(inventory, ensure_ascii=False, indent=2), encoding="utf-8")
    logger.info("Inventaire sauvegardé dans %s", destination)
    return destination


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Extract input zones from DSF template")
    parser.add_argument("--template", default="DSF Normal standard.xlsx", help="Chemin du template DSF")
    parser.add_argument("--output", default="data/dsf_inventory.json", help="Fichier JSON de sortie")
    parser.add_argument(
        "--sheet",
        action="append",
        dest="sheets",
        help="Limite l'analyse à une feuille (option répétable)",
    )
    parser.add_argument(
        "--include-locked",
        action="store_true",
        help="Inclut les cellules verrouillées dans l'inventaire",
    )
    parser.add_argument(
        "--include-column-type",
        action="append",
        dest="column_types",
        help="Types de colonnes à inclure (ex: exercice_n1)",
    )
    parser.add_argument("--verbose", action="store_true")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    logging.basicConfig(level=logging.DEBUG if args.verbose else logging.INFO)
    inventory = extract_template_inventory(
        Path(args.template),
        sheet_whitelist=args.sheets,
        include_locked=args.include_locked,
        include_column_types=args.column_types,
    )
    save_inventory(inventory, Path(args.output))


if __name__ == "__main__":
    main()
