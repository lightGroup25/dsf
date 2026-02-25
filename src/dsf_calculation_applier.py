# -*- coding: utf-8 -*-
"""DSF Calculation Applier - Integrates calculation manager into pipeline."""

import logging
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)


class DSFCalculationApplier:
    """
    Applies calculations (TOTALS, VARIATIONS, RATIOS, FORMULAS) to filled DSF.
    
    This applies the calculation system to generate Excel formulas for:
    - TOTAL rows (=SUM(...))
    - VARIATION columns (=Closing - Opening)
    - PERCENTAGE columns (=(New-Old)/Old*100)
    - RATIO columns (=Value1/Value2)
    - COMPLEX FORMULAS from template (=A+B+D-H adapted to each row)
    """
    
    def __init__(
        self,
        dsf_file_path: Path,
        use_formulas: bool = True,
        has_previous_balance_n1: bool = True,
        enforce_n1_from_previous_only: bool = True,
        overwrite_formulas: bool = False,
    ):
        """
        Initialize calculator.

        Args:
            dsf_file_path: Path to DSF file (already filled with balance data)
            use_formulas: If True, create Excel formulas; if False, create numeric values
            overwrite_formulas: If False (default), never overwrite existing Excel formulas
                in cells — this protects =SUM() and other template formulas from being
                replaced by calculated numeric values.
        """
        self.dsf_file_path = Path(dsf_file_path)
        self.use_formulas = use_formulas
        self.has_previous_balance_n1 = bool(has_previous_balance_n1)
        self.enforce_n1_from_previous_only = bool(enforce_n1_from_previous_only)
        self.overwrite_formulas = bool(overwrite_formulas)
        self.wb = None
        self.formulas_created = 0
        self.preserved_formulas = 0
        
    def apply(self) -> int:
        """
        Apply calculations to all sheets in DSF.
        
        Returns:
            Number of formulas/calculations created
        """
        try:
            # Import here to avoid circular dependency and allow missing module
            from calculation_manager import get_calculation_manager
        except ImportError:
            logger.warning("calculation_manager not available - skipping calculations")
            return 0
        
        self.wb = load_workbook(self.dsf_file_path)
        self.formulas_created = 0
        
        logger.info("Applying calculations to DSF: %s", self.dsf_file_path)
        
        for sheet_name in self.wb.sheetnames:
            ws = self.wb[sheet_name]
            
            # Skip empty sheets
            if not ws.max_row or not ws.max_column:
                continue
            
            logger.info("  Processing sheet: %s", sheet_name)
            
            # Detect column types for this sheet
            column_info = self._detect_column_types(ws)
            
            if not column_info.get('data_columns'):
                logger.debug("    No data columns detected in %s", sheet_name)
                continue
            
            # Create calculation manager for this sheet
            try:
                calc_mgr = get_calculation_manager(ws, column_info)
            except Exception as e:
                logger.warning("Error creating calculator for %s: %s", sheet_name, e)
                continue
            
            # Apply calculations to all data cells
            sheet_formulas = self._apply_sheet_calculations(
                ws, calc_mgr, column_info
            )
            
            self.formulas_created += sheet_formulas
            logger.info("    %d calculations applied", sheet_formulas)
        
        # Save with formulas
        try:
            self.wb.save(self.dsf_file_path)
            logger.info(
                "DSF saved with %d formulas applied, %d formules Excel préservées",
                self.formulas_created,
                self.preserved_formulas,
            )
        except Exception as e:
            logger.error("Error saving DSF: %s", e)
            raise
        finally:
            if self.wb:
                self.wb.close()

        return self.formulas_created
    
    def _detect_column_types(self, ws):
        """Detect column types (OPENING, CLOSING, VARIATION, TOTAL, etc)."""
        column_info = {
            'label_columns': [],
            'data_columns': [],
            'column_types': {},
            'year_hints': {},
        }
        
        # Try to find headers and detect column types
        # Look for row with common headers
        for row_idx in range(1, min(20, ws.max_row + 1)):
            row_data = [ws.cell(row_idx, col).value for col in range(1, ws.max_column + 1)]
            
            if not row_data or all(not v for v in row_data):
                continue
            
            # Check if this looks like a header row
            header_str = ' '.join(str(v or '').upper() for v in row_data)
            
            if any(kw in header_str for kw in [
                'LIBELLE', 'ACCOUNT', 'DESCRIPTION', 
                'OUVERTURE', 'OPENING', 'CLÔTURE', 'CLOSING',
                'VARIATION', 'MOUVEMENT', 'PERCENTAGE', '%'
            ]):
                # Found header row - analyze columns
                for col_idx, col_val in enumerate(row_data, 1):
                    if not col_val:
                        continue
                    
                    col_val_upper = str(col_val).upper()
                    col_letter = get_column_letter(col_idx)
                    
                    # Classify column
                    if any(kw in col_val_upper for kw in ['LIBELLE', 'LABEL', 'ACCOUNT', 'DESCRIPTION']):
                        column_info['label_columns'].append(col_letter)
                    elif any(kw in col_val_upper for kw in ['OUVERTURE', 'OPENING', 'SOLDE INITIAL']):
                        column_info['data_columns'].append(col_letter)
                        column_info['column_types'][col_letter] = 'OPENING'
                    elif any(kw in col_val_upper for kw in ['CLÔTURE', 'CLOSING', 'SOLDE FINAL']):
                        column_info['data_columns'].append(col_letter)
                        column_info['column_types'][col_letter] = 'CLOSING'
                    elif any(kw in col_val_upper for kw in ['MOUVEMENT', 'MOVEMENT', 'ÉCHEANCE']):
                        column_info['data_columns'].append(col_letter)
                        column_info['column_types'][col_letter] = 'MOVEMENT'
                    elif any(kw in col_val_upper for kw in ['VARIATION', 'ÉCART', 'RÉSULTAT']):
                        column_info['data_columns'].append(col_letter)
                        column_info['column_types'][col_letter] = 'VARIATION'
                    elif any(kw in col_val_upper for kw in ['%', 'POURCENTAGE', 'TAUX', 'RATIO']):
                        column_info['data_columns'].append(col_letter)
                        if '%' in col_val_upper:
                            column_info['column_types'][col_letter] = 'PERCENTAGE'
                        else:
                            column_info['column_types'][col_letter] = 'RATIO'
                    else:
                        # Default numeric column
                        column_info['data_columns'].append(col_letter)
                        column_info['column_types'][col_letter] = 'VALUE'
                    year_hint = self._infer_year_hint_from_header(col_val_upper)
                    if year_hint:
                        column_info['year_hints'][col_letter] = year_hint
                
                break
        
        # If no headers detected, assume all non-label columns are data
        if not column_info['data_columns']:
            column_info['data_columns'] = [
                get_column_letter(col_idx) 
                for col_idx in range(2, ws.max_column + 1)
            ]
        
        return column_info
    
    def _apply_sheet_calculations(self, ws, calc_mgr, column_info):
        """Apply calculations to data rows."""
        try:
            from calculation_manager import get_calculation_manager
        except ImportError:
            return 0

        formulas_created = 0
        sheet_preserved = 0

        # Find data range (skip header rows)
        start_row = 15  # Typical start for DSF data
        end_row = ws.max_row

        # Process each data row
        year_hints = column_info.get('year_hints', {})
        strict_no_n1 = self.enforce_n1_from_previous_only and not self.has_previous_balance_n1
        for row_idx in range(start_row, end_row + 1):
            # Get label for this row (check first label column)
            row_label = ""
            if column_info['label_columns']:
                label_cell = ws[f"{column_info['label_columns'][0]}{row_idx}"]
                row_label = str(label_cell.value or "").strip()

            # Process each data column
            for col_letter in column_info['data_columns']:
                col_type = column_info['column_types'].get(col_letter, 'VALUE')

                try:
                    if strict_no_n1 and year_hints.get(col_letter) == "n1":
                        continue

                    # P3-M5 : Protéger les formules Excel existantes du template
                    cell_ref = f'{col_letter}{row_idx}'
                    existing_val = ws[cell_ref].value
                    if (
                        not self.overwrite_formulas
                        and isinstance(existing_val, str)
                        and existing_val.startswith('=')
                    ):
                        logger.debug(
                            "Formule Excel préservée : %s!%s = '%s'",
                            ws.title, cell_ref, existing_val[:40],
                        )
                        sheet_preserved += 1
                        continue

                    # Check if calculation needed
                    should_calc, calc_type, extra_info = calc_mgr.should_calculate_cell(
                        row_idx, col_letter, col_type
                    )

                    if should_calc:
                        # Apply calculation
                        value = calc_mgr.apply_calculation(
                            row_idx, col_letter, col_type, calc_type,
                            extra_info=extra_info,
                            use_formulas=self.use_formulas,
                            use_balance=False  # Already filled by balance
                        )

                        if value is not None:
                            if (
                                strict_no_n1
                                and isinstance(value, str)
                                and value.startswith("=")
                                and self._formula_references_n1(value, year_hints)
                            ):
                                continue
                            ws[cell_ref] = value
                            formulas_created += 1

                except Exception as e:
                    logger.debug("Error calculating %s%d: %s", col_letter, row_idx, e)
                    continue

        if sheet_preserved:
            logger.info(
                "  %s : %d formules Excel préservées (overwrite_formulas=False)",
                ws.title, sheet_preserved,
            )
            self.preserved_formulas += sheet_preserved

        return formulas_created

    @staticmethod
    def _infer_year_hint_from_header(header_upper: str):
        txt = (header_upper or "").replace(" ", "")
        if "N-1" in header_upper or "N-1" in txt or "N1" in txt:
            return "n1"
        if "EXERCICEN" in txt or txt.endswith("N"):
            return "n"
        return None

    @staticmethod
    def _formula_references_n1(formula: str, year_hints: dict) -> bool:
        import re

        refs = re.findall(r"\$?([A-Z]{1,3})\$?\d+", (formula or "").upper())
        for col in refs:
            if year_hints.get(col) == "n1":
                return True
        return False


def apply_calculations_to_dsf(dsf_file_path: Path, use_formulas: bool = True) -> int:
    """
    Convenience function to apply calculations to a DSF file.
    
    Args:
        dsf_file_path: Path to DSF file
        use_formulas: If True, create Excel formulas; else numeric values
        
    Returns:
        Number of calculations created
    """
    applier = DSFCalculationApplier(dsf_file_path, use_formulas=use_formulas)
    return applier.apply()


if __name__ == "__main__":
    import sys
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s %(levelname)s %(name)s %(message)s"
    )
    
    if len(sys.argv) > 1:
        dsf_path = Path(sys.argv[1])
    else:
        dsf_path = Path("output/DSF_OUTPUT_2024.xlsx")
    
    if dsf_path.exists():
        count = apply_calculations_to_dsf(dsf_path, use_formulas=True)
        print(f"\n✓ {count} calculations applied to {dsf_path}")
    else:
        print(f"File not found: {dsf_path}")
