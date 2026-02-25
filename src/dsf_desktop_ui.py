#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""LIGHT DSF MAPPER – Design GULFCAM avec Landing Page Animée et Excel Intégré (CORRIGÉ - FINAL)"""

from __future__ import annotations

import json
import math
import os
import shutil
import subprocess
import time
import tempfile
from pathlib import Path
from typing import Callable, Dict, List, Optional
from dataclasses import dataclass, asdict

IS_WINDOWS = os.name == "nt"

if IS_WINDOWS:
    try:
        import pythoncom
        import win32com.client
        import win32gui
        import win32con
        import win32process
        from win32com.client import Dispatch
        WINDOWS_COM_AVAILABLE = True
    except ImportError:
        pythoncom = None
        win32com = None
        win32gui = None
        win32con = None
        win32process = None
        WINDOWS_COM_AVAILABLE = False
else:
    pythoncom = None
    win32com = None
    win32gui = None
    win32con = None
    win32process = None
    WINDOWS_COM_AVAILABLE = False

from PySide6.QtCore import (
    Qt,
    QTimer,
    QSize,
    QPropertyAnimation,
    QEasingCurve,
    QUrl,
    QObject,
    QThread,
    Signal,
    Slot,
)
from PySide6.QtGui import (
    QDesktopServices,
    QFont,
    QAction,
    QColor,
    QIcon,
    QKeySequence,
    QShortcut,
    QPixmap,
    QPainter,
    QWindow,
)
from PySide6.QtWidgets import (
    QApplication,
    QComboBox,
    QFileDialog,
    QHeaderView,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QSizePolicy,
    QStackedWidget,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
    QFrame,
    QProgressBar,
    QGraphicsDropShadowEffect,
    QStatusBar,
    QStyle,
    QGridLayout,
    QScrollArea,
)

# FontAwesome (qtawesome)
try:
    import qtawesome as qta
    FONTS_AVAILABLE = True
except ImportError:
    qta = None
    FONTS_AVAILABLE = False

# Lazy imports
try:
    from fill_dsf_from_balance import BalanceReader
except ImportError:
    BalanceReader = None

try:
    from dsf_pipeline import DSFPipeline, DSFPipelineConfig, PipelineArtifacts
except ImportError as e:
    print(f"Warning: Could not import DSF modules: {e}")
    DSFPipeline = None
    DSFPipelineConfig = None
    PipelineArtifacts = None

# CHARTE GRAPHIQUE GULFCAM
GULFCAM_COLORS = {
    'primary': '#118E55',        # Vert renforcé pour meilleure visibilité
    'primary_light': '#1CAC69',  # Vert lumineux pour hover/active
    'secondary': '#D8173A',      # Rouge logo (anneau)
    'secondary_light': '#EF3E5D',# Rouge adouci
    'accent': '#F0C419',         # Jaune logo (swoosh)
    'success': '#28B273',        # Vert de confirmation
    'danger': '#C81E35',         # Rouge d'erreur
    'warning': '#D8A90E',        # Jaune d'alerte
    'light': '#F3F9F4',          # Fond clair légèrement verdoyant
    'dark': '#102D21',           # Texte sombre lisible
    'gray': '#5E7166',           # Gris chaud neutre
    'light-gray': '#D6E3DA',     # Bordures claires
    'card-bg': '#FFFFFF',        # Fond de carte
    'gradient_start': '#11784A', # Dégradé splash/header
    'gradient_end': '#1CAC69',
}

# Pour le mode sombre
GULFCAM_DARK_COLORS = {
    'primary': '#4DD38E',        # Vert renforcé pour fond sombre
    'primary_light': '#73E8AB',
    'secondary': '#FF5A74',      # Rouge réhaussé pour contraste
    'secondary_light': '#FF8095',
    'accent': '#F7D14A',         # Jaune réchauffé
    'success': '#44E0A5',
    'danger': '#FB7185',
    'warning': '#F4C74F',
    'light': '#101A14',          # Fond global sombre
    'dark': '#F2F7F3',           # Texte principal clair
    'gray': '#A6B8AE',           # Texte secondaire
    'light-gray': '#2A3D33',     # Bordures sombres
    'card-bg': '#18251E',        # Cartes en mode sombre
    'gradient_start': '#0E2A1B',
    'gradient_end': '#1E6B47',
}

# Logo GULFCAM
GULFCAM_LOGO_TEXT = "GULFCAM"
GULFCAM_TAGLINE = "S'investir pour vous"
PRIMARY_LOGO_FILENAME = "OIP.png"
PRIMARY_LOGO_FALLBACK = "OIP.webp"
LIGHTGROUP_LOGO_FILENAME = "lightgroup.png"
APP_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = APP_DIR.parent
DEFAULT_TEMPLATE_PATH = PROJECT_ROOT / "templates" / "DSF Normal standard.xlsx"
DEFAULT_INVENTORY_PATH = PROJECT_ROOT / "data" / "dsf_inventory.json"
DEFAULT_RULES_PATH = PROJECT_ROOT / "config" / "dsf_rule_generated.yaml"
DEFAULT_PREFILL_MAPPING_PATH = PROJECT_ROOT / "config" / "dsf_prefill_mapping.json"
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "output"
DEFAULT_REPORTS_DIR = DEFAULT_OUTPUT_DIR / "reports"

# App Settings
SETTINGS_FILE = Path.home() / ".light_dsf_mapper_settings.json"

@dataclass
class AppSettings:
    """Paramètres de l'application persistants."""
    dark_mode: bool = False
    window_width: int = 1400
    window_height: int = 900
    recent_files: List[str] = None
    recent_dsf_files: List[str] = None  # Historique des DSF générés
    
    def __post_init__(self):
        if self.recent_files is None:
            self.recent_files = []
        if self.recent_dsf_files is None:
            self.recent_dsf_files = []
    
    def add_recent_dsf(self, path: str, max_entries: int = 30) -> None:
        """Ajoute un DSF à l'historique (dédoublonné, limité)."""
        path_str = str(path).replace("\\", "/")
        if path_str in self.recent_dsf_files:
            self.recent_dsf_files.remove(path_str)
        self.recent_dsf_files.insert(0, path_str)
        self.recent_dsf_files = self.recent_dsf_files[:max_entries]
    
    def save(self):
        try:
            with open(SETTINGS_FILE, 'w') as f:
                json.dump(asdict(self), f, indent=2)
        except Exception as e:
            print(f"Erreur sauvegarde settings: {e}")
    
    @staticmethod
    def load():
        if SETTINGS_FILE.exists():
            try:
                with open(SETTINGS_FILE, 'r') as f:
                    data = json.load(f)
                    valid_keys = ['dark_mode', 'window_width', 'window_height', 'recent_files', 'recent_dsf_files']
                    filtered_data = {k: v for k, v in data.items() if k in valid_keys}
                    return AppSettings(**filtered_data)
            except Exception as e:
                print(f"Erreur chargement settings: {e}")
        return AppSettings()

def get_gulfcam_colors(dark_mode=False):
    """Retourne les couleurs GULFCAM selon le thème"""
    return GULFCAM_DARK_COLORS if dark_mode else GULFCAM_COLORS


def resolve_image_path(filename: str) -> Optional[Path]:
    """Résout le chemin d'une image dans le projet."""
    here = Path(__file__).resolve().parent
    candidates = (
        Path.cwd() / filename,
        here / filename,
        here.parent / filename,
    )
    for candidate in candidates:
        if candidate.exists():
            return candidate
    return None


def load_image_pixmap(filename: str, width: int, height: int) -> QPixmap:
    """Charge une image du projet, redimensionnée proprement."""
    image_path = resolve_image_path(filename)
    if image_path is None:
        return QPixmap()
    pixmap = QPixmap(str(image_path))
    if pixmap.isNull():
        return QPixmap()
    return pixmap.scaled(width, height, Qt.KeepAspectRatio, Qt.SmoothTransformation)


def load_logo_pixmap(width: int, height: int) -> QPixmap:
    """Charge le logo principal (OIP.png), puis fallback."""
    pixmap = load_image_pixmap(PRIMARY_LOGO_FILENAME, width, height)
    if not pixmap.isNull():
        return pixmap
    return load_image_pixmap(PRIMARY_LOGO_FALLBACK, width, height)


def load_lightgroup_pixmap(width: int, height: int) -> QPixmap:
    """Charge le logo Light Group."""
    return load_image_pixmap(LIGHTGROUP_LOGO_FILENAME, width, height)


def ui_icon(
    widget: QWidget,
    fa_name: str,
    color: Optional[str] = None,
    fallback: Optional[QStyle.StandardPixmap] = None,
) -> QIcon:
    """Retourne une icône FontAwesome (qtawesome), ou une fallback Qt."""
    if FONTS_AVAILABLE and qta is not None:
        kwargs = {}
        if color:
            kwargs["color"] = color
        try:
            return qta.icon(fa_name, **kwargs)
        except Exception:
            pass
    if fallback is not None:
        return widget.style().standardIcon(fallback)
    return QIcon()


class PipelineWorker(QObject):
    """Exécute le pipeline DSF en thread de fond."""

    finished = Signal(object)
    failed = Signal(str)
    status = Signal(str)
    progress = Signal(int, str)

    def __init__(self, config: "DSFPipelineConfig"):
        super().__init__()
        self.config = config

    def _on_progress(self, percent: int, message: str):
        self.progress.emit(percent, message)
        self.status.emit(message)

    @Slot()
    def run(self):
        try:
            if DSFPipeline is None:
                raise RuntimeError("DSFPipeline indisponible")
            self.progress.emit(0, "Préparation de la génération")
            try:
                pipeline = DSFPipeline(self.config, progress_callback=self._on_progress)
            except TypeError:
                pipeline = DSFPipeline(self.config)
            artifacts = pipeline.run()
            self.finished.emit(artifacts)
        except Exception as exc:
            self.failed.emit(str(exc))

class SplashScreen(QWidget):
    """Splash screen avec progression réelle des étapes de démarrage."""

    def __init__(self, callback, startup_tasks=None, min_duration_ms: int = 5000):
        super().__init__()
        self.callback = callback
        self.startup_tasks = startup_tasks or []
        self.min_duration_ms = max(2000, min_duration_ms)
        self.current_task_index = 0
        self.started_at_ms = 0.0
        self.task_failed = False

        self.splash_logo_label: Optional[QLabel] = None
        self.splash_logo_base = QPixmap()
        self.logo_glow_effect: Optional[QGraphicsDropShadowEffect] = None
        self.logo_glow_timer = QTimer(self)
        self.logo_glow_timer.setInterval(40)
        self.logo_glow_timer.timeout.connect(self._animate_logo_glow)
        self.logo_glow_elapsed_ms = 0

        self.progress_bar_width = 360
        self.progress_value = 0
        self.progress_fill: Optional[QFrame] = None
        self.loading_label: Optional[QLabel] = None

        self.setWindowFlags(Qt.FramelessWindowHint | Qt.WindowStaysOnTopHint)
        self.setAttribute(Qt.WA_TranslucentBackground)

        screen = QApplication.primaryScreen().geometry()
        self.setGeometry(screen)

        self.opacity_animation = QPropertyAnimation(self, b"windowOpacity")
        self.opacity_animation.setDuration(700)
        self.opacity_animation.setStartValue(0.0)
        self.opacity_animation.setEndValue(1.0)
        self.opacity_animation.setEasingCurve(QEasingCurve.OutCubic)

        self.setup_ui()
        self.start_animations()

    def setup_ui(self):
        """Configure l'interface du splash screen."""
        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setAlignment(Qt.AlignCenter)

        container = QFrame()
        container.setStyleSheet(f"""
            QFrame {{
                background: qlineargradient(x1:0, y1:0, x2:1, y2:1,
                    stop:0 #1a2320,
                    stop:0.55 #25342f,
                    stop:1 #31443d);
                border-radius: 0px;
            }}
        """)
        container.setFixedSize(self.width(), self.height())

        container_layout = QVBoxLayout(container)
        container_layout.setAlignment(Qt.AlignCenter)
        container_layout.setSpacing(18)

        logo_container = QFrame()
        logo_layout = QVBoxLayout(logo_container)
        logo_layout.setAlignment(Qt.AlignCenter)
        logo_layout.setSpacing(10)

        logo_card = QFrame()
        logo_card.setStyleSheet("""
            QFrame {
                background: transparent;
                border: none;
            }
        """)
        logo_card_layout = QVBoxLayout(logo_card)
        logo_card_layout.setContentsMargins(0, 0, 0, 0)
        logo_card_layout.setAlignment(Qt.AlignCenter)

        self.splash_logo_base = load_logo_pixmap(185, 195)
        self.splash_logo_label = QLabel()
        self.splash_logo_label.setFixedSize(185, 195)
        self.splash_logo_label.setAlignment(Qt.AlignCenter)

        if not self.splash_logo_base.isNull():
            self.splash_logo_label.setPixmap(self.splash_logo_base)
            self.splash_logo_label.setStyleSheet("background: transparent;")
            self.logo_glow_effect = QGraphicsDropShadowEffect(self.splash_logo_label)
            self.logo_glow_effect.setOffset(0, 0)
            self.logo_glow_effect.setBlurRadius(26)
            glow_color = QColor("#f0c419")
            glow_color.setAlpha(130)
            self.logo_glow_effect.setColor(glow_color)
            self.splash_logo_label.setGraphicsEffect(self.logo_glow_effect)
        else:
            self.splash_logo_label.setText(GULFCAM_LOGO_TEXT)
            self.splash_logo_label.setStyleSheet("""
                QLabel {
                    color: white;
                    font-size: 56px;
                    font-weight: 800;
                    letter-spacing: 6px;
                    background: transparent;
                }
            """)

        tagline_label = QLabel(GULFCAM_TAGLINE)
        tagline_label.setStyleSheet("""
            QLabel {
                color: rgba(255,255,255,0.95);
                font-size: 24px;
                font-weight: 700;
                letter-spacing: 1px;
                background: transparent;
            }
        """)
        tagline_label.setAlignment(Qt.AlignCenter)

        logo_card_layout.addWidget(self.splash_logo_label, 0, Qt.AlignCenter)
        logo_layout.addWidget(logo_card, 0, Qt.AlignCenter)
        logo_layout.addWidget(tagline_label, 0, Qt.AlignCenter)

        footer_widget = QWidget()
        footer_layout = QHBoxLayout(footer_widget)
        footer_layout.setContentsMargins(0, 0, 0, 0)
        footer_layout.setSpacing(12)
        footer_layout.setAlignment(Qt.AlignCenter)

        dsf_by_label = QLabel("DSF by")
        dsf_by_label.setStyleSheet("""
            QLabel {
                color: rgba(255,255,255,0.98);
                font-size: 34px;
                font-weight: 700;
                letter-spacing: 2px;
                background: transparent;
            }
        """)

        lightgroup_badge = QFrame()
        lightgroup_badge.setStyleSheet("""
            QFrame {
                background: transparent;
                border: none;
            }
        """)
        lightgroup_badge_layout = QHBoxLayout(lightgroup_badge)
        lightgroup_badge_layout.setContentsMargins(0, 0, 0, 0)
        lightgroup_badge_layout.setAlignment(Qt.AlignCenter)

        lightgroup_footer_logo = QLabel()
        lightgroup_footer_logo.setAlignment(Qt.AlignCenter)
        lightgroup_footer_pixmap = load_lightgroup_pixmap(150, 60)
        if not lightgroup_footer_pixmap.isNull():
            lightgroup_footer_logo.setPixmap(lightgroup_footer_pixmap)
        else:
            lightgroup_footer_logo.setText("LIGHTGROUP")
            lightgroup_footer_logo.setStyleSheet("""
                QLabel {
                    color: rgba(255,255,255,0.95);
                    font-size: 20px;
                    font-weight: 700;
                    letter-spacing: 2px;
                    background: transparent;
                }
            """)
        lightgroup_badge_layout.addWidget(lightgroup_footer_logo)

        footer_layout.addWidget(dsf_by_label)
        footer_layout.addWidget(lightgroup_badge)

        self.progress_bar = QFrame()
        self.progress_bar.setFixedSize(self.progress_bar_width, 6)
        self.progress_bar.setStyleSheet("""
            QFrame {
                background-color: rgba(255,255,255,0.22);
                border-radius: 3px;
            }
        """)

        self.progress_fill = QFrame(self.progress_bar)
        self.progress_fill.setFixedSize(0, 6)
        self.progress_fill.setStyleSheet(f"""
            QFrame {{
                background-color: {GULFCAM_COLORS['accent']};
                border-radius: 3px;
            }}
        """)

        self.loading_label = QLabel("Initialisation...")
        self.loading_label.setStyleSheet("""
            QLabel {
                color: rgba(255,255,255,0.82);
                font-size: 15px;
                font-weight: 500;
                background: transparent;
            }
        """)
        self.loading_label.setAlignment(Qt.AlignCenter)

        container_layout.addWidget(logo_container, 0, Qt.AlignCenter)
        container_layout.addSpacing(24)
        container_layout.addWidget(footer_widget, 0, Qt.AlignCenter)
        container_layout.addSpacing(44)
        container_layout.addWidget(self.progress_bar, 0, Qt.AlignCenter)
        container_layout.addWidget(self.loading_label, 0, Qt.AlignCenter)

        layout.addWidget(container)

    def start_animations(self):
        """Démarre les animations et le chargement réel."""
        self.opacity_animation.start()
        self.logo_glow_elapsed_ms = 0
        self.logo_glow_timer.start()
        self.started_at_ms = time.time() * 1000.0
        self._set_progress(3)
        QTimer.singleShot(80, self._run_next_task)

    def _set_progress(self, value: int):
        """Met à jour la barre de progression."""
        self.progress_value = max(0, min(100, value))
        fill_width = int((self.progress_bar_width * self.progress_value) / 100)
        if self.progress_fill is not None:
            self.progress_fill.setGeometry(0, 0, fill_width, 6)

    def _run_next_task(self):
        """Exécute les tâches de démarrage en séquence."""
        if self.task_failed:
            return

        if self.current_task_index >= len(self.startup_tasks):
            self._finalize_startup()
            return

        step_name, step_fn = self.startup_tasks[self.current_task_index]
        if self.loading_label is not None:
            self.loading_label.setText(step_name)
        QApplication.processEvents()

        try:
            step_fn()
        except Exception as exc:
            self.task_failed = True
            if self.loading_label is not None:
                self.loading_label.setText(f"Erreur au démarrage: {exc}")
            QTimer.singleShot(1200, self.close_splash)
            return

        self.current_task_index += 1
        task_ratio = self.current_task_index / max(1, len(self.startup_tasks))
        self._set_progress(int(5 + (task_ratio * 90)))
        QTimer.singleShot(60, self._run_next_task)

    def _finalize_startup(self):
        """Attend la durée minimale (5s), puis termine à 100%."""
        elapsed = (time.time() * 1000.0) - self.started_at_ms
        remaining = max(0, int(self.min_duration_ms - elapsed))

        if self.loading_label is not None:
            self.loading_label.setText("Chargement de l'application...")

        if remaining > 0:
            self._set_progress(97)
            QTimer.singleShot(remaining, self._finish_and_close)
        else:
            self._finish_and_close()

    def _finish_and_close(self):
        self._set_progress(100)
        if self.loading_label is not None:
            self.loading_label.setText("Application prête")
        QTimer.singleShot(150, self.close_splash)

    def _animate_logo_glow(self):
        """Fait varier un halo lumineux autour du logo."""
        self.logo_glow_elapsed_ms += self.logo_glow_timer.interval()
        if self.logo_glow_effect is None:
            return

        period_ms = 1050.0
        wave = (math.sin((2.0 * math.pi * self.logo_glow_elapsed_ms) / period_ms) + 1.0) / 2.0
        blur = 18.0 + (22.0 * wave)
        alpha = int(70 + (95 * wave))

        glow_color = QColor(GULFCAM_COLORS["accent"])
        glow_color.setAlpha(alpha)
        self.logo_glow_effect.setBlurRadius(blur)
        self.logo_glow_effect.setColor(glow_color)
        self.logo_glow_effect.setOffset(0, 0)

    def close_splash(self):
        """Ferme le splash screen et affiche l'application."""
        if self.logo_glow_timer.isActive():
            self.logo_glow_timer.stop()
        self.close()
        self.callback()

class ModernCard(QFrame):
    """Carte moderne avec effets visuels et branding GULFCAM"""
    
    def __init__(self, parent=None, colors=None):
        super().__init__(parent)
        self.colors = colors or GULFCAM_COLORS
        self.setup_ui()
        
    def setup_ui(self):
        self.setFrameShape(QFrame.NoFrame)
        
        # Ombre portée
        shadow = QGraphicsDropShadowEffect()
        shadow.setBlurRadius(20)
        shadow_color = QColor(self.colors['primary'])
        shadow_color.setAlpha(32)
        shadow.setColor(shadow_color)
        shadow.setOffset(0, 4)
        self.setGraphicsEffect(shadow)
        
        self.update_style()
        
    def update_style(self):
        """Met à jour le style avec les couleurs actuelles"""
        self.setStyleSheet(f"""
            ModernCard {{
                background-color: {self.colors['card-bg']};
                border-radius: 16px;
                border: 1px solid {self.colors['light-gray']};
            }}
        """)
        
    def set_colors(self, colors):
        """Met à jour les couleurs"""
        self.colors = colors
        self.update_style()

class AnimatedButton(QPushButton):
    """Bouton avec animations au survol et couleurs GULFCAM"""
    
    def __init__(self, text="", parent=None, primary=True, colors=None):
        super().__init__(text, parent)
        self.colors = colors or GULFCAM_COLORS
        self.primary = primary
        self.setCursor(Qt.PointingHandCursor)
        self.setMinimumHeight(44)
        self.setMinimumWidth(120)
        self.update_style()
        
    def update_style(self):
        """Met à jour le style selon le type de bouton"""
        if self.primary:
            self.setStyleSheet(f"""
                QPushButton {{
                    background-color: {self.colors['primary']};
                    color: white;
                    border: none;
                    border-radius: 8px;
                    padding: 12px 24px;
                    font-weight: 600;
                    font-size: 14px;
                }}
                QPushButton:hover {{
                    background-color: {self.colors['primary_light']};
                }}
                QPushButton:pressed {{
                    background-color: {self.colors['secondary']};
                }}
                QPushButton:disabled {{
                    background-color: {self.colors['light-gray']};
                    color: {self.colors['gray']};
                }}
            """)
        else:
            self.setStyleSheet(f"""
                QPushButton {{
                    background-color: transparent;
                    color: {self.colors['primary']};
                    border: 2px solid {self.colors['primary']};
                    border-radius: 8px;
                    padding: 12px 24px;
                    font-weight: 600;
                    font-size: 14px;
                }}
                QPushButton:hover {{
                    background-color: {self.colors['primary']}10;
                    border-color: {self.colors['primary_light']};
                }}
                QPushButton:pressed {{
                    background-color: {self.colors['primary']}20;
                }}
            """)
            
    def set_colors(self, colors):
        """Met à jour les couleurs"""
        self.colors = colors
        self.update_style()

class ExcelIntegratedWidget(QWidget):
    """Widget Excel avec intégration forcée et couleurs GULFCAM (VERSION FINALE CORRIGÉE)"""
    
    def __init__(self, parent=None, colors=None):
        super().__init__(parent)
        self.colors = colors or GULFCAM_COLORS
        self.excel = None
        self.workbook = None
        self.preview_workbook = None
        self.preview_sheet_name = None
        self.current_file = None
        self.is_modified = False
        self.excel_window_hwnd = None
        self.excel_container_widget = None
        self.embed_timer = None
        self.resize_timer = None
        self.retry_count = 0
        self.max_retries = 30
        self.rpc_error_count = 0
        self.max_rpc_errors = 3
        self.preview_max_rows = 350
        self.preview_max_cols = 80
        self._updating_table = False
        self.status_color_key = "success"
        self.use_windows_com = WINDOWS_COM_AVAILABLE
        self.excel_process_id = None
        self.spreadsheet_command = self._detect_spreadsheet_command()
        self.native_can_embed, self.native_embed_block_reason = self._detect_native_embed_support()
        self.external_editor_process = None
        self.embedded_process = None
        self.embedded_profile_dir: Optional[Path] = None
        self.embedded_window_id: Optional[int] = None
        self.embedded_qwindow: Optional[QWindow] = None
        self.embedded_container: Optional[QWidget] = None
        self.embedded_stderr_log: Optional[Path] = None
        self.embed_launch_candidates: List[List[str]] = []
        self.embed_launch_index = 0
        self.pending_embed_pid: Optional[int] = None
        self.pending_embed_file_name = ""
        self.pending_embed_tries = 0
        self.max_embed_tries = 90
        self.setup_ui()

    def _detect_spreadsheet_command(self) -> Optional[List[str]]:
        """Détecte un tableur disponible sur la machine."""
        for binary in ("libreoffice", "soffice"):
            if shutil.which(binary):
                return [binary]
        return None

    def _detect_native_embed_support(self) -> tuple[bool, str]:
        """Vérifie si l'intégration native d'une fenêtre tableur est possible."""
        if IS_WINDOWS:
            return True, ""
        display = os.environ.get("DISPLAY", "").strip()
        session_type = os.environ.get("XDG_SESSION_TYPE", "").strip().lower()
        wayland_display = os.environ.get("WAYLAND_DISPLAY", "").strip()
        if not display:
            if session_type == "wayland" or wayland_display:
                return False, "Session Wayland sans bridge X11 détecté"
            return False, "Aucun serveur d'affichage X11 détecté"
        return True, ""

    def _set_status(self, text: str, color_key: str):
        """Met à jour l'indicateur de statut."""
        self.status_color_key = color_key
        self.status_indicator.setText(text)
        self.status_indicator.setStyleSheet(
            f"color: {self.colors[color_key]}; font-weight: 600; font-size: 12px;"
        )

    def _apply_theme_styles(self):
        """Applique le thème courant à tous les widgets visibles."""
        self.toolbar.setStyleSheet(f"""
            QFrame {{
                background-color: {self.colors['light']};
                border-top-left-radius: 12px;
                border-top-right-radius: 12px;
                border-bottom: 2px solid {self.colors['primary']}20;
                padding: 8px;
            }}
        """)

        self.file_indicator.setStyleSheet(f"""
            QLabel {{
                color: {self.colors['primary']};
                font-weight: 600;
                font-size: 13px;
                padding: 4px 12px;
                background-color: {self.colors['primary']}10;
                border-radius: 16px;
            }}
        """)

        for btn in self.toolbar_buttons:
            btn.setStyleSheet(f"""
                QPushButton {{
                    background-color: transparent;
                    color: {self.colors['primary']};
                    border: 1px solid {self.colors['primary']}30;
                    border-radius: 6px;
                    padding: 5px 12px;
                    font-weight: 500;
                    font-size: 12px;
                }}
                QPushButton:hover {{
                    background-color: {self.colors['primary']}10;
                }}
                QPushButton:pressed {{
                    background-color: {self.colors['primary']}20;
                }}
            """)

        self.excel_container.setStyleSheet(f"""
            QFrame {{
                background-color: {self.colors['card-bg']};
                border: 2px solid {self.colors['primary']}18;
                border-bottom-left-radius: 12px;
                border-bottom-right-radius: 12px;
            }}
        """)

        self.placeholder.setStyleSheet(f"""
            QFrame {{
                background-color: {self.colors['primary']}08;
                border: 2px dashed {self.colors['primary']}30;
                border-radius: 8px;
            }}
        """)
        self.placeholder_title.setStyleSheet(
            f"color: {self.colors['primary']}; font-size: 20px; font-weight: 600;"
        )
        self.placeholder_desc.setStyleSheet(f"color: {self.colors['gray']}; font-size: 14px;")
        self.placeholder_tagline.setStyleSheet(
            f"color: {self.colors['secondary']}; font-size: 12px; font-style: italic;"
        )

        self.preview_toolbar.setStyleSheet(f"""
            QFrame {{
                background-color: {self.colors['card-bg']};
                border: 1px solid {self.colors['light-gray']};
                border-radius: 8px;
            }}
        """)
        self.native_embed_banner.setStyleSheet(f"""
            QFrame {{
                background-color: {self.colors['card-bg']};
                border: 1px solid {self.colors['light-gray']};
                border-radius: 8px;
            }}
        """)
        self.native_embed_title.setStyleSheet(
            f"color: {self.colors['primary']}; font-weight: 700;"
        )
        self.native_embed_info.setStyleSheet(f"color: {self.colors['gray']}; font-size: 12px;")
        self.native_embed_host.setStyleSheet(f"""
            QFrame {{
                background-color: {self.colors['light']};
                border: 1px solid {self.colors['light-gray']};
                border-radius: 10px;
            }}
        """)
        self.sheet_label.setStyleSheet(f"color: {self.colors['primary']}; font-weight: 600;")
        self.preview_info.setStyleSheet(f"color: {self.colors['gray']}; font-size: 12px;")

        self.sheet_selector.setStyleSheet(f"""
            QComboBox {{
                background-color: {self.colors['card-bg']};
                color: {self.colors['dark']};
                border: 1px solid {self.colors['light-gray']};
                border-radius: 6px;
                padding: 4px 8px;
                min-height: 28px;
            }}
            QComboBox::drop-down {{
                border: none;
                width: 24px;
            }}
            QComboBox QAbstractItemView {{
                background-color: {self.colors['card-bg']};
                color: {self.colors['dark']};
                border: 1px solid {self.colors['light-gray']};
                selection-background-color: {self.colors['primary']}35;
                selection-color: {self.colors['dark']};
            }}
        """)

        self.preview_table.setStyleSheet(f"""
            QTableWidget {{
                background-color: {self.colors['card-bg']};
                color: {self.colors['dark']};
                alternate-background-color: {self.colors['light']};
                gridline-color: {self.colors['light-gray']};
                border: 1px solid {self.colors['light-gray']};
                border-radius: 8px;
                selection-background-color: {self.colors['primary']}55;
                selection-color: white;
            }}
            QHeaderView::section {{
                background-color: {self.colors['light']};
                color: {self.colors['primary']};
                border: 1px solid {self.colors['light-gray']};
                border-left: none;
                border-top: none;
                padding: 6px;
                font-weight: 600;
            }}
            QTableCornerButton::section {{
                background-color: {self.colors['light']};
                border: 1px solid {self.colors['light-gray']};
            }}
        """)

    def _set_loaded_mode(self, loaded: bool):
        """Agrandit la zone de prévisualisation quand un fichier est chargé."""
        # HAUTEUR AUGMENTÉE à 800px pour mieux voir Excel
        self.excel_container.setMinimumHeight(800 if loaded else 420)

    def _clear_embedded_container(self):
        """Supprime le conteneur Qt qui héberge une fenêtre externe."""
        if self.embedded_container is not None:
            self.native_embed_layout.removeWidget(self.embedded_container)
            self.embedded_container.deleteLater()
            self.embedded_container = None
        self.embedded_qwindow = None
        self.embedded_window_id = None

    def _cleanup_embedded_profile(self):
        if self.embedded_profile_dir and self.embedded_profile_dir.exists():
            shutil.rmtree(self.embedded_profile_dir, ignore_errors=True)
        self.embedded_profile_dir = None
        if self.embedded_stderr_log and self.embedded_stderr_log.exists():
            try:
                self.embedded_stderr_log.unlink()
            except Exception:
                pass
        self.embedded_stderr_log = None

    def _build_embed_launch_candidates(self, file_path: Path) -> List[List[str]]:
        """Construit plusieurs variantes de commande pour maximiser la compatibilité."""
        if not self.spreadsheet_command:
            return []
        binary = self.spreadsheet_command[0]
        base = ["--nologo", "--norestore", "--nodefault", "--nolockcheck", "--calc", str(file_path.absolute())]

        candidates: List[List[str]] = []
        if self.embedded_profile_dir:
            candidates.append(
                [
                    binary,
                    f"-env:UserInstallation={self.embedded_profile_dir.as_uri()}",
                    "--nologo",
                    "--norestore",
                    "--nodefault",
                    "--nolockcheck",
                    "--calc",
                    str(file_path.absolute()),
                ]
            )
        candidates.append([binary] + base)
        if binary != "soffice" and shutil.which("soffice"):
            candidates.append(["soffice"] + base)
        return candidates

    def _read_embedded_launch_error(self) -> str:
        if not self.embedded_stderr_log or not self.embedded_stderr_log.exists():
            return ""
        try:
            lines = self.embedded_stderr_log.read_text(errors="ignore").splitlines()
        except Exception:
            return ""
        if not lines:
            return ""
        tail = " | ".join(line.strip() for line in lines[-4:] if line.strip())
        return tail[:500]

    def _launch_embed_candidate(self, cmd: List[str]) -> Optional[subprocess.Popen]:
        env = os.environ.copy()
        session_type = env.get("XDG_SESSION_TYPE", "").strip().lower()
        if session_type == "wayland" and env.get("DISPLAY"):
            # Force X11 backend sous Wayland pour permettre l'intégration de fenêtre.
            env.setdefault("GDK_BACKEND", "x11")
            env.setdefault("SAL_USE_VCLPLUGIN", "gtk3")

        if self.embedded_stderr_log and self.embedded_stderr_log.exists():
            try:
                self.embedded_stderr_log.unlink()
            except Exception:
                pass
        self.embedded_stderr_log = Path(tempfile.mkstemp(prefix="dsf_lo_err_", suffix=".log")[1])
        log_handle = None
        try:
            log_handle = open(self.embedded_stderr_log, "wb")
        except Exception:
            log_handle = subprocess.DEVNULL
        try:
            proc = subprocess.Popen(
                cmd,
                stdout=subprocess.DEVNULL,
                stderr=log_handle,
                env=env,
            )
            if log_handle not in (None, subprocess.DEVNULL):
                log_handle.close()
            return proc
        except Exception:
            try:
                if log_handle not in (None, subprocess.DEVNULL):
                    log_handle.close()
            except Exception:
                pass
            return None

    def _try_next_embed_candidate(self) -> bool:
        if self.embed_launch_index >= len(self.embed_launch_candidates):
            return False
        cmd = self.embed_launch_candidates[self.embed_launch_index]
        self.embed_launch_index += 1
        proc = self._launch_embed_candidate(cmd)
        if proc is None:
            return self._try_next_embed_candidate()
        self.embedded_process = proc
        self.pending_embed_pid = proc.pid
        self.pending_embed_tries = 0
        return True

    def _stop_embedded_process(self):
        """Arrête le processus de tableur intégré lancé par l'application."""
        if hasattr(self, "native_embed_timer"):
            self.native_embed_timer.stop()
        self.pending_embed_tries = 0
        self.pending_embed_pid = None
        self.pending_embed_file_name = ""
        self.embed_launch_candidates = []
        self.embed_launch_index = 0

        if self.embedded_process and self.embedded_process.poll() is None:
            try:
                self.embedded_process.terminate()
                self.embedded_process.wait(timeout=1.2)
            except Exception:
                try:
                    self.embedded_process.kill()
                except Exception:
                    pass
        self.embedded_process = None
        self._cleanup_embedded_profile()
        self._clear_embedded_container()

    def _find_window_with_wmctrl(self, file_name: str, pid: Optional[int]) -> Optional[int]:
        if not shutil.which("wmctrl"):
            return None
        try:
            output = subprocess.check_output(
                ["wmctrl", "-lp"], text=True, stderr=subprocess.DEVNULL
            )
        except Exception:
            return None

        file_name_l = file_name.lower()
        candidates = []
        for line in output.splitlines():
            parts = line.split(None, 4)
            if len(parts) < 5:
                continue
            wid_hex, _desktop, pid_str, _host, title = parts
            try:
                wid = int(wid_hex, 16)
            except ValueError:
                continue

            title_l = title.lower()
            if "libreoffice" not in title_l and "calc" not in title_l and file_name_l not in title_l:
                continue

            score = 0
            if file_name_l and file_name_l in title_l:
                score += 5
            if "libreoffice" in title_l or "calc" in title_l:
                score += 3
            if pid and pid_str.isdigit() and int(pid_str) == pid:
                score += 4
            candidates.append((score, wid))

        if not candidates:
            return None
        candidates.sort(key=lambda item: item[0], reverse=True)
        return candidates[0][1]

    def _find_window_with_xdotool(self, file_name: str, pid: Optional[int]) -> Optional[int]:
        if not shutil.which("xdotool"):
            return None
        queries: List[List[str]] = []
        if pid:
            queries.append(["xdotool", "search", "--all", "--pid", str(pid), "--name", "LibreOffice|Calc"])
            queries.append(["xdotool", "search", "--all", "--pid", str(pid), "--name", file_name])
        queries.append(["xdotool", "search", "--all", "--name", file_name])
        queries.append(["xdotool", "search", "--all", "--name", "LibreOffice|Calc"])

        for cmd in queries:
            try:
                output = subprocess.check_output(
                    cmd,
                    text=True,
                    stderr=subprocess.DEVNULL,
                )
            except Exception:
                continue
            ids = [line.strip() for line in output.splitlines() if line.strip().isdigit()]
            if ids:
                return int(ids[-1])
        return None

    def _find_spreadsheet_window_id(self, file_name: str, pid: Optional[int]) -> Optional[int]:
        """Trouve la fenêtre du tableur pour intégration Qt (X11)."""
        if not file_name:
            return None
        wid = self._find_window_with_wmctrl(file_name=file_name, pid=pid)
        if wid is not None:
            return wid
        return self._find_window_with_xdotool(file_name=file_name, pid=pid)

    def _fallback_to_external_spreadsheet(self, reason: str) -> bool:
        """Ouvre le tableur en externe si l'intégration native échoue."""
        opened = self.open_in_external_spreadsheet(show_message=False)
        if opened:
            self._set_status("Tableur externe actif", "warning")
            self.native_embed_info.setText(f"Intégration non aboutie: {reason}. Tableur ouvert en externe.")
        return opened

    def _attach_embedded_window(self, window_id: int) -> bool:
        """Attache la fenêtre tableur externe dans la zone UI."""
        try:
            self._clear_embedded_container()
            foreign = QWindow.fromWinId(window_id)
            if foreign is None:
                return False
            container = QWidget.createWindowContainer(foreign, self.native_embed_host)
            container.setFocusPolicy(Qt.StrongFocus)
            container.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
            self.native_embed_layout.addWidget(container, 1)
            self.embedded_qwindow = foreign
            self.embedded_container = container
            self.embedded_window_id = window_id
            self.content_stack.setCurrentWidget(self.native_embed_page)
            return True
        except Exception:
            return False

    def _poll_native_embed(self):
        """Attend la création de la fenêtre LibreOffice puis l'intègre."""
        if not self.current_file:
            self.native_embed_timer.stop()
            return

        self.pending_embed_tries += 1
        if self.embedded_process and self.embedded_process.poll() is not None and self.embedded_window_id is None:
            if self._try_next_embed_candidate():
                self.native_embed_info.setText("Nouvelle tentative de lancement du tableur...")
                return
            self.native_embed_timer.stop()
            self._set_status("Échec lancement tableur", "danger")
            error_details = self._read_embedded_launch_error()
            reason = "Le tableur s'est fermé avant l'intégration"
            if error_details:
                reason = f"{reason} ({error_details})"
            self._fallback_to_external_spreadsheet(reason)
            self._fallback_to_local_preview(reason)
            return

        window_id = self._find_spreadsheet_window_id(
            file_name=self.pending_embed_file_name,
            pid=self.pending_embed_pid,
        )
        if window_id is not None and self._attach_embedded_window(window_id):
            self.native_embed_timer.stop()
            self._set_status("Tableur intégré", "success")
            self.native_embed_info.setText(
                "Fenêtre tableur intégrée. Styles/formules/liens Excel conservés."
            )
            self.is_modified = False
            return

        if self.pending_embed_tries >= self.max_embed_tries:
            self.native_embed_timer.stop()
            reason = "Fenêtre tableur non détectée dans le délai"
            if self.embedded_process and self.embedded_process.poll() is None:
                self._set_status("Tableur ouvert (non intégré)", "warning")
                self.native_embed_info.setText(
                    "Le tableur est lancé mais non intégrable dans cette session graphique."
                )
                self._load_preview_workbook(self.current_file)
                return
            self._set_status("Intégration impossible", "danger")
            self._fallback_to_external_spreadsheet(reason)
            self._fallback_to_local_preview(reason)

    def _fallback_to_local_preview(self, reason: str) -> bool:
        """Fallback contrôlé vers la table locale si l'intégration native échoue."""
        if not self.current_file:
            return False
        self._stop_embedded_process()
        self.native_embed_info.setText(f"Intégration native indisponible: {reason}.")
        ok = self._load_preview_workbook(self.current_file)
        if ok:
            self._set_status("Prévisualisation locale active", "warning")
        return ok

    def _start_native_embed(self, file_path: Path) -> bool:
        """Lance LibreOffice Calc et intègre sa fenêtre dans l'UI Linux."""
        if not self.spreadsheet_command:
            self._set_status("Aucun tableur détecté", "danger")
            return self._fallback_to_local_preview("LibreOffice/soffice non installé")
        if not self.native_can_embed:
            self._set_status("Intégration native indisponible", "warning")
            self._fallback_to_external_spreadsheet(self.native_embed_block_reason)
            return self._fallback_to_local_preview(self.native_embed_block_reason)

        self._stop_embedded_process()
        profile_dir = Path(tempfile.mkdtemp(prefix="dsf_lo_embed_"))
        self.embedded_profile_dir = profile_dir

        try:
            self.embed_launch_candidates = self._build_embed_launch_candidates(file_path)
            self.embed_launch_index = 0
            if not self._try_next_embed_candidate():
                raise RuntimeError("Aucune commande de lancement valide n'a pu démarrer")
            self.pending_embed_file_name = file_path.name
            self.pending_embed_tries = 0
            self.native_embed_info.setText("Ouverture du tableur en cours...")
            self.content_stack.setCurrentWidget(self.native_embed_page)
            self._set_loaded_mode(True)
            self._set_status("Lancement du tableur...", "warning")
            self.native_embed_timer.start()
            return True
        except Exception as e:
            self._set_status("Ouverture tableur échouée", "danger")
            self._cleanup_embedded_profile()
            return self._fallback_to_local_preview(str(e))

    def _is_native_embed_active(self) -> bool:
        return (
            not self.use_windows_com
            and self.embedded_process is not None
            and self.embedded_process.poll() is None
            and self.embedded_window_id is not None
        )

    def _trigger_native_save(self) -> bool:
        """Déclenche Ctrl+S dans la fenêtre intégrée si xdotool est disponible."""
        if not self._is_native_embed_active() or not self.embedded_window_id:
            return False
        if not shutil.which("xdotool"):
            return False
        try:
            subprocess.run(
                ["xdotool", "key", "--window", str(self.embedded_window_id), "ctrl+s"],
                check=True,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
            )
            return True
        except Exception:
            return False

    def open_in_external_spreadsheet(self, show_message: bool = True) -> bool:
        """Ouvre le fichier dans un tableur externe (Linux/macOS fallback)."""
        if not self.current_file or not self.current_file.exists():
            if show_message:
                QMessageBox.warning(self, "Erreur GULFCAM", "Aucun fichier valide à ouvrir.")
            return False

        file_path = str(self.current_file.absolute())
        try:
            if IS_WINDOWS and hasattr(os, "startfile"):
                os.startfile(file_path)
            elif self.spreadsheet_command:
                self.external_editor_process = subprocess.Popen(
                    self.spreadsheet_command + ["--calc", file_path]
                )
            elif shutil.which("xdg-open"):
                self.external_editor_process = subprocess.Popen(["xdg-open", file_path])
            elif shutil.which("open"):
                self.external_editor_process = subprocess.Popen(["open", file_path])
            else:
                raise RuntimeError("Aucun tableur détecté. Installez LibreOffice (libreoffice-calc).")

            self._set_status("Tableur externe ouvert", "success")
            if show_message:
                QMessageBox.information(
                    self,
                    "GULFCAM",
                    "Le fichier est ouvert dans le tableur externe.\n"
                    "Utilisez Ctrl+S dans le tableur pour enregistrer.",
                )
            return True
        except Exception as e:
            self._set_status("Ouverture impossible", "danger")
            if show_message:
                QMessageBox.critical(self, "Erreur GULFCAM", f"Impossible d'ouvrir le tableur:\n{str(e)}")
            return False

    def _load_preview_workbook(self, file_path: Path) -> bool:
        """Charge le classeur en mode prévisualisation/édition intégrée."""
        try:
            from openpyxl import load_workbook as ox_load_workbook

            if self.preview_workbook is not None:
                try:
                    self.preview_workbook.close()
                except Exception:
                    pass

            self.preview_workbook = ox_load_workbook(file_path, data_only=False)
            sheet_names = list(self.preview_workbook.sheetnames)
            if not sheet_names:
                raise RuntimeError("Le fichier ne contient aucune feuille.")

            self.sheet_selector.blockSignals(True)
            self.sheet_selector.clear()
            self.sheet_selector.addItems(sheet_names)
            self.sheet_selector.blockSignals(False)
            self.sheet_selector.setCurrentIndex(0)
            self._populate_preview_table(sheet_names[0])

            self.content_stack.setCurrentWidget(self.preview_page)
            self._set_loaded_mode(True)
            self.is_modified = False
            return True
        except Exception as e:
            self.preview_workbook = None
            self.content_stack.setCurrentWidget(self.placeholder)
            self._set_loaded_mode(False)
            QMessageBox.critical(self, "Erreur GULFCAM", f"Prévisualisation impossible:\n{str(e)}")
            return False

    def _on_sheet_changed(self, sheet_name: str):
        if not sheet_name:
            return
        self._populate_preview_table(sheet_name)

    def _populate_preview_table(self, sheet_name: str):
        """Affiche une feuille dans la grille éditable."""
        if self.preview_workbook is None or sheet_name not in self.preview_workbook.sheetnames:
            return

        ws = self.preview_workbook[sheet_name]
        self.preview_sheet_name = sheet_name

        max_row = max(1, min(ws.max_row or 1, self.preview_max_rows))
        max_col = max(1, min(ws.max_column or 1, self.preview_max_cols))

        from openpyxl.utils import get_column_letter

        self._updating_table = True
        self.preview_table.clear()
        self.preview_table.setRowCount(max_row)
        self.preview_table.setColumnCount(max_col)
        self.preview_table.setHorizontalHeaderLabels([get_column_letter(i) for i in range(1, max_col + 1)])
        self.preview_table.setVerticalHeaderLabels([str(i) for i in range(1, max_row + 1)])

        for r in range(1, max_row + 1):
            for c in range(1, max_col + 1):
                value = ws.cell(row=r, column=c).value
                if value is None:
                    continue
                item = QTableWidgetItem(str(value))
                self.preview_table.setItem(r - 1, c - 1, item)

        self._updating_table = False

        clipped = ""
        if ws.max_row > self.preview_max_rows or ws.max_column > self.preview_max_cols:
            clipped = f" (affichage limité à {self.preview_max_rows} x {self.preview_max_cols})"
        self.preview_info.setText(
            f"{sheet_name} • {ws.max_row} lignes x {ws.max_column} colonnes{clipped}"
        )

    def _coerce_cell_value(self, text: str):
        value = text.strip()
        if value == "":
            return None
        if value.startswith("="):
            return value
        try:
            if value.isdigit() or (value.startswith("-") and value[1:].isdigit()):
                return int(value)
            normalized = value.replace(",", ".")
            return float(normalized)
        except Exception:
            return value

    def _on_preview_item_changed(self, item: QTableWidgetItem):
        """Synchronise les modifications de la grille vers le workbook."""
        if self._updating_table or self.preview_workbook is None or not self.preview_sheet_name:
            return

        try:
            ws = self.preview_workbook[self.preview_sheet_name]
            row = item.row() + 1
            col = item.column() + 1
            ws.cell(row=row, column=col).value = self._coerce_cell_value(item.text())
            self.is_modified = True
            self._set_status("Modifications locales", "warning")
        except Exception as e:
            QMessageBox.warning(self, "Erreur GULFCAM", f"Impossible d'écrire la cellule:\n{str(e)}")
        
    def setup_ui(self):
        """Configure l'interface Excel avec couleurs GULFCAM"""
        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(0)
        
        # Barre d'outils Excel
        self.toolbar = QFrame()
        self.toolbar.setMinimumHeight(68)
        toolbar_layout = QHBoxLayout(self.toolbar)
        toolbar_layout.setContentsMargins(12, 10, 12, 10)
        toolbar_layout.setSpacing(10)
        
        # Indicateur de fichier
        self.file_indicator = QLabel("Aucun fichier chargé")
        self.file_indicator.setMinimumHeight(34)
        self.file_indicator.setAlignment(Qt.AlignVCenter | Qt.AlignLeft)
        
        # Statut
        self.status_indicator = QLabel("Prêt")
        self.status_indicator.setStyleSheet(f"""
            QLabel {{
                color: {self.colors['success']};
                font-weight: 600;
                font-size: 12px;
            }}
        """)
        self.status_indicator.setMinimumHeight(30)
        self.status_indicator.setAlignment(Qt.AlignVCenter | Qt.AlignLeft)
        
        # Boutons
        self.save_btn = QPushButton("Enregistrer")
        self.save_btn.setIcon(
            ui_icon(
                self,
                "fa5s.save",
                color=self.colors["primary"],
                fallback=QStyle.StandardPixmap.SP_DialogSaveButton,
            )
        )
        self.save_btn.setCursor(Qt.PointingHandCursor)
        self.save_btn.clicked.connect(self.save_changes)
        
        self.refresh_btn = QPushButton("Recharger")
        self.refresh_btn.setIcon(
            ui_icon(
                self,
                "fa5s.sync-alt",
                color=self.colors["primary"],
                fallback=QStyle.StandardPixmap.SP_BrowserReload,
            )
        )
        self.refresh_btn.setCursor(Qt.PointingHandCursor)
        self.refresh_btn.clicked.connect(self.refresh)
        
        embed_label = "Forcer intégration" if self.use_windows_com else "Intégrer le tableur"
        self.embed_btn = QPushButton(embed_label)
        embed_fa_icon = "fa5s.object-group" if self.use_windows_com else "fa5s.window-maximize"
        embed_fallback = (
            QStyle.StandardPixmap.SP_CommandLink if self.use_windows_com else QStyle.StandardPixmap.SP_DialogOpenButton
        )
        self.embed_btn.setIcon(
            ui_icon(
                self,
                embed_fa_icon,
                color=self.colors["primary"],
                fallback=embed_fallback,
            )
        )
        self.embed_btn.setCursor(Qt.PointingHandCursor)
        self.embed_btn.clicked.connect(self.force_embed_excel_now)
        self.toolbar_buttons = [self.save_btn, self.refresh_btn, self.embed_btn]
        for btn in self.toolbar_buttons:
            btn.setMinimumHeight(34)
            btn.setIconSize(QSize(16, 16))
        
        toolbar_layout.addWidget(self.file_indicator)
        toolbar_layout.addSpacing(20)
        toolbar_layout.addWidget(self.status_indicator)
        toolbar_layout.addStretch()
        toolbar_layout.addWidget(self.save_btn)
        toolbar_layout.addWidget(self.refresh_btn)
        toolbar_layout.addWidget(self.embed_btn)
        
        # Conteneur Excel - HAUTEUR AUGMENTÉE
        self.excel_container = QFrame()
        self.excel_container.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
        self.excel_container.setMinimumHeight(800)  # Augmenté de 600 à 800
        
        # Layout du conteneur
        container_layout = QVBoxLayout(self.excel_container)
        container_layout.setContentsMargins(0, 0, 0, 0)
        
        # Message placeholder
        self.placeholder = QFrame()
        
        placeholder_layout = QVBoxLayout(self.placeholder)
        placeholder_layout.setAlignment(Qt.AlignCenter)
        placeholder_layout.setSpacing(15)
        
        title_text = "GULFCAM Excel Intégré" if self.use_windows_com else "GULFCAM Tableur Linux"
        self.placeholder_title = QLabel(title_text)
        self.placeholder_title.setAlignment(Qt.AlignCenter)
        
        desc_text = (
            "Déposez votre fichier Excel ici\nou utilisez le bouton \"Charger\""
            if self.use_windows_com
            else "Chargez un fichier .xlsx pour l'afficher dans le tableur intégré\n(et conserver styles, formules et liaisons)"
        )
        self.placeholder_desc = QLabel(desc_text)
        self.placeholder_desc.setAlignment(Qt.AlignCenter)
        
        self.placeholder_tagline = QLabel(GULFCAM_TAGLINE)
        self.placeholder_tagline.setAlignment(Qt.AlignCenter)
        
        placeholder_layout.addWidget(self.placeholder_title)
        placeholder_layout.addWidget(self.placeholder_desc)
        placeholder_layout.addWidget(self.placeholder_tagline)

        # Vue d'intégration native Linux (LibreOffice)
        self.native_embed_page = QFrame()
        native_layout = QVBoxLayout(self.native_embed_page)
        native_layout.setContentsMargins(6, 6, 6, 6)
        native_layout.setSpacing(8)

        self.native_embed_banner = QFrame()
        native_banner_layout = QHBoxLayout(self.native_embed_banner)
        native_banner_layout.setContentsMargins(10, 6, 10, 6)
        native_banner_layout.setSpacing(10)
        self.native_embed_title = QLabel("Fenêtre tableur intégrée")
        self.native_embed_info = QLabel("Prêt à intégrer le fichier")
        native_banner_layout.addWidget(self.native_embed_title)
        native_banner_layout.addWidget(self.native_embed_info, 1)

        self.native_embed_host = QFrame()
        self.native_embed_layout = QVBoxLayout(self.native_embed_host)
        self.native_embed_layout.setContentsMargins(0, 0, 0, 0)
        self.native_embed_layout.setSpacing(0)

        # Vue d'édition intégrée (fallback Linux/macOS)
        self.preview_page = QFrame()
        preview_layout = QVBoxLayout(self.preview_page)
        preview_layout.setContentsMargins(6, 6, 6, 6)
        preview_layout.setSpacing(8)

        self.preview_toolbar = QFrame()
        preview_toolbar_layout = QHBoxLayout(self.preview_toolbar)
        preview_toolbar_layout.setContentsMargins(10, 6, 10, 6)
        preview_toolbar_layout.setSpacing(10)

        self.sheet_label = QLabel("Feuille :")
        self.sheet_selector = QComboBox()
        self.sheet_selector.setMinimumWidth(220)
        self.sheet_selector.currentTextChanged.connect(self._on_sheet_changed)

        self.preview_info = QLabel("Prévisualisation locale")

        preview_toolbar_layout.addWidget(self.sheet_label)
        preview_toolbar_layout.addWidget(self.sheet_selector)
        preview_toolbar_layout.addWidget(self.preview_info, 1)

        self.preview_table = QTableWidget()
        self.preview_table.setAlternatingRowColors(True)
        self.preview_table.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
        self.preview_table.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive)
        self.preview_table.horizontalHeader().setDefaultSectionSize(130)
        self.preview_table.verticalHeader().setDefaultSectionSize(24)
        self.preview_table.itemChanged.connect(self._on_preview_item_changed)

        preview_layout.addWidget(self.preview_toolbar)
        preview_layout.addWidget(self.preview_table, 1)

        self.content_stack = QStackedWidget()
        
        # Page dédiée pour l'intégration COM (Windows)
        self.excel_com_page = QWidget()
        self.excel_com_layout = QVBoxLayout(self.excel_com_page)
        self.excel_com_layout.setContentsMargins(0, 0, 0, 0)
        
        self.content_stack.addWidget(self.placeholder)
        self.content_stack.addWidget(self.native_embed_page)
        self.content_stack.addWidget(self.preview_page)
        self.content_stack.addWidget(self.excel_com_page)
        self.content_stack.setCurrentWidget(self.placeholder)

        native_layout.addWidget(self.native_embed_banner)
        native_layout.addWidget(self.native_embed_host, 1)

        container_layout.addWidget(self.content_stack, 1)
        
        layout.addWidget(self.toolbar)
        layout.addWidget(self.excel_container, 1)
        self._apply_theme_styles()
        self._set_status("Prêt", self.status_color_key)
        
        # Timer pour l'intégration COM
        if self.use_windows_com:
            self.embed_timer = QTimer()
            self.embed_timer.timeout.connect(self._attempt_excel_integration)
            self.resize_timer = QTimer()
            self.resize_timer.setInterval(100)
            self.resize_timer.timeout.connect(self._force_excel_resize)
            
        self.native_embed_timer = QTimer(self)
        self.native_embed_timer.setInterval(350)
        self.native_embed_timer.timeout.connect(self._poll_native_embed)

    def _close_all_other_excel_windows(self, keep_hwnd=None):
        """Ferme toutes les fenêtres Excel sauf celle à garder"""
        if not win32gui:
            return
            
        def enum_windows_callback(hwnd, ctx):
            try:
                # Vérifier si c'est une fenêtre Excel
                class_name = win32gui.GetClassName(hwnd)
                if class_name != "XLMAIN":
                    return True
                    
                # Ne pas fermer la fenêtre à garder
                if keep_hwnd and hwnd == keep_hwnd:
                    return True
                    
                # Vérifier si c'est une fenêtre sans fichier ou de récupération
                window_text = win32gui.GetWindowText(hwnd)
                if "récupéré" in window_text.lower() or "recovered" in window_text.lower():
                    print(f"Fermeture fenêtre de récupération: {window_text}")
                    win32gui.PostMessage(hwnd, win32con.WM_CLOSE, 0, 0)
                elif "Microsoft Excel" in window_text and not self.current_file:
                    # Fenêtre Excel vide
                    print(f"Fermeture fenêtre Excel vide: {window_text}")
                    win32gui.PostMessage(hwnd, win32con.WM_CLOSE, 0, 0)
                    
            except:
                pass
            return True
            
        win32gui.EnumWindows(enum_windows_callback, None)
        
    def force_embed_excel_now(self):
        """Force l'intégration d'Excel immédiatement"""
        if not self.use_windows_com:
            if not self.current_file or not self.current_file.exists():
                QMessageBox.warning(self, "Erreur GULFCAM", "Aucun fichier chargé.")
                return
            self._start_native_embed(self.current_file)
            return
        
        if not self.excel or not self.workbook:
            QMessageBox.warning(self, "Erreur GULFCAM", "Aucun fichier Excel chargé.")
            return
            
        self.retry_count = 0
        self.embed_timer.start(500)
        self._attempt_excel_integration()
        
    def _check_excel_connection(self) -> bool:
        """Vérifie si la connexion Excel est toujours valide et la réinitialise si nécessaire."""
        if self.excel is None:
            return False
            
        try:
            # Test simple pour vérifier la connexion
            _ = self.excel.Name
            self.rpc_error_count = 0
            return True
        except Exception as e:
            self.rpc_error_count += 1
            print(f"Erreur de connexion Excel (tentative {self.rpc_error_count}/{self.max_rpc_errors}): {e}")
            
            if self.rpc_error_count >= self.max_rpc_errors:
                print("Trop d'erreurs RPC, réinitialisation complète...")
                self._reset_excel_connection()
                return False
            
            # Attendre un peu avant de réessayer
            QTimer.singleShot(1000, lambda: self._check_excel_connection())
            return False
            
    def _reset_excel_connection(self):
        """Réinitialise complètement la connexion Excel."""
        try:
            # Nettoyer l'ancienne connexion
            if self.workbook:
                try:
                    self.workbook.Close(SaveChanges=False)
                except:
                    pass
                self.workbook = None
                
            if self.excel:
                try:
                    self.excel.Quit()
                except:
                    pass
                self.excel = None
                
            # Nettoyer les handles
            self.excel_window_hwnd = None
            if self.excel_container_widget:
                self.excel_com_layout.removeWidget(self.excel_container_widget)
                self.excel_container_widget.deleteLater()
                self.excel_container_widget = None
                
            # Réinitialiser le compteur
            self.rpc_error_count = 0
            
            # Recharger le fichier si nécessaire
            if self.current_file and self.current_file.exists():
                print("Tentative de reconnexion à Excel...")
                QTimer.singleShot(2000, lambda: self.load_file(self.current_file))
                
        except Exception as e:
            print(f"Erreur lors de la réinitialisation Excel: {e}")
            
    def _attempt_excel_integration(self):
        """Tente d'intégrer la fenêtre Excel (appelé par le timer)"""
        if self.excel_window_hwnd and self.excel_container_widget:
            # Déjà intégré, arrêter le timer et lancer le redimensionnement
            self.embed_timer.stop()
            if not self.resize_timer.isActive():
                self.resize_timer.start()
            return
            
        if not self.excel or not self.workbook:
            self.embed_timer.stop()
            return
            
        # Vérifier la connexion avant de continuer
        if not self._check_excel_connection():
            return
            
        try:
            # Essayer de rendre Excel visible (protégé car certains COM ne le permettent pas)
            try:
                self.excel.Visible = True
            except Exception:
                # Visible peut déjà être à True par défaut, ignorer l'erreur
                pass
            
            # Chercher la fenêtre Excel
            def find_excel_window(hwnd, ctx):
                try:
                    if win32gui.GetClassName(hwnd) == "XLMAIN":
                        ctx['hwnd'] = hwnd
                        return False
                except:
                    pass
                return True
                
            ctx = {'hwnd': None}
            win32gui.EnumWindows(find_excel_window, ctx)
            
            if not ctx['hwnd']:
                self.retry_count += 1
                if self.retry_count >= self.max_retries:
                    print("Fenêtre Excel introuvable après plusieurs tentatives")
                    self.embed_timer.stop()
                    self._set_status("Intégration échouée - prévisualisation locale", "warning")
                    self._fallback_to_preview_windows("fenêtre Excel introuvable")
                return
                
            excel_hwnd = ctx['hwnd']
            
            # Fermer toutes les autres fenêtres Excel (récupération, etc.)
            self._close_all_other_excel_windows(keep_hwnd=excel_hwnd)
            
            # Vérifier que la fenêtre n'est pas déjà intégrée
            if excel_hwnd == self.excel_window_hwnd and self.excel_container_widget:
                self.embed_timer.stop()
                if not self.resize_timer.isActive():
                    self.resize_timer.start()
                return
                
            # S'assurer que la fenêtre est restaurée
            win32gui.ShowWindow(excel_hwnd, win32con.SW_RESTORE)
            
            # CRITIQUE: Créer d'abord le conteneur Qt, puis faire le reparenting
            # 1. Créer la QWindow à partir du HWND
            self.excel_window = QWindow.fromWinId(excel_hwnd)
            
            # 2. Nettoyer l'ancien conteneur s'il existe
            if self.excel_container_widget:
                self.excel_com_layout.removeWidget(self.excel_container_widget)
                self.excel_container_widget.deleteLater()
                self.excel_container_widget = None
                
            # 3. Créer le nouveau conteneur
            self.excel_container_widget = QWidget.createWindowContainer(self.excel_window)
            self.excel_container_widget.setParent(self.excel_com_page)
            
            # 4. Nettoyer le layout et ajouter le nouveau conteneur
            while self.excel_com_layout.count():
                child = self.excel_com_layout.takeAt(0)
                if child.widget():
                    child.widget().deleteLater()
                    
            self.excel_com_layout.addWidget(self.excel_container_widget)
            
            # 5. Appliquer le style WS_CLIPCHILDREN
            try:
                container_hwnd = int(self.excel_com_page.winId())
                style = win32gui.GetWindowLong(container_hwnd, win32con.GWL_STYLE)
                win32gui.SetWindowLong(container_hwnd, win32con.GWL_STYLE, style | win32con.WS_CLIPCHILDREN)
            except:
                pass
                
            # 6. Configurer Excel (sans DisplayAlerts qui cause l'erreur)
            try:
                self.excel.DisplayFormulaBar = False
                self.excel.DisplayStatusBar = False
                self.excel.WindowState = -4137  # Maximized
            except:
                # Ignorer les erreurs de configuration mineures
                pass
            
            # 7. Activer la page
            self.content_stack.setCurrentWidget(self.excel_com_page)
            
            # 8. Stocker le HWND et arrêter le timer
            self.excel_window_hwnd = excel_hwnd
            self.retry_count = 0
            self.embed_timer.stop()
            
            # 9. Démarrer le timer de redimensionnement
            if not self.resize_timer.isActive():
                self.resize_timer.start()
            
            self._set_status("Excel intégré", "success")
            print(f"Excel intégré avec succès: {excel_hwnd}")
            
        except Exception as e:
            self.retry_count += 1
            if "RPC" in str(e) or "-2147023174" in str(e):
                # Erreur RPC, réinitialiser la connexion
                self._reset_excel_connection()
                
            if self.retry_count >= self.max_retries:
                print(f"Erreur intégration Excel: {e}")
                self.embed_timer.stop()
                self._set_status("Intégration échouée - prévisualisation locale", "warning")
                self._fallback_to_preview_windows(str(e))
                
    def _fallback_to_preview_windows(self, reason: str) -> bool:
        """Sur Windows: bascule vers la prévisualisation openpyxl si l'intégration COM échoue."""
        if not self.current_file or not self.current_file.exists():
            return False
        try:
            if self.workbook:
                try:
                    self.workbook.Close(SaveChanges=False)
                except Exception:
                    pass
                self.workbook = None
            if self.excel:
                try:
                    self.excel.Quit()
                except Exception:
                    pass
                self.excel = None
            self.excel_window_hwnd = None
            self.excel_container_widget = None
        except Exception:
            pass
        ok = self._load_preview_workbook(self.current_file)
        if ok:
            self._set_status("Prévisualisation locale (Excel non intégré)", "warning")
        return ok

    def _force_excel_resize(self):
        """Force le redimensionnement de la fenêtre Excel pour remplir tout l'espace"""
        if not self.excel_window_hwnd or not self.excel_container_widget:
            self.resize_timer.stop()
            return
            
        # Vérifier que la connexion est toujours valide
        if not self._check_excel_connection():
            self.resize_timer.stop()
            return
            
        try:
            # Obtenir la taille du conteneur
            container_size = self.excel_container_widget.size()
            if container_size.width() <= 10 or container_size.height() <= 10:
                # Si le conteneur n'a pas encore de taille valide, réessayer plus tard
                return
                
            # Obtenir la taille actuelle de la fenêtre Excel
            try:
                rect = win32gui.GetWindowRect(self.excel_window_hwnd)
                current_width = rect[2] - rect[0]
                current_height = rect[3] - rect[1]
            except:
                # Fenêtre peut-être détruite, réinitialiser
                self.excel_window_hwnd = None
                self.resize_timer.stop()
                return
            
            # Vérifier si un redimensionnement est nécessaire (marge de 5 pixels)
            if abs(current_width - container_size.width()) > 5 or abs(current_height - container_size.height()) > 5:
                # Redimensionner la fenêtre Excel pour remplir le conteneur
                try:
                    win32gui.SetWindowPos(
                        self.excel_window_hwnd,
                        win32con.HWND_TOP,
                        0, 0,
                        container_size.width(),
                        container_size.height(),
                        win32con.SWP_NOACTIVATE | win32con.SWP_SHOWWINDOW | win32con.SWP_FRAMECHANGED
                    )
                    
                    # Forcer le rafraîchissement
                    win32gui.InvalidateRect(self.excel_window_hwnd, None, True)
                    
                    print(f"Excel redimensionné à {container_size.width()}x{container_size.height()}")
                except Exception as e:
                    if "RPC" in str(e) or "-2147023174" in str(e):
                        # Erreur RPC, réinitialiser la connexion
                        self._reset_excel_connection()
                        self.resize_timer.stop()
                    else:
                        print(f"Erreur lors du redimensionnement: {e}")
                
        except Exception as e:
            print(f"Erreur dans _force_excel_resize: {e}")
                
    def resizeEvent(self, event):
        """Redimensionne Excel quand le widget change de taille"""
        super().resizeEvent(event)
        # Déclencher un redimensionnement après le changement de taille
        if self.use_windows_com and self.excel_window_hwnd:
            QTimer.singleShot(50, self._force_excel_resize)
        
    def showEvent(self, event):
        """Quand le widget devient visible"""
        super().showEvent(event)
        # Forcer un redimensionnement après l'affichage
        if self.use_windows_com and self.excel_window_hwnd:
            QTimer.singleShot(100, self._force_excel_resize)
        
    def load_file(self, file_path: Path) -> bool:
        """Charge un fichier Excel - GARANTIT QUE C'EST LE BON FICHIER QUI EST AFFICHÉ"""
        print(f"[DEBUG] load_file appelé avec: {file_path}")
        print(f"[DEBUG] WINDOWS_COM_AVAILABLE = {WINDOWS_COM_AVAILABLE}")
        
        try:
            if not file_path or not file_path.exists():
                print("[DEBUG] Fichier n'existe pas ou chemin vide!")
                return False
            
            self.current_file = file_path
            self.file_indicator.setText(file_path.name)

            if not self.use_windows_com:
                return self._start_native_embed(file_path)
            
            # Arrêter les timers existants
            if self.embed_timer and self.embed_timer.isActive():
                self.embed_timer.stop()
            if self.resize_timer and self.resize_timer.isActive():
                self.resize_timer.stop()
            
            # Initialiser/Vérifier Excel
            need_dispatch = False
            if self.excel is None:
                need_dispatch = True
            else:
                try:
                    # Test de connexion RPC actif
                    _ = self.excel.Name
                except Exception as e:
                    print(f"Connexion RPC perdue: {e}. Réinitialisation...")
                    need_dispatch = True
                    self.excel = None
                    self.workbook = None
                    self.excel_window_hwnd = None
                    self.excel_container_widget = None
            
            if need_dispatch:
                try:
                    # Réinitialiser le compteur d'erreurs
                    self.rpc_error_count = 0
                    self.excel = Dispatch("Excel.Application")
                    # Vérifier que l'objet est valide
                    _ = self.excel.Name
                    print("Nouvelle instance Excel créée avec succès")
                except Exception as e:
                    self.excel = None
                    QMessageBox.critical(self, "Erreur GULFCAM", f"Impossible de démarrer Excel :\n{str(e)}")
                    return False
                
            # Fermer l'ancien classeur s'il existe
            if self.workbook:
                try:
                    # Vérifier si c'est le même fichier
                    if hasattr(self.workbook, 'FullName') and self.workbook.FullName == str(file_path.absolute()):
                        print(f"Le fichier {file_path.name} est déjà ouvert")
                    else:
                        self.workbook.Close(SaveChanges=False)
                        self.workbook = None
                except Exception as e:
                    print(f"Erreur lors de la fermeture de l'ancien classeur: {e}")
                    self.workbook = None
                    
            # Ouvrir le nouveau fichier - GARANTIE QUE C'EST LE BON FICHIER
            try:
                self.excel.Visible = True
            except Exception:
                # Visible est souvent défini par défaut, ignorer les erreurs
                pass
            
            # Si le classeur est déjà ouvert avec le bon fichier, on le garde
            if self.workbook is None:
                try:
                    self.workbook = self.excel.Workbooks.Open(str(file_path.absolute()))
                    print(f"Fichier ouvert avec succès: {file_path.name}")
                except Exception as e:
                    QMessageBox.critical(self, "Erreur GULFCAM", f"Impossible d'ouvrir le fichier Excel:\n{str(e)}")
                    return False
            
            # Vérifier que c'est bien le bon fichier qui est ouvert
            if hasattr(self.workbook, 'FullName'):
                if self.workbook.FullName != str(file_path.absolute()):
                    print(f"ERREUR: Mauvais fichier ouvert! Attendu: {file_path}, Obtenu: {self.workbook.FullName}")
                    # Forcer la réouverture
                    try:
                        self.workbook.Close(SaveChanges=False)
                        self.workbook = self.excel.Workbooks.Open(str(file_path.absolute()))
                    except Exception as e:
                        QMessageBox.critical(self, "Erreur GULFCAM", f"Impossible d'ouvrir le bon fichier:\n{str(e)}")
                        return False
            
            # Fermer toutes les fenêtres Excel de récupération
            self._close_all_other_excel_windows()
            
            # Réinitialiser les handles
            self.excel_window_hwnd = None
            self.excel_container_widget = None
            self.retry_count = 0
            
            # Afficher la page d'intégration
            self.content_stack.setCurrentWidget(self.excel_com_page)
            self._set_loaded_mode(True)
            
            # Lancer l'intégration après un court délai
            QTimer.singleShot(1000, self.force_embed_excel_now)
            
            self._set_status(f"Chargé: {file_path.name}", "success")
            
            return True
            
        except Exception as e:
            QMessageBox.critical(self, "Erreur GULFCAM", f"Impossible de charger le fichier:\n{str(e)}")
            return False
            
    def save_changes(self):
        """Enregistre les modifications"""
        if not self.current_file:
            QMessageBox.warning(self, "Erreur GULFCAM", "Aucun fichier chargé.")
            return

        if not self.use_windows_com:
            if self._is_native_embed_active():
                if self._trigger_native_save():
                    self._set_status("Enregistré (tableur intégré)", "success")
                    QMessageBox.information(self, "GULFCAM", "Enregistrement demandé au tableur intégré.")
                else:
                    QMessageBox.information(
                        self,
                        "GULFCAM",
                        "Utilisez Ctrl+S dans la fenêtre tableur intégrée pour enregistrer.",
                    )
                return
            if self.preview_workbook is None:
                QMessageBox.warning(self, "Erreur GULFCAM", "Aucune prévisualisation à enregistrer.")
                return
            try:
                self.preview_workbook.save(self.current_file)
                self.is_modified = False
                self._set_status("Enregistré", "success")
                QMessageBox.information(self, "GULFCAM", "Modifications enregistrées avec succès!")
            except Exception as e:
                QMessageBox.critical(self, "Erreur GULFCAM", f"Impossible d'enregistrer:\n{str(e)}")
            return

        if self.workbook:
            try:
                # Vérifier la connexion avant de sauvegarder
                if self._check_excel_connection():
                    self.workbook.Save()
                    self._set_status("Enregistré", "success")
                    QMessageBox.information(self, "GULFCAM", "Modifications enregistrées avec succès!")
                else:
                    QMessageBox.warning(self, "Erreur GULFCAM", "Connexion Excel perdue. Veuillez recharger le fichier.")
            except Exception as e:
                QMessageBox.critical(self, "Erreur GULFCAM", f"Impossible d'enregistrer:\n{str(e)}")
                
    def refresh(self):
        """Recharge le fichier"""
        if self.current_file and self.current_file.exists():
            reply = QMessageBox.question(
                self, "Confirmation GULFCAM",
                "Recharger le fichier ? Les modifications non enregistrées seront perdues.",
                QMessageBox.Yes | QMessageBox.No
            )
            if reply == QMessageBox.Yes:
                if self.use_windows_com:
                    self.load_file(self.current_file)
                else:
                    if self._start_native_embed(self.current_file):
                        self._set_status("Rechargé", "success")
                    else:
                        self._set_status("Rechargement échoué", "danger")

    def clear_loaded_file(self):
        """Réinitialise le fichier chargé sans fermer le widget."""
        if self.resize_timer and self.resize_timer.isActive():
            self.resize_timer.stop()
        if self.embed_timer and self.embed_timer.isActive():
            self.embed_timer.stop()
            
        if self.use_windows_com:
            if self.workbook:
                try:
                    self.workbook.Close(SaveChanges=False)
                except Exception:
                    pass
                self.workbook = None
            self.excel_window_hwnd = None
            self.excel_container_widget = None
            self.rpc_error_count = 0
        else:
            self._stop_embedded_process()
            if self.preview_workbook is not None:
                try:
                    self.preview_workbook.close()
                except Exception:
                    pass
                self.preview_workbook = None
            self.preview_sheet_name = None
            self.sheet_selector.blockSignals(True)
            self.sheet_selector.clear()
            self.sheet_selector.blockSignals(False)
            self.preview_table.clear()

        self.current_file = None
        self.is_modified = False
        self.file_indicator.setText("Aucun fichier chargé")
        self.content_stack.setCurrentWidget(self.placeholder)
        self._set_loaded_mode(False)
        self._set_status("Prêt", "success")

    def shutdown(self):
        """Libère proprement les ressources COM / workbook."""
        if hasattr(self, 'embed_timer') and self.embed_timer:
            self.embed_timer.stop()
        if hasattr(self, 'resize_timer') and self.resize_timer:
            self.resize_timer.stop()
            
        if self.use_windows_com:
            if self.workbook:
                try:
                    self.workbook.Close(SaveChanges=False)
                except Exception:
                    pass
                self.workbook = None
            if self.excel:
                try:
                    self.excel.Quit()
                except Exception:
                    pass
                self.excel = None
        else:
            self._stop_embedded_process()
            if self.preview_workbook is not None:
                try:
                    self.preview_workbook.close()
                except Exception:
                    pass
                self.preview_workbook = None
                
    def closeEvent(self, event):
        """Fermeture propre"""
        self.shutdown()
        event.accept()

    def set_colors(self, colors):
        """Met à jour les couleurs et réapplique le thème."""
        self.colors = colors
        self._apply_theme_styles()
        self._set_status(self.status_indicator.text(), self.status_color_key)

class CompanyDashboard(QWidget):
    """Dashboard entreprise avec branding GULFCAM"""
    
    def __init__(self, colors, action_callbacks: Optional[Dict[str, Callable[[], None]]] = None):
        super().__init__()
        self.colors = colors
        self.action_callbacks = action_callbacks or {}
        self.setup_ui()
        
    def setup_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(24, 24, 24, 24)
        layout.setSpacing(24)
        
        # En-tête avec logo
        header = QFrame()
        header_layout = QGridLayout(header)
        header_layout.setContentsMargins(0, 0, 0, 0)
        header_layout.setHorizontalSpacing(16)
        header_layout.setColumnStretch(0, 1)
        header_layout.setColumnStretch(1, 1)
        header_layout.setColumnStretch(2, 1)
        
        # Logo GULFCAM avec tagline
        logo_container = QFrame()
        logo_layout = QHBoxLayout(logo_container)
        logo_layout.setSpacing(16)

        logo_label = QLabel()
        logo_pixmap = load_logo_pixmap(160, 110)
        if not logo_pixmap.isNull():
            logo_label.setPixmap(logo_pixmap)
            logo_label.setFixedSize(164, 114)
            logo_label.setAlignment(Qt.AlignCenter)
        else:
            logo_label.setText("G")
            logo_label.setStyleSheet(f"""
                font-size: 32px;
                font-weight: 800;
                color: white;
                background-color: {self.colors['primary']};
                border-radius: 10px;
                padding: 8px 16px;
            """)

        brand_text = QVBoxLayout()
        brand_text.setSpacing(2)
        brand_name = QLabel("GULFCAM DSF")
        brand_name.setStyleSheet(f"""
            font-size: 24px;
            font-weight: 800;
            color: {self.colors['primary']};
            letter-spacing: 1px;
        """)
        tagline = QLabel(GULFCAM_TAGLINE)
        tagline.setStyleSheet(f"""
            color: {self.colors['secondary']};
            font-size: 14px;
            font-style: italic;
        """)
        brand_text.addWidget(brand_name)
        brand_text.addWidget(tagline)

        logo_layout.addWidget(logo_label)
        logo_layout.addLayout(brand_text)
        
        # Date et session
        date_label = QLabel(time.strftime("%d %B %Y"))
        date_label.setStyleSheet(f"""
            color: {self.colors['gray']};
            font-size: 15px;
            font-weight: 600;
            padding: 8px 16px;
            background-color: {self.colors['light-gray']};
            border-radius: 12px;
        """)

        header_layout.addWidget(logo_container, 0, 0, Qt.AlignLeft | Qt.AlignVCenter)
        header_layout.addWidget(date_label, 0, 1, Qt.AlignRight | Qt.AlignVCenter)
        
        layout.addWidget(header)
        
        # Message de bienvenue avec design enrichi
        welcome_card = ModernCard(colors=self.colors)
        welcome_card.setStyleSheet(welcome_card.styleSheet() + f"""
            ModernCard {{
                background: qlineargradient(x1:0, y1:0, x2:1, y2:0, 
                    stop:0 {self.colors['primary']}15, 
                    stop:1 {self.colors['card-bg']});
                border-left: 5px solid {self.colors['primary']};
            }}
        """)
        welcome_layout = QHBoxLayout(welcome_card)
        welcome_layout.setContentsMargins(20, 20, 20, 20)
        
        welcome_icon = QLabel()
        welcome_icon.setPixmap(
            ui_icon(
                self,
                "fa5s.gem",
                color=self.colors["primary"],
                fallback=QStyle.StandardPixmap.SP_ComputerIcon,
            ).pixmap(32, 32)
        )
        welcome_icon.setFixedSize(40, 40)
        welcome_icon.setAlignment(Qt.AlignCenter)
        
        welcome_text_layout = QVBoxLayout()
        welcome_text_layout.setSpacing(4)
        
        welcome_title = QLabel(f"Bienvenue sur GULFCAM DSF Mapper")
        welcome_title.setStyleSheet(f"""
            font-size: 22px;
            font-weight: 800;
            color: {self.colors['primary']};
        """)
        
        welcome_subtitle = QLabel("L'excellence digitale au service de votre conformité fiscale.")
        welcome_subtitle.setStyleSheet(f"color: {self.colors['gray']}; font-size: 14px; font-weight: 500;")
        
        welcome_text_layout.addWidget(welcome_title)
        welcome_text_layout.addWidget(welcome_subtitle)
        
        welcome_layout.addWidget(welcome_icon)
        welcome_layout.addSpacing(15)
        welcome_layout.addLayout(welcome_text_layout)
        welcome_layout.addStretch()
        
        layout.addWidget(welcome_card)
        
        # Actions rapides
        actions_card = ModernCard(colors=self.colors)
        actions_layout = QVBoxLayout(actions_card)
        
        actions_title = QLabel("Actions rapides GULFCAM")
        actions_title.setStyleSheet(f"""
            font-size: 18px;
            font-weight: 600;
            color: {self.colors['primary']};
            margin-bottom: 16px;
        """)
        
        actions_grid = QGridLayout()
        actions_grid.setSpacing(12)
        
        actions = [
            ("Nouveau DSF", "new_dsf", "fa5s.file-alt", QStyle.StandardPixmap.SP_FileIcon, self.colors['primary']),
            ("Charger balance", "load_balance", "fa5s.file-import", QStyle.StandardPixmap.SP_DirOpenIcon, self.colors['success']),
            ("Rapports", "reports", "fa5s.chart-line", QStyle.StandardPixmap.SP_FileDialogDetailedView, self.colors['secondary']),
            ("Paramètres", "settings", "fa5s.cogs", QStyle.StandardPixmap.SP_FileDialogContentsView, self.colors['gray']),
        ]
        
        for i, (text, action_key, fa_name, icon_sp, color) in enumerate(actions):
            btn = QPushButton(text)
            btn.setIcon(ui_icon(self, fa_name, color=color, fallback=icon_sp))
            btn.setIconSize(QSize(18, 18))
            btn.setCursor(Qt.PointingHandCursor)
            btn.setStyleSheet(f"""
                QPushButton {{
                    background-color: {self.colors['card-bg']};
                    color: {self.colors['dark']};
                    border: 1px solid {self.colors['light-gray']};
                    border-radius: 12px;
                    padding: 24px;
                    font-weight: 700;
                    text-align: left;
                    font-size: 15px;
                }}
                QPushButton:hover {{
                    background-color: {color}10;
                    border: 2px solid {color};
                    padding: 23px;
                }}
                QPushButton:pressed {{
                    background-color: {color}20;
                }}
            """)
            callback = self.action_callbacks.get(action_key)
            if callback:
                btn.clicked.connect(callback)
                
            actions_grid.addWidget(btn, i // 2, i % 2)
            
        actions_layout.addWidget(actions_title)
        actions_layout.addLayout(actions_grid)
        
        layout.addWidget(actions_card)

    def set_colors(self, colors):
        """Conserve la compatibilité lors d'un refresh de thème."""
        self.colors = colors

class MainWorkView(QWidget):
    """Vue de travail principale avec branding GULFCAM - SECTION 4 AGRANDIE"""
    
    def __init__(self, controller, colors):
        super().__init__()
        self.controller = controller
        self.colors = colors
        self.balance_file: Optional[str] = None
        self.balance_file_n1: Optional[str] = None
        self.generated_dsf_file: Optional[Path] = None
        self.latest_artifacts: Optional[PipelineArtifacts] = None
        self.latest_report_path: Optional[Path] = None
        self.section_widgets: Dict[str, QWidget] = {}
        self.scroll_area: Optional[QScrollArea] = None
        self.pipeline_thread: Optional[QThread] = None
        self.pipeline_worker: Optional[PipelineWorker] = None
        self.pipeline_running: bool = False
        self.setup_ui()
        
    def setup_ui(self):
        self.setObjectName("mainWorkView")
        self.setStyleSheet(f"""
            QWidget#mainWorkView {{
                background-color: {self.colors['light']};
            }}
        """)

        main_layout = QVBoxLayout(self)
        main_layout.setContentsMargins(0, 0, 0, 0)
        main_layout.setSpacing(0)
        
        # Barre de navigation
        nav_bar = QFrame()
        nav_bar.setObjectName("mainNavBar")
        nav_bar.setStyleSheet(f"""
            QFrame#mainNavBar {{
                background-color: {self.colors['card-bg']};
                border-bottom: 2px solid {self.colors['primary']}20;
                padding: 8px 24px;
            }}
        """)
        nav_layout = QHBoxLayout(nav_bar)
        nav_layout.setContentsMargins(0, 0, 0, 0)
        
        # Logo GULFCAM
        logo_container = QFrame()
        logo_container.setStyleSheet("QFrame { background: transparent; border: none; }")
        logo_layout = QHBoxLayout(logo_container)
        logo_layout.setSpacing(8)

        logo_image = QLabel()
        logo_pixmap = load_logo_pixmap(104, 74)
        if not logo_pixmap.isNull():
            logo_image.setPixmap(logo_pixmap)
            logo_image.setAlignment(Qt.AlignCenter)
            logo_image.setFixedSize(logo_pixmap.size())
        else:
            logo_image.setText("G")
            logo_image.setStyleSheet(f"""
                font-size: 24px;
                font-weight: 800;
                color: white;
                background-color: {self.colors['primary']};
                padding: 8px 16px;
                border-radius: 8px;
            """)
        
        logo_full = QLabel("GULFCAM DSF")
        logo_full.setStyleSheet(f"""
            font-size: 18px;
            font-weight: 700;
            color: {self.colors['primary']};
            letter-spacing: 1px;
        """)
        
        logo_layout.addWidget(logo_image)
        logo_layout.addWidget(logo_full)
        
        # Navigation
        nav_buttons = [
            ("Accueil", 0, "fa5s.home", QStyle.StandardPixmap.SP_DirHomeIcon),
            ("Balance", 1, "fa5s.balance-scale", QStyle.StandardPixmap.SP_DriveHDIcon),
            ("DSF", 2, "fa5s.file-signature", QStyle.StandardPixmap.SP_FileIcon),
            ("Rapports", 3, "fa5s.chart-bar", QStyle.StandardPixmap.SP_FileDialogDetailedView),
        ]
        
        nav_layout.addWidget(logo_container)
        nav_layout.addSpacing(40)
        
        for text, idx, fa_name, icon_sp in nav_buttons:
            btn = QPushButton(text)
            btn.setIcon(ui_icon(self, fa_name, color=self.colors["gray"], fallback=icon_sp))
            btn.setIconSize(QSize(16, 16))
            btn.setCursor(Qt.PointingHandCursor)
            btn.setStyleSheet(f"""
                QPushButton {{
                    background-color: transparent;
                    color: {self.colors['gray']};
                    border: none;
                    padding: 8px 16px;
                    font-weight: 500;
                    border-radius: 20px;
                }}
                QPushButton:hover {{
                    background-color: {self.colors['primary']}10;
                    color: {self.colors['primary']};
                }}
                QPushButton:pressed {{
                    background-color: {self.colors['secondary']}20;
                    color: {self.colors['secondary']};
                }}
            """)
            btn.clicked.connect(lambda checked, x=idx: self.controller.switch_view(x))
            nav_layout.addWidget(btn)
            
        nav_layout.addStretch()
        
        # Thème toggle
        theme_btn = QPushButton("Thème")
        theme_btn.setIcon(
            ui_icon(
                self,
                "fa5s.adjust",
                color=self.colors["primary"],
                fallback=QStyle.StandardPixmap.SP_DesktopIcon,
            )
        )
        theme_btn.setIconSize(QSize(16, 16))
        theme_btn.setCursor(Qt.PointingHandCursor)
        theme_btn.setMinimumHeight(40)
        theme_btn.setStyleSheet(f"""
            QPushButton {{
                background-color: {self.colors['primary']}10;
                color: {self.colors['primary']};
                border: none;
                border-radius: 20px;
                font-size: 13px;
                font-weight: 600;
                padding: 0 12px;
            }}
            QPushButton:hover {{
                background-color: {self.colors['primary']};
                color: white;
            }}
        """)
        theme_btn.clicked.connect(self.controller.toggle_dark_mode)
        nav_layout.addWidget(theme_btn)
        
        main_layout.addWidget(nav_bar)
        
        # Zone de contenu avec scroll
        self.scroll_area = QScrollArea()
        self.scroll_area.setWidgetResizable(True)
        self.scroll_area.setStyleSheet(f"""
            QScrollArea {{
                border: none;
                background-color: {self.colors['light']};
            }}
        """)
        
        content = QWidget()
        content.setObjectName("scrollContent")
        content.setStyleSheet(f"""
            QWidget#scrollContent {{
                background-color: {self.colors['light']};
            }}
        """)
        content_layout = QVBoxLayout(content)
        content_layout.setContentsMargins(24, 24, 24, 24)
        content_layout.setSpacing(24)
        
        # Dashboard
        self.dashboard = CompanyDashboard(
            self.colors,
            action_callbacks={
                "new_dsf": self.start_new_dsf,
                "load_balance": self.pick_and_load_balance,
                "reports": self.open_latest_report,
                "settings": self.show_settings,
            },
        )
        self.section_widgets["dashboard"] = self.dashboard
        content_layout.addWidget(self.dashboard)
        
        # Section Balance
        balance_card = ModernCard(colors=self.colors)
        balance_layout = QVBoxLayout(balance_card)
        
        balance_title = QLabel("1. Chargement de la balance")
        balance_title.setStyleSheet(f"""
            font-size: 20px;
            font-weight: 600;
            color: {self.colors['primary']};
            margin-bottom: 16px;
        """)
        
        # Zone de sélection de fichier N
        file_row = QHBoxLayout()

        self.balance_path = QLineEdit()
        self.balance_path.setPlaceholderText("Balance exercice N (.xlsx)")
        self.balance_path.setStyleSheet(f"""
            QLineEdit {{
                border: 2px solid {self.colors['light-gray']};
                border-radius: 8px;
                padding: 12px;
                font-size: 14px;
                background-color: {self.colors['card-bg']};
                color: {self.colors['dark']};
            }}
            QLineEdit:focus {{
                border-color: {self.colors['secondary']};
            }}
        """)

        self.balance_path_n1 = QLineEdit()
        self.balance_path_n1.setPlaceholderText("Balance exercice N-1 (.xlsx) - recommandé")
        self.balance_path_n1.setStyleSheet(self.balance_path.styleSheet())

        self.browse_btn = AnimatedButton("Parcourir N", primary=False, colors=self.colors)
        self.browse_btn_n1 = AnimatedButton("Parcourir N-1", primary=False, colors=self.colors)
        self.load_btn = AnimatedButton("Charger", primary=True, colors=self.colors)
        self.load_btn.clicked.connect(self.load_balance)

        self.browse_btn.clicked.connect(self.pick_balance)
        self.browse_btn_n1.clicked.connect(self.pick_balance_n1)

        file_row_n1 = QHBoxLayout()
        file_row.addWidget(self.balance_path, 1)
        file_row.addWidget(self.browse_btn)
        file_row_n1.addWidget(self.balance_path_n1, 1)
        file_row_n1.addWidget(self.browse_btn_n1)
        file_row.addWidget(self.load_btn)

        year_hint = QLabel("N est requis. N-1 améliore fortement les colonnes comparatives DSF.")
        year_hint.setStyleSheet(f"color: {self.colors['gray']}; font-size: 12px;")
        
        balance_layout.addWidget(balance_title)
        balance_layout.addLayout(file_row)
        balance_layout.addLayout(file_row_n1)
        balance_layout.addWidget(year_hint)
        
        self.section_widgets["balance"] = balance_card
        content_layout.addWidget(balance_card)
        
        # Section Excel intégré
        excel_card = ModernCard(colors=self.colors)
        excel_layout = QVBoxLayout(excel_card)
        
        excel_title = QLabel("2. Visualisation et édition Excel")
        excel_title.setStyleSheet(f"""
            font-size: 20px;
            font-weight: 600;
            color: {self.colors['primary']};
            margin-bottom: 16px;
        """)
        
        excel_layout.addWidget(excel_title)
        
        # Widget Excel intégré
        self.excel_widget = ExcelIntegratedWidget(colors=self.colors)
        excel_layout.addWidget(self.excel_widget, 1)
        
        self.section_widgets["excel"] = excel_card
        content_layout.addWidget(excel_card)
        
        # Section DSF
        dsf_card = ModernCard(colors=self.colors)
        dsf_layout = QVBoxLayout(dsf_card)
        
        dsf_title = QLabel("3. Génération DSF")
        dsf_title.setStyleSheet(f"""
            font-size: 20px;
            font-weight: 600;
            color: {self.colors['primary']};
            margin-bottom: 16px;
        """)
        
        dsf_actions = QHBoxLayout()
        
        self.generate_btn = AnimatedButton("Générer le DSF", primary=True, colors=self.colors)
        self.generate_btn.clicked.connect(self.generate_dsf)
        
        self.export_btn = AnimatedButton("Exporter", primary=False, colors=self.colors)
        self.export_btn.clicked.connect(self.export_dsf)

        self.report_btn = AnimatedButton("Voir rapport", primary=False, colors=self.colors)
        self.report_btn.clicked.connect(self.open_latest_report)
        self.report_btn.setEnabled(False)
        
        dsf_actions.addWidget(self.generate_btn)
        dsf_actions.addWidget(self.export_btn)
        dsf_actions.addWidget(self.report_btn)
        dsf_actions.addStretch()

        self.pipeline_step_label = QLabel("Prêt")
        self.pipeline_step_label.setStyleSheet(f"""
            color: {self.colors['gray']};
            font-size: 12px;
            font-weight: 500;
            margin-top: 4px;
        """)

        self.pipeline_progress = QProgressBar()
        self.pipeline_progress.setRange(0, 100)
        self.pipeline_progress.setValue(0)
        self.pipeline_progress.setTextVisible(True)
        self.pipeline_progress.setFormat("%p%")
        self.pipeline_progress.setStyleSheet(f"""
            QProgressBar {{
                border: 1px solid {self.colors['light-gray']};
                border-radius: 8px;
                text-align: center;
                height: 14px;
                background-color: {self.colors['card-bg']};
                color: {self.colors['dark']};
            }}
            QProgressBar::chunk {{
                border-radius: 7px;
                background-color: {self.colors['primary']};
            }}
        """)
        
        dsf_layout.addWidget(dsf_title)
        dsf_layout.addLayout(dsf_actions)
        dsf_layout.addWidget(self.pipeline_step_label)
        dsf_layout.addWidget(self.pipeline_progress)
        
        self.section_widgets["dsf"] = dsf_card
        content_layout.addWidget(dsf_card)

        # Section Rapports et historique DSF - AGRANDIE
        reports_card = ModernCard(colors=self.colors)
        reports_card.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Preferred)
        reports_layout = QVBoxLayout(reports_card)
        reports_layout.setContentsMargins(20, 20, 20, 20)
        reports_layout.setSpacing(12)

        reports_title = QLabel("4. Rapports et historique")
        reports_title.setStyleSheet(f"""
            font-size: 20px;
            font-weight: 600;
            color: {self.colors['primary']};
            margin-bottom: 12px;
        """)
        
        # Zone des rapports récents avec plus d'espace
        reports_frame = QFrame()
        reports_frame.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Minimum)
        reports_frame.setStyleSheet(f"""
            QFrame {{
                background-color: {self.colors['light']};
                border: 1px solid {self.colors['light-gray']};
                border-radius: 10px;
            }}
        """)
        reports_frame_layout = QVBoxLayout(reports_frame)
        reports_frame_layout.setContentsMargins(14, 12, 14, 12)
        
        self.report_status_label = QLabel("Aucun rapport généré pour le moment.")
        self.report_status_label.setStyleSheet(f"""
            color: {self.colors['dark']};
            font-size: 14px;
            font-weight: 500;
            padding: 18px;
            background-color: {self.colors['card-bg']};
            border: 1px solid {self.colors['light-gray']};
            border-radius: 8px;
        """)
        reports_frame_layout.addWidget(self.report_status_label)
        
        reports_layout.addWidget(reports_title)
        reports_layout.addWidget(reports_frame)

        # Historique des DSF générés - PLUS GRAND
        history_header = QHBoxLayout()
        history_header.setContentsMargins(0, 6, 0, 0)
        history_header.setSpacing(12)
        history_label = QLabel("Historique des DSF produits (30 derniers)")
        history_label.setStyleSheet(f"""
            font-size: 20px;
            font-weight: 700;
            color: {self.colors['primary']};
        """)
        history_refresh_btn = QPushButton("Actualiser")
        history_refresh_btn.setCursor(Qt.PointingHandCursor)
        history_refresh_btn.setMinimumHeight(36)
        history_refresh_btn.setMinimumWidth(118)
        history_refresh_btn.setSizePolicy(QSizePolicy.Fixed, QSizePolicy.Fixed)
        history_refresh_btn.setStyleSheet(f"""
            QPushButton {{ 
                font-size: 13px; 
                padding: 6px 16px; 
                color: {self.colors['card-bg']}; 
                background-color: {self.colors['primary']}; 
                border: 1px solid {self.colors['primary']}; 
                border-radius: 8px; 
                font-weight: 700;
            }}
            QPushButton:hover {{ 
                background-color: {self.colors['primary_light']}; 
                border-color: {self.colors['primary_light']};
            }}
            QPushButton:pressed {{
                background-color: {self.colors['primary']}CC;
            }}
        """)
        history_refresh_btn.clicked.connect(self._refresh_dsf_history)
        history_header.addWidget(history_label, 0, Qt.AlignLeft | Qt.AlignVCenter)
        history_header.addStretch(1)
        history_header.addWidget(history_refresh_btn, 0, Qt.AlignRight | Qt.AlignVCenter)
        reports_layout.addLayout(history_header)
        
        # Table d'historique agrandie
        self.dsf_history_list = QTableWidget(0, 3)
        self.dsf_history_list.setHorizontalHeaderLabels(["Fichier", "Date", ""])
        self.dsf_history_list.horizontalHeader().setStretchLastSection(False)
        self.dsf_history_list.horizontalHeader().setSectionResizeMode(0, QHeaderView.Stretch)
        self.dsf_history_list.horizontalHeader().setSectionResizeMode(1, QHeaderView.Fixed)
        self.dsf_history_list.horizontalHeader().setSectionResizeMode(2, QHeaderView.ResizeToContents)
        self.dsf_history_list.setColumnWidth(1, 190)
        self.dsf_history_list.setColumnWidth(2, 160)
        self.dsf_history_list.setAlternatingRowColors(True)
        self.dsf_history_list.setShowGrid(False)
        self.dsf_history_list.verticalHeader().setVisible(False)
        self.dsf_history_list.verticalHeader().setDefaultSectionSize(44)
        self.dsf_history_list.horizontalHeader().setMinimumHeight(40)
        self.dsf_history_list.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        self.dsf_history_list.setHorizontalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        self.dsf_history_list.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
        self.dsf_history_list.setStyleSheet(f"""
            QTableWidget {{
                background-color: {self.colors['card-bg']};
                color: {self.colors['dark']};
                border: 1px solid {self.colors['light-gray']};
                border-radius: 8px;
                font-size: 13px;
                selection-background-color: {self.colors['primary']}22;
                selection-color: {self.colors['dark']};
            }}
            QTableWidget::item {{
                padding: 8px 10px;
                border-bottom: 1px solid {self.colors['light-gray']};
            }}
            QTableWidget::item:alternate {{
                background-color: {self.colors['light']}B0;
            }}
            QHeaderView::section {{
                background-color: {self.colors['primary']}14;
                color: {self.colors['primary']};
                border: none;
                border-bottom: 2px solid {self.colors['primary']};
                padding: 10px;
                font-weight: 700;
                font-size: 14px;
            }}
            QTableWidget QScrollBar:vertical {{
                background: {self.colors['light']};
                width: 11px;
                border-radius: 5px;
            }}
            QTableWidget QScrollBar::handle:vertical {{
                background: {self.colors['primary']}80;
                min-height: 28px;
                border-radius: 5px;
            }}
            QTableWidget QScrollBar::add-line:vertical,
            QTableWidget QScrollBar::sub-line:vertical {{
                height: 0px;
            }}
        """)
        reports_layout.addWidget(self.dsf_history_list, 1)
        self._refresh_dsf_history()

        self.section_widgets["reports"] = reports_card
        content_layout.addWidget(reports_card, 1)
        
        # Footer
        footer = QLabel(f"{GULFCAM_LOGO_TEXT} • {GULFCAM_TAGLINE} • DSF Mapper v1.0")
        footer.setAlignment(Qt.AlignCenter)
        footer.setStyleSheet(f"""
            color: {self.colors['gray']};
            font-size: 12px;
            padding: 16px;
            border-top: 1px solid {self.colors['light-gray']};
        """)
        content_layout.addWidget(footer)
        
        self.scroll_area.setWidget(content)
        main_layout.addWidget(self.scroll_area, 1)
        
    def _refresh_dsf_history(self) -> None:
        """Rafraîchit la liste des DSF produits (output + historique sauvegardé)."""
        if not hasattr(self, "dsf_history_list"):
            return
        self.dsf_history_list.setRowCount(0)
        seen: set = set()
        entries: List[tuple] = []  # (path, mtime)

        # 1. Fichiers dans output/ (DSF_OUTPUT_UI_*.xlsx)
        if DEFAULT_OUTPUT_DIR.exists():
            for p in DEFAULT_OUTPUT_DIR.glob("DSF_OUTPUT_UI_*.xlsx"):
                if p.is_file() and str(p) not in seen:
                    seen.add(str(p))
                    try:
                        mtime = p.stat().st_mtime
                        entries.append((p, mtime))
                    except OSError:
                        pass

        # 2. Historique sauvegardé (fichiers encore existants)
        for path_str in getattr(self.controller.settings, "recent_dsf_files", []) or []:
            p = Path(path_str)
            if p.exists() and str(p) not in seen:
                seen.add(str(p))
                try:
                    entries.append((p, p.stat().st_mtime))
                except OSError:
                    pass

        entries.sort(key=lambda x: x[1], reverse=True)
        max_action_col_width = 0
        for path, mtime in entries[:30]:
            row = self.dsf_history_list.rowCount()
            self.dsf_history_list.insertRow(row)
            file_item = QTableWidgetItem(path.name)
            file_item.setTextAlignment(Qt.AlignLeft | Qt.AlignVCenter)
            self.dsf_history_list.setItem(row, 0, file_item)

            date_item = QTableWidgetItem(time.strftime("%d/%m/%Y %H:%M", time.localtime(mtime)))
            date_item.setTextAlignment(Qt.AlignCenter)
            self.dsf_history_list.setItem(row, 1, date_item)
            
            # Bouton d'ouverture stylisé
            btn_widget = QWidget()
            btn_layout = QHBoxLayout(btn_widget)
            btn_layout.setContentsMargins(8, 4, 8, 4)
            btn_layout.setSpacing(0)
            btn_layout.setAlignment(Qt.AlignCenter)
            
            btn = QPushButton("Ouvrir DSF")
            btn.setCursor(Qt.PointingHandCursor)
            btn.setMinimumHeight(22)
            btn.setMinimumWidth(65)
            btn.setSizePolicy(QSizePolicy.Minimum, QSizePolicy.Fixed)
            btn.setStyleSheet(f"""
                QPushButton {{ 
                    background-color: {self.colors['primary']}; 
                    color: white;
                    border: 1px solid {self.colors['primary']}; 
                    border-radius: 8px; 
                    padding: 6px 14px; 
                    font-weight: 700;
                    font-size: 8px;
                }}
                QPushButton:hover {{ 
                    background-color: {self.colors['primary_light']}; 
                    border-color: {self.colors['primary_light']};
                }}
                QPushButton:pressed {{
                    background-color: {self.colors['primary']}CC;
                }}
            """)
            btn.clicked.connect(lambda checked, fp=path: self._open_dsf_from_history(fp))
            btn_layout.addWidget(btn, 0, Qt.AlignCenter)
            self.dsf_history_list.setCellWidget(row, 2, btn_widget)

            row_height = max(
                self.dsf_history_list.verticalHeader().defaultSectionSize(),
                btn.sizeHint().height() + btn_layout.contentsMargins().top() + btn_layout.contentsMargins().bottom() + 2,
            )
            self.dsf_history_list.setRowHeight(row, row_height)

            action_col_width = (
                btn.sizeHint().width()
                + btn_layout.contentsMargins().left()
                + btn_layout.contentsMargins().right()
                + 8
            )
            max_action_col_width = max(max_action_col_width, action_col_width)

        if max_action_col_width > 0:
            self.dsf_history_list.setColumnWidth(2, max(152, max_action_col_width))
            self.dsf_history_list.resizeColumnToContents(2)
            self.dsf_history_list.setColumnWidth(
                2, max(self.dsf_history_list.columnWidth(2), max(152, max_action_col_width))
            )

        self._resize_history_table_to_content()

    def _resize_history_table_to_content(self) -> None:
        """Adapte la hauteur du tableau d'historique au nombre de lignes."""
        if not hasattr(self, "dsf_history_list"):
            return
        row_count = self.dsf_history_list.rowCount()
        header_h = self.dsf_history_list.horizontalHeader().height() or 40
        row_h = self.dsf_history_list.verticalHeader().defaultSectionSize() or 44
        frame = self.dsf_history_list.frameWidth() * 2
        target_rows = max(4, row_count)
        rows_height = sum(self.dsf_history_list.rowHeight(i) for i in range(row_count))
        rows_height = max(rows_height, target_rows * row_h)
        content_margins = self.dsf_history_list.contentsMargins()
        margins_height = content_margins.top() + content_margins.bottom()
        h_scroll_h = self.dsf_history_list.horizontalScrollBar().sizeHint().height()
        total_height = header_h + rows_height + frame + margins_height + h_scroll_h + 6
        self.dsf_history_list.setMinimumHeight(total_height)

    def _open_dsf_from_history(self, path: Path) -> None:
        """Ouvre un DSF depuis l'historique dans le widget Excel."""
        if not path.exists():
            QMessageBox.warning(self, "GULFCAM", f"Fichier introuvable: {path.name}")
            return
        if self.excel_widget.load_file(path):
            self.generated_dsf_file = path
            self.controller.statusBar().showMessage(f"GULFCAM - DSF ouvert: {path.name}")
            self.navigate_to(2)

    def pick_balance(self):
        default_dir = PROJECT_ROOT / "input"
        if not default_dir.exists():
            default_dir = Path.home() / "Desktop"
        path, _ = QFileDialog.getOpenFileName(
            self,
            "Sélectionner le bilan - GULFCAM",
            str(default_dir),
            "Fichiers Excel (*.xlsx)"
        )
        if path:
            self.balance_path.setText(path)

    def pick_balance_n1(self):
        default_dir = PROJECT_ROOT / "input"
        if not default_dir.exists():
            default_dir = Path.home() / "Desktop"
        path, _ = QFileDialog.getOpenFileName(
            self,
            "Sélectionner la balance N-1 - GULFCAM",
            str(default_dir),
            "Fichiers Excel (*.xlsx)"
        )
        if path:
            self.balance_path_n1.setText(path)

    def pick_and_load_balance(self):
        """Action rapide: ouvrir le sélecteur puis charger immédiatement."""
        self.pick_balance()
        if self.balance_path.text().strip():
            self.load_balance()
            
    def load_balance(self):
        if self.is_pipeline_running():
            QMessageBox.information(self, "GULFCAM", "Attendez la fin de la génération en cours.")
            return
        path_n = self.balance_path.text().strip()
        path_n1 = self.balance_path_n1.text().strip()
        if not path_n or not Path(path_n).exists():
            QMessageBox.warning(self, "Erreur GULFCAM", "Veuillez sélectionner un fichier valide")
            return
        if path_n1 and not Path(path_n1).exists():
            QMessageBox.warning(self, "Erreur GULFCAM", "Le fichier N-1 est introuvable.")
            return
            
        if self.excel_widget.load_file(Path(path_n)):
            self.balance_file = path_n
            self.balance_file_n1 = path_n1 or None
            self.generated_dsf_file = None
            self.latest_artifacts = None
            self.latest_report_path = None
            self.report_btn.setEnabled(False)
            self.report_status_label.setText("Balance chargée. Lancez la génération pour créer un rapport.")
            self.pipeline_progress.setValue(0)
            self.pipeline_step_label.setText("Prêt")
            if self.balance_file_n1:
                self.controller.statusBar().showMessage(
                    f"GULFCAM - Balances chargées N: {Path(path_n).name} | N-1: {Path(self.balance_file_n1).name}"
                )
            else:
                self.controller.statusBar().showMessage(
                    f"GULFCAM - Balance N chargée: {Path(path_n).name} (N-1 non fourni)"
                )

    def start_new_dsf(self):
        """Réinitialise la session courante."""
        if self.is_pipeline_running():
            QMessageBox.information(self, "GULFCAM", "Attendez la fin de la génération en cours.")
            return
        self.balance_file = None
        self.balance_file_n1 = None
        self.generated_dsf_file = None
        self.latest_artifacts = None
        self.latest_report_path = None
        self.balance_path.clear()
        self.balance_path_n1.clear()
        self.excel_widget.clear_loaded_file()
        self.report_btn.setEnabled(False)
        self.report_status_label.setText("Aucun rapport généré pour le moment.")
        self.pipeline_progress.setValue(0)
        self.pipeline_step_label.setText("Prêt")
        self.controller.statusBar().showMessage("GULFCAM - Nouvelle session initialisée")
        self.navigate_to(0)

    def show_settings(self):
        """Affiche un résumé des paramètres actifs."""
        theme_name = "Sombre" if self.controller.settings.dark_mode else "Clair"
        QMessageBox.information(
            self,
            "Paramètres GULFCAM",
            f"Thème: {theme_name}\n"
            f"Template: {DEFAULT_TEMPLATE_PATH}\n"
            f"Inventaire: {DEFAULT_INVENTORY_PATH}\n"
            f"Règles: {DEFAULT_RULES_PATH}",
        )

    def navigate_to(self, index: int):
        """Navigation vers les sections via la barre haute."""
        mapping = {
            0: "dashboard",
            1: "balance",
            2: "excel",
            3: "reports",
        }
        key = mapping.get(index, "dashboard")
        target = self.section_widgets.get(key)
        if target and self.scroll_area is not None:
            self.scroll_area.ensureWidgetVisible(target, 0, 24)
        if key == "reports" and self.latest_report_path:
            self.controller.statusBar().showMessage(f"Dernier rapport: {self.latest_report_path.name}")

    def is_pipeline_running(self) -> bool:
        return bool(self.pipeline_running and self.pipeline_thread and self.pipeline_thread.isRunning())

    def _set_pipeline_controls_enabled(self, enabled: bool):
        self.pipeline_running = not enabled
        if hasattr(self, "generate_btn"):
            self.generate_btn.setEnabled(enabled)
        if hasattr(self, "export_btn"):
            self.export_btn.setEnabled(enabled)
        if hasattr(self, "report_btn"):
            self.report_btn.setEnabled(enabled and self.latest_report_path is not None)
        if hasattr(self, "browse_btn"):
            self.browse_btn.setEnabled(enabled)
        if hasattr(self, "browse_btn_n1"):
            self.browse_btn_n1.setEnabled(enabled)
        if hasattr(self, "load_btn"):
            self.load_btn.setEnabled(enabled)
        self.balance_path.setReadOnly(not enabled)
        if hasattr(self, "balance_path_n1"):
            self.balance_path_n1.setReadOnly(not enabled)

    @Slot(str)
    def _on_pipeline_status(self, text: str):
        self.controller.statusBar().showMessage(text)

    @Slot(int, str)
    def _on_pipeline_progress(self, percent: int, message: str):
        self.pipeline_progress.setRange(0, 100)
        self.pipeline_progress.setValue(max(0, min(100, int(percent))))
        if message:
            self.pipeline_step_label.setText(message)
            self.controller.statusBar().showMessage(f"GULFCAM - {message}")

    def _build_pipeline_config(self) -> Optional["DSFPipelineConfig"]:
        required_paths = [
            ("Template DSF", DEFAULT_TEMPLATE_PATH),
            ("Inventaire", DEFAULT_INVENTORY_PATH),
            ("Règles", DEFAULT_RULES_PATH),
            ("Mapping préremplissage", DEFAULT_PREFILL_MAPPING_PATH),
        ]
        missing = [name for name, path in required_paths if not path.exists()]
        if missing:
            QMessageBox.critical(
                self,
                "Erreur GULFCAM",
                "Fichiers requis manquants:\n- " + "\n- ".join(missing),
            )
            return None

        timestamp = time.strftime("%Y%m%d_%H%M%S")
        DEFAULT_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        DEFAULT_REPORTS_DIR.mkdir(parents=True, exist_ok=True)
        prefilled_path = DEFAULT_OUTPUT_DIR / f"dsf_prefilled_ui_{timestamp}.xlsx"
        output_path = DEFAULT_OUTPUT_DIR / f"DSF_OUTPUT_UI_{timestamp}.xlsx"
        report_json = DEFAULT_REPORTS_DIR / f"dsf_report_ui_{timestamp}.json"
        report_html = DEFAULT_REPORTS_DIR / f"dsf_report_ui_{timestamp}.html"

        return DSFPipelineConfig(
            template_dsf=DEFAULT_TEMPLATE_PATH,
            balance_input=Path(self.balance_file),
            previous_balance_input=Path(self.balance_file_n1) if self.balance_file_n1 else None,
            inventory_path=DEFAULT_INVENTORY_PATH,
            rules_path=DEFAULT_RULES_PATH,
            prefill_mapping_path=DEFAULT_PREFILL_MAPPING_PATH,
            prefilled_template_path=prefilled_path,
            dsf_output=output_path,
            report_json_path=report_json,
            report_html_path=report_html,
            filling_method="hybrid",
            fuzzy_threshold=0.6,
            apply_calculations=True,
            use_calculation_formulas=True,
            use_smart_general_filler=False,
            apply_note_rules_in_semantic=True,
            enable_notes_filler=True,
            enable_cash_flow=True,
            enable_fiscal_controls=False,
        )

    def _start_pipeline_async(self, config: "DSFPipelineConfig"):
        self._set_pipeline_controls_enabled(False)
        self.pipeline_progress.setRange(0, 100)
        self.pipeline_progress.setValue(0)
        self.pipeline_step_label.setText("Préparation de la génération...")
        self.controller.statusBar().showMessage("GULFCAM - Génération DSF en cours...")
        self.report_status_label.setText("Génération en cours... veuillez patienter.")

        self.pipeline_thread = QThread(self)
        self.pipeline_worker = PipelineWorker(config)
        self.pipeline_worker.moveToThread(self.pipeline_thread)

        self.pipeline_thread.started.connect(self.pipeline_worker.run)
        self.pipeline_worker.status.connect(self._on_pipeline_status)
        self.pipeline_worker.progress.connect(self._on_pipeline_progress)
        self.pipeline_worker.finished.connect(self._on_pipeline_finished)
        self.pipeline_worker.failed.connect(self._on_pipeline_failed)
        self.pipeline_worker.finished.connect(self.pipeline_thread.quit)
        self.pipeline_worker.failed.connect(self.pipeline_thread.quit)
        self.pipeline_thread.finished.connect(self._on_pipeline_thread_finished)
        self.pipeline_thread.start()

    @Slot(object)
    def _on_pipeline_finished(self, artifacts):
        self.pipeline_progress.setRange(0, 100)
        self.pipeline_progress.setValue(100)
        self.pipeline_step_label.setText("Génération terminée")
        self.latest_artifacts = artifacts
        self.generated_dsf_file = artifacts.dsf_output
        self.latest_report_path = artifacts.report_html or artifacts.report_json
        if self.latest_report_path:
            self.report_status_label.setText(f"Rapport disponible: {self.latest_report_path.name}")
        else:
            self.report_status_label.setText("DSF généré, sans rapport attaché.")

        # Historique des DSF
        self.controller.settings.add_recent_dsf(str(artifacts.dsf_output))
        self.controller.settings.save()
        self._refresh_dsf_history()
        self.report_btn.setEnabled(True)

        loaded = self.excel_widget.load_file(artifacts.dsf_output)
        if not loaded:
            QMessageBox.warning(
                self,
                "GULFCAM",
                f"DSF généré mais prévisualisation non chargée:\n{artifacts.dsf_output}",
            )
        self.controller.statusBar().showMessage(f"GULFCAM - DSF généré: {artifacts.dsf_output.name}")
        QMessageBox.information(
            self,
            "Génération terminée",
            "Le DSF a été généré avec succès.\n"
            "Vous pouvez maintenant le prévisualiser et le modifier dans la section 2.",
        )
        self.navigate_to(2)

    @Slot(str)
    def _on_pipeline_failed(self, error: str):
        self.pipeline_progress.setRange(0, 100)
        self.pipeline_progress.setValue(0)
        self.pipeline_step_label.setText("Échec de génération")
        self.report_status_label.setText("Échec de génération. Vérifiez les logs et réessayez.")
        self.controller.statusBar().showMessage("GULFCAM - Échec génération DSF")
        QMessageBox.critical(self, "Erreur GULFCAM", f"Échec génération DSF:\n{error}")

    @Slot()
    def _on_pipeline_thread_finished(self):
        if self.pipeline_worker is not None:
            self.pipeline_worker.deleteLater()
            self.pipeline_worker = None
        if self.pipeline_thread is not None:
            self.pipeline_thread.deleteLater()
            self.pipeline_thread = None
        self._set_pipeline_controls_enabled(True)

    def generate_dsf(self):
        if self.is_pipeline_running():
            QMessageBox.information(self, "GULFCAM", "Une génération est déjà en cours.")
            return

        if not self.balance_file:
            QMessageBox.warning(self, "Erreur GULFCAM", "Veuillez d'abord charger une balance")
            return
        if not self.balance_file_n1:
            reply = QMessageBox.question(
                self,
                "Balance N-1 absente",
                "La balance N-1 n'est pas renseignée.\n"
                "Le DSF sera généré avec une couverture comparative réduite.\n\n"
                "Continuer quand même ?",
                QMessageBox.Yes | QMessageBox.No,
            )
            if reply != QMessageBox.Yes:
                return

        if DSFPipeline is None or DSFPipelineConfig is None:
            QMessageBox.critical(
                self,
                "Erreur GULFCAM",
                "Le module pipeline n'est pas disponible.\nVérifiez l'installation des dépendances du projet.",
            )
            return

        config = self._build_pipeline_config()
        if config is None:
            return
        self._start_pipeline_async(config)
        
    def export_dsf(self):
        if self.is_pipeline_running():
            QMessageBox.information(self, "GULFCAM", "Export indisponible pendant la génération.")
            return
        source_file = self.generated_dsf_file or self.excel_widget.current_file
        if not source_file or not Path(source_file).exists():
            QMessageBox.warning(self, "Erreur GULFCAM", "Aucun fichier à exporter")
            return
            
        path, _ = QFileDialog.getSaveFileName(
            self,
            "Exporter le DSF - GULFCAM",
            str(DEFAULT_OUTPUT_DIR),
            "Fichiers Excel (*.xlsx)"
        )
        
        if path:
            if not path.endswith('.xlsx'):
                path += '.xlsx'
            try:
                self.excel_widget.save_changes()
                shutil.copy(str(source_file), path)
                QMessageBox.information(self, "GULFCAM", f"Fichier exporté avec succès vers:\n{path}")
            except Exception as e:
                QMessageBox.critical(self, "Erreur GULFCAM", f"Export impossible:\n{str(e)}")

    def open_latest_report(self):
        """Ouvre le dernier rapport HTML/JSON généré."""
        report_path = self.latest_report_path
        if report_path is None and self.latest_artifacts is not None:
            report_path = self.latest_artifacts.report_html or self.latest_artifacts.report_json

        if report_path is None or not Path(report_path).exists():
            QMessageBox.information(self, "Rapports GULFCAM", "Aucun rapport disponible pour le moment.")
            return

        local_url = QUrl.fromLocalFile(str(Path(report_path).resolve()))
        if QDesktopServices.openUrl(local_url):
            self.controller.statusBar().showMessage(f"GULFCAM - Rapport ouvert: {Path(report_path).name}")
            return

        try:
            if IS_WINDOWS and hasattr(os, "startfile"):
                os.startfile(str(report_path))
            elif shutil.which("xdg-open"):
                subprocess.Popen(["xdg-open", str(report_path)])
            elif shutil.which("open"):
                subprocess.Popen(["open", str(report_path)])
            else:
                raise RuntimeError("Aucun ouvreur de fichier détecté.")
            self.controller.statusBar().showMessage(f"GULFCAM - Rapport ouvert: {Path(report_path).name}")
        except Exception as e:
            QMessageBox.warning(self, "Rapports GULFCAM", f"Impossible d'ouvrir le rapport:\n{str(e)}")
                
    def set_colors(self, colors):
        """Met à jour les couleurs"""
        self.colors = colors
        self.dashboard.set_colors(colors)
        self.excel_widget.set_colors(colors)

    def shutdown(self) -> bool:
        """Libère les ressources. Retourne False si une génération est en cours."""
        if self.is_pipeline_running():
            return False
        self.excel_widget.shutdown()
        return True

class LightDSFMapperWindow(QMainWindow):
    """Fenêtre principale avec branding GULFCAM"""
    
    def __init__(self):
        super().__init__()
        
        # Charger les paramètres
        self.settings = AppSettings.load()
        self.colors = get_gulfcam_colors(self.settings.dark_mode)
        
        self.setWindowTitle("GULFCAM DSF Mapper")
        self.setWindowIcon(self.create_gulfcam_icon())
        self.resize(self.settings.window_width, self.settings.window_height)
        
        # Widget central
        central = QWidget()
        self.setCentralWidget(central)
        
        # Stack pour les différentes vues
        self.stack = QStackedWidget()
        layout = QVBoxLayout(central)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.addWidget(self.stack)
        
        # Vue de travail principale
        self.main_view = MainWorkView(self, self.colors)
        self.stack.addWidget(self.main_view)
        
        # Barre d'état
        self.setup_status_bar()
        
    def create_gulfcam_icon(self):
        """Crée l'icône de l'application depuis OIP.png."""
        logo_pixmap = load_logo_pixmap(128, 128)
        if not logo_pixmap.isNull():
            return QIcon(logo_pixmap)

        icon = QIcon()
        pixmap = QPixmap(64, 64)
        pixmap.fill(QColor(self.colors['primary']))
        painter = QPainter(pixmap)
        painter.setPen(Qt.white)
        painter.setFont(QFont("Arial", 24, QFont.Bold))
        painter.drawText(pixmap.rect(), Qt.AlignCenter, "G")
        painter.end()
        icon.addPixmap(pixmap)
        return icon
        
    def setup_status_bar(self):
        status_bar = QStatusBar()
        self.setStatusBar(status_bar)
        status_bar.setStyleSheet(f"""
            QStatusBar {{
                background-color: {self.colors['primary']};
                color: white;
                padding: 8px;
                font-size: 12px;
            }}
        """)
        status_bar.showMessage(f"{GULFCAM_LOGO_TEXT} • {GULFCAM_TAGLINE} • Application prête")
        
    def switch_view(self, index):
        """Change la vue active"""
        if hasattr(self.main_view, "navigate_to"):
            self.main_view.navigate_to(index)
            return
        if 0 <= index < self.stack.count():
            self.stack.setCurrentIndex(index)
        
    def toggle_dark_mode(self):
        """Bascule le mode sombre/clair"""
        if hasattr(self.main_view, "is_pipeline_running") and self.main_view.is_pipeline_running():
            QMessageBox.information(
                self,
                "GULFCAM",
                "Le changement de thème est temporairement bloqué pendant la génération DSF.",
            )
            return
        self.settings.dark_mode = not self.settings.dark_mode
        self.colors = get_gulfcam_colors(self.settings.dark_mode)
        self.settings.save()
        
        # Rafraîchir l'interface
        old_view = self.main_view
        self.main_view = MainWorkView(self, self.colors)
        self.stack.addWidget(self.main_view)
        self.stack.setCurrentWidget(self.main_view)
        self.stack.removeWidget(old_view)
        old_view.deleteLater()
        
        # Mettre à jour la barre d'état
        self.statusBar().setStyleSheet(f"""
            QStatusBar {{
                background-color: {self.colors['primary']};
                color: white;
                padding: 8px;
                font-size: 12px;
            }}
        """)
        
        self.statusBar().showMessage(f"Mode {'sombre' if self.settings.dark_mode else 'clair'} GULFCAM activé")
        
    def closeEvent(self, event):
        """Sauvegarde les paramètres à la fermeture"""
        if hasattr(self.main_view, "is_pipeline_running") and self.main_view.is_pipeline_running():
            QMessageBox.warning(
                self,
                "GULFCAM",
                "Une génération DSF est en cours.\nAttendez la fin avant de fermer l'application.",
            )
            event.ignore()
            return

        self.settings.window_width = self.width()
        self.settings.window_height = self.height()
        self.settings.save()
        
        # Fermer Excel proprement
        if hasattr(self.main_view, "shutdown") and not self.main_view.shutdown():
            event.ignore()
            return
            
        event.accept()

def main():
    """Point d'entrée principal"""
    com_initialized = False
    if WINDOWS_COM_AVAILABLE and pythoncom is not None:
        # Initialiser COM uniquement sur Windows avec pywin32 disponible
        pythoncom.CoInitialize()
        com_initialized = True
    
    app = QApplication([])
    
    # Police par défaut
    font = QFont("Segoe UI", 10)
    app.setFont(font)
    
    # Variable pour stocker la fenêtre principale
    main_window = None
    startup_state = {"window": None}

    def load_settings_step():
        AppSettings.load()

    def load_assets_step():
        load_logo_pixmap(185, 195)
        load_lightgroup_pixmap(150, 60)

    def build_ui_step():
        startup_state["window"] = LightDSFMapperWindow()

    startup_tasks = [
        ("Chargement des paramètres", load_settings_step),
        ("Chargement des ressources graphiques", load_assets_step),
        ("Initialisation de l'interface", build_ui_step),
    ]
    
    def show_main_window():
        nonlocal main_window
        main_window = startup_state.get("window")
        if main_window is None:
            main_window = LightDSFMapperWindow()
        main_window.show()
    
    # Afficher le splash screen
    splash = SplashScreen(show_main_window, startup_tasks=startup_tasks, min_duration_ms=5000)
    splash.show()
    
    result = app.exec()
    
    if com_initialized:
        pythoncom.CoUninitialize()
    
    return result

if __name__ == "__main__":
    main()