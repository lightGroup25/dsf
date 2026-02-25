# -*- coding: utf-8 -*-
"""Prefill ENTETE/R1/R2/R3/NOTE 13 with DSF_InfosGenerales mappings."""
from __future__ import annotations

import json
import logging
from dataclasses import asdict, is_dataclass
from datetime import date
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from dsf_general_info import DSF_InfosGenerales, format_date

logger = logging.getLogger(__name__)


class DSFGeneralPrefiller:
    SAFE_HEADER_SHEETS = {
        "ENTETE",
        "ENTÊTE",
        "FICHE R1",
        "R1",
        "FICHE R2",
        "R2",
        "FICHE R3",
        "R3",
        "NOTE 13",
        "NOTE13",
        "PAGE DE GARDE",
    }

    # P3-Prefill : Zones protégées par feuille.
    # Toute cellule dont le numéro de ligne est >= la limite définie ici
    # ne sera PAS écrite par _fill_cells (évite écriture hors zones dans tableaux).
    PROTECTED_ZONES: Dict[str, int] = {
        "PAGE DE GARDE": 50,  # Éviter écriture hors zone identité
        "FICHE R1": 25,
        "R1":       25,
        "FICHE R2": 48,
        "R2":       48,
        "FICHE R3": 12,
        "R3":       12,
    }

    # Limites des tableaux : max nombre de lignes à écrire (évite débordement)
    TABLE_ROW_LIMITS: Dict[Tuple[str, str], int] = {
        ("FICHE R1", "activites"): 0,   # R1 n'a pas activites
        ("FICHE R2", "activites"): 25,
        ("FICHE R3", "dirigeants"): 13,
        ("FICHE R3", "conseil_administration"): 9,
        ("FICHE R3", "actionnaires"): 0,
        ("R3", "actionnaires"): 0,
        ("NOTE 13", "actionnaires"): 20,
    }

    def __init__(self, template_path: Path | str, mapping_path: Path | str = "dsf_prefill_mapping.json"):
        self.template_path = Path(template_path)
        self.mapping_path = Path(mapping_path)
        self.wb = None
        self.mapping: Dict[str, Any] = {}

    def load(self) -> None:
        if not self.template_path.exists():
            raise FileNotFoundError(self.template_path)
        if not self.mapping_path.exists():
            raise FileNotFoundError(self.mapping_path)
        self.wb = load_workbook(self.template_path)
        self.mapping = json.loads(self.mapping_path.read_text(encoding="utf-8"))

    def fill(self, info: DSF_InfosGenerales) -> int:
        if self.wb is None or not self.mapping:
            raise RuntimeError("Prefiller not loaded. Call load() first.")
        
        info_dict = asdict(info)
        filled = 0
        
        # 1. Automatic Universal Header Filling
        filled += self._fill_universal_headers(self.wb, info)

        # 2. Configured Mapping Filling
        for sheet_name, spec in self.mapping.items():
            ws = self._find_sheet(sheet_name)
            if ws is None:
                continue
            sheet_key = self._normalize_sheet_name(ws.title)
            filled += self._fill_cells(ws, sheet_name, spec.get("cells", {}), info, info_dict)
            for table_name, table_spec in (spec.get("tables", {}) or {}).items():
                filled += self._fill_table(ws, sheet_name, sheet_key, table_name, table_spec, info, info_dict)
        # Fallback hardening for PAGE DE GARDE so critical identity fields are always set.
        filled += self._fill_page_de_garde_fallback(info)
        return filled
    
    def _fill_universal_headers(self, wb, info: DSF_InfosGenerales) -> int:
        """Scan all sheets for header patterns and fill them."""
        count = 0
        
        # Valeurs depuis la balance N (exercice_fin)
        designation = info.denomination_sociale
        cloture = format_date(info.exercice_fin) if info.exercice_fin else "31/12/2024"
        cloture_dash = info.exercice_fin.strftime("%d-%m-%Y") if info.exercice_fin else "31-12-2024"
        annee = str(info.exercice_fin.year) if info.exercice_fin else "2024"
        niu = info.num_identification_fiscale
        duree = str(info.duree_mois)

        search_patterns = [
            ("Désignation entité", "Exercice clos le", f"Désignation entité : {designation}                                                                Exercice clos le {cloture_dash}"),
            ("Désignation entité", "exercice clos", f"Désignation entité : {designation}                                                                Exercice clos le {cloture_dash}"),
            ("N° d'identification", "Exercice clos le", f"N° d'identification fiscal : {niu}                                          Exercice clos le : {cloture_dash}                                           Durée en mois : {duree}"),
            ("N° d'identification", "Durée (en mois)", f"Numéro d'identification : {niu}                                                                 Durée (en mois) : {duree}"),
            ("Numéro d’identification", "Durée (en mois)", f"Numéro d'identification : {niu}                                                                 Durée (en mois) : {duree}")
        ]

        # Toutes les feuilles (sauf SOMMAIRE) : en-têtes sur BILAN, CR, Notes, R1, R2, R3, etc.
        EXCLUDED = {"SOMMAIRE"}
        for ws in wb.worksheets:
            if self._normalize_sheet_name(ws.title) in EXCLUDED:
                continue
            for r in range(1, min(21, ws.max_row + 1)):
                for c in range(1, min(9, ws.max_column + 1)):
                    cell = ws.cell(row=r, column=c)
                    val = cell.value
                    if not isinstance(val, str) or not val.strip():
                        continue
                    text = val
                    text_lower = text.lower()
                    # 1. Désignation entité / N° identification
                    for start_marker, end_marker, replacement in search_patterns:
                        if start_marker.lower() in text_lower and end_marker.lower() in text_lower:
                            if self._write_cell(ws, cell.coordinate, replacement):
                                count += 1
                            break
                    else:
                        # 2. BILAN AU 31 DECEMBRE / COMPTE DE RESULTAT (date = exercice N)
                        if "BILAN AU 31 DECEMBRE" in text.upper():
                            self._write_cell(ws, cell.coordinate, f"BILAN AU 31 DECEMBRE {annee}")
                            count += 1
                        elif "COMPTE" in text_lower and "RESULTAT" in text_lower and "31 DECEMBRE" in text.upper():
                            self._write_cell(ws, cell.coordinate, f"COMPTE DE RESULTAT AU 31 DECEMBRE {annee}")
                            count += 1
        return count

    def _fill_page_de_garde_fallback(self, info: DSF_InfosGenerales) -> int:
        if self.wb is None:
            return 0

        ws = None
        for candidate in self.wb.worksheets:
            if self._normalize_sheet_name(candidate.title) == "PAGE DE GARDE":
                ws = candidate
                break
        if ws is None:
            return 0

        fallback_values = {
            "B10": f"CENTRE DE DEPOT DE : {info.centre_depot}",
            "B18": f"EXERCICE CLOS LE : {format_date(info.exercice_fin)}",
            "A27": f"Dénomination sociale : {info.denomination_sociale}",
            "A29": f"Sigle usuel : {info.sigle_usuel}",
            "A31": f"Adresse complète : {info.adresse_complete}",
            "A33": f"N° d'identification fiscale : {info.num_identification_fiscale}",
            "C36": f"Système : {info.systeme_comptable}",
        }

        count = 0
        for cell_ref, value in fallback_values.items():
            master = self._master_cell(ws, cell_ref)
            current = master.value
            should_write = False
            if current is None or (isinstance(current, str) and not current.strip()):
                should_write = True
            elif isinstance(current, str):
                txt = current.strip()
                # If placeholder survived mapping, force replacement.
                if "…" in txt or "..." in txt or "_" in txt:
                    should_write = True
            if should_write and self._write_cell(ws, cell_ref, value):
                count += 1
        return count

    def save(self, output_path: Path | str) -> Path:
        if self.wb is None:
            raise RuntimeError("Workbook not loaded")
        output = Path(output_path)
        output.parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(output)
        return output

    # ------------------------------------------------------------------
    # Helpers
    # ------------------------------------------------------------------
    def _fill_cells(
        self,
        ws: Worksheet,
        sheet_name: str,
        cells_spec: Dict[str, Dict[str, Any]],
        info_obj: DSF_InfosGenerales,
        info_dict: Dict[str, Any],
    ) -> int:
        filled = 0
        for cell_ref, spec in cells_spec.items():
            # P3-Prefill : Skip si la cellule est dans une zone protégée (évite écriture hors zones)
            if self._is_in_protected_zone(ws, cell_ref):
                logger.debug("Zone protégée : skip %s!%s", ws.title, cell_ref)
                continue

            attribute = spec.get("attribute")
            fmt = spec.get("format")
            optional = spec.get("optional", False)
            value = self._resolve_value(info_obj, info_dict, attribute, fmt)

            value_template = spec.get("value_template")
            if value_template and value is not None:
                try:
                    value = value_template.format(value)
                except Exception:
                    pass  # Keep original value if format fails

            if optional and (value is None or value == ""):
                logger.debug("Champ optionnel vide : skip %s", cell_ref)
                continue

            if self._write_cell(ws, cell_ref, value):
                filled += 1
            else:
                logger.debug("Impossible d'écrire %s!%s = '%s'", ws.title, cell_ref, value)
        return filled

    def _fill_table(
        self,
        ws: Worksheet,
        sheet_name: str,
        sheet_key: str,
        table_name: str,
        table_spec: Dict[str, Any],
        info_obj: DSF_InfosGenerales,
        info_dict: Dict[str, Any],
    ) -> int:
        records = getattr(info_obj, table_name, [])
        if not isinstance(records, Iterable):
            return 0

        # Limite de lignes pour éviter débordement hors zone
        limit_key = (sheet_key, table_name)
        alt_key = (sheet_name.strip().upper(), table_name)
        max_rows = self.TABLE_ROW_LIMITS.get(limit_key) or self.TABLE_ROW_LIMITS.get(alt_key)
        if max_rows is not None and max_rows <= 0:
            return 0
        records = list(records)[:max_rows] if max_rows else list(records)
        
        # Check if table uses sections (for split table layouts)
        sections = table_spec.get("sections")
        if sections:
            return self._fill_table_sections(ws, sections, records, info_obj, info_dict, table_name)
        
        # Original logic for single-section tables
        explicit_rows = table_spec.get("explicit_rows")
        data_start = table_spec.get("data_start_row")
        
        if not data_start and not explicit_rows:
            return 0
            
        if data_start:
            start_col = ''.join(ch for ch in data_start if ch.isalpha())
            start_row = int(''.join(ch for ch in data_start if ch.isdigit()))
        else:
            start_row = 0

        count = 0
        for index, record in enumerate(records, start=1):
            record_dict = asdict(record) if is_dataclass(record) else dict(record)
            
            if explicit_rows and index - 1 < len(explicit_rows):
                row_number = explicit_rows[index - 1]
            else:
                row_number = start_row + index - 1

            for column, field_name in table_spec.get("fields", {}).items():
                cell_ref = f"{column}{row_number}"
                value = self._resolve_table_value(field_name, record_dict, index, info_obj)
                self._write_cell(ws, cell_ref, value)
            count += 1
        if count:
            self._write_table_totals(ws, table_spec, start_row, count, info_obj, info_dict)
        return count

    def _fill_table_sections(
        self,
        ws: Worksheet,
        sections: List[Dict[str, Any]],
        records: List,
        info_obj: DSF_InfosGenerales,
        info_dict: Dict[str, Any],
        table_name: str,
    ) -> int:
        """Fill a table that has multiple sections with different layouts."""
        total_filled = 0
        record_index = 0
        
        for section in sections:
            section_rows = section.get("explicit_rows") or []
            row_count = section.get("row_count", len(section_rows))
            fields = section.get("fields", {})
            
            for section_idx in range(row_count):
                if record_index >= len(records):
                    break
                
                record = records[record_index]
                record_dict = asdict(record) if is_dataclass(record) else dict(record)
                
                # Determine which row to write to
                if section_rows and section_idx < len(section_rows):
                    row_number = section_rows[section_idx]
                else:
                    # For data_start_row sections
                    data_start = section.get("data_start_row")
                    if not data_start:
                        record_index += 1
                        continue
                    start_col = ''.join(ch for ch in data_start if ch.isalpha())
                    start_row = int(''.join(ch for ch in data_start if ch.isdigit()))
                    row_number = start_row + section_idx
                
                # Write fields for this record
                for column, field_name in fields.items():
                    cell_ref = f"{column}{row_number}"
                    value = self._resolve_table_value(field_name, record_dict, record_index + 1, info_obj)
                    self._write_cell(ws, cell_ref, value)
                
                record_index += 1
                total_filled += 1
        
        return total_filled


    def _write_table_totals(
        self,
        ws: Worksheet,
        table_spec: Dict[str, Any],
        start_row: int,
        count: int,
        info_obj: DSF_InfosGenerales,
        info_dict: Dict[str, Any],
    ) -> None:
        current_row = start_row + count
        totals = table_spec.get("totals_row")
        if totals:
            ws[f"{totals.get('label_column', 'A')}{current_row}"] = totals.get("label", "TOTAL")
            for column in totals.get("sum_columns", []):
                column_tot = 0
                for row_idx in range(start_row, start_row + count):
                    value = ws[f"{column}{row_idx}"].value or 0
                    column_tot += value if isinstance(value, (int, float)) else 0
                ws[f"{column}{current_row}"] = column_tot
            current_row += 1
        for optional_row in table_spec.get("optional_rows", []) or []:
            condition = optional_row.get("condition")
            if not self._evaluate_condition(condition, info_dict):
                continue
            label_col = optional_row.get("label_column", "A")
            ws[f"{label_col}{current_row}"] = optional_row.get("label", "")
            attr = condition.split()[0] if condition else ""
            attr_value = getattr(info_obj, attr, 0)
            for column in optional_row.get("sum_columns", []):
                ws[f"{column}{current_row}"] = attr_value
            if optional_row.get("bold"):
                cell = ws[f"{label_col}{current_row}"]
                cell.font = cell.font + Font(bold=True)
            current_row += 1

    def _resolve_value(
        self,
        info_obj: DSF_InfosGenerales,
        info_dict: Dict[str, Any],
        attribute: Optional[str],
        fmt: Optional[str],
    ) -> Any:
        value = getattr(info_obj, attribute, None) if attribute else None
        if attribute == "exercice_precedent_fin" and value is None:
            current_end = getattr(info_obj, "exercice_fin", None)
            if isinstance(current_end, date):
                try:
                    value = current_end.replace(year=current_end.year - 1)
                except ValueError:
                    # Handle leap-year edge case (e.g., 29/02 -> 28/02).
                    value = current_end.replace(year=current_end.year - 1, day=28)
        if fmt in {"date", "currency"}:
            return self._format_value(value, fmt)
        if fmt and hasattr(info_obj, fmt):
            return self._format_value(getattr(info_obj, fmt), _infer_format(getattr(info_obj, fmt)))
        return self._format_value(value, _infer_format(value))

    def _resolve_table_value(
        self,
        field_name: str,
        record_dict: Dict[str, Any],
        index: int,
        info_obj: DSF_InfosGenerales,
    ) -> Any:
        if field_name == "index":
            return index
        if field_name == "nom_prenoms":
            return " ".join(filter(None, [record_dict.get("nom"), record_dict.get("prenoms")])).strip()
        value = record_dict.get(field_name)
        return self._format_value(value, _infer_format(value))

    def _format_value(self, value: Any, fmt: Optional[str]) -> Any:
        if value is None:
            return ""
        if fmt == "date" and isinstance(value, date):
            return format_date(value)
        if fmt == "currency":
            return float(value)
        return value

    def _write_cell(self, ws: Worksheet, cell_ref: str, value: Any) -> bool:
        try:
            cell = ws[cell_ref]
            master = self._master_cell(ws, cell.coordinate)
            master.value = value
            return True
        except Exception:
            return False

    @staticmethod
    def _master_cell(ws: Worksheet, coord: str):
        cell = ws[coord]
        for merged in ws.merged_cells.ranges:
            if coord in merged:
                return ws[merged.start_cell.coordinate]
        return cell

    @classmethod
    def _is_in_protected_zone(cls, ws: Worksheet, cell_ref: str) -> bool:
        """
        P3-Prefill : vérifie si une cellule est dans la zone protégée d'une feuille.
        
        Les zones protégées sont définies dans PROTECTED_ZONES par nom de feuille normalisé.
        Une cellule est protégée si son numéro de ligne >= la limite définie.
        """
        sheet_key = cls._normalize_sheet_name(ws.title)
        min_protected_row = cls.PROTECTED_ZONES.get(sheet_key)
        if min_protected_row is None:
            return False
        try:
            row_num = int(''.join(ch for ch in cell_ref if ch.isdigit()))
            return row_num >= min_protected_row
        except (ValueError, TypeError):
            return False

    @staticmethod
    def _evaluate_condition(condition: Optional[str], info_dict: Dict[str, Any]) -> bool:
        if not condition:
            return False
        try:
            return bool(eval(condition, {}, info_dict))
        except Exception:
            return False

    @staticmethod
    def _normalize_sheet_name(name: str) -> str:
        return " ".join((name or "").upper().split())

    def _find_sheet(self, name: str):
        """Trouve une feuille par nom exact ou normalisé (tolère espaces finaux)."""
        if self.wb is None:
            return None
        if name in self.wb.sheetnames:
            return self.wb[name]
        key = self._normalize_sheet_name(name)
        for ws in self.wb.worksheets:
            if self._normalize_sheet_name(ws.title) == key:
                return ws
        return None


def _infer_format(value: Any) -> Optional[str]:
    if isinstance(value, date):
        return "date"
    if isinstance(value, (int, float)):
        return "currency"
    return None


__all__ = ["DSFGeneralPrefiller"]
