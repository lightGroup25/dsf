#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Test du remplissage des cellules composites du BILAN PAYSAGE"""

import sys
from pathlib import Path

# Ajouter src au path
sys.path.insert(0, str(Path(__file__).parent / "src"))

from dsf_bilan_composite_filler import fill_bilan_composite_cells, _resolve_note_reference, _find_value_in_note
from openpyxl import load_workbook

print("=" * 80)
print("TEST: REMPLISSAGE CELLULES COMPOSITES BILAN PAYSAGE")
print("=" * 80)

# Test du DSF 2023
dsf_2023 = Path(r"C:\Users\Emmaneul Ambadiang\Desktop\dsf\input\DSF GULFCAM 2023 V3.xlsx")

if dsf_2023.exists():
    print(f"\n[OK] Fichier trouve: {dsf_2023.name}")
    
    wb = load_workbook(dsf_2023)
    
    # Test 1: Résolution de références
    print("\n### Test 1: Resolution de references ###")
    test_refs = ["13", "3e", "3A", "14", "NOTE 7"]
    for ref in test_refs:
        resolved = _resolve_note_reference(ref)
        print(f"  {ref!r:8} -> {resolved!r}")
    
    # Test 2: Remplissage effectif
    print("\n### Test 2: Remplissage des cellules ###")
    n_filled = fill_bilan_composite_cells(wb)
    print(f"  Cellules remplies: {n_filled}")
    
    # Test 3: Vérifier les cellules
    print("\n### Test 3: Verification des cellules remplies ###")
    ws = wb["BILAN PAYSAGE"]
    cells_to_check = ["L12", "L13", "L14", "L15", "L16"]
    for cell_addr in cells_to_check:
        cell = ws[cell_addr]
        label = ws[f"I{cell_addr[1]}"].value
        val = cell.value
        filled = "[OK]" if val else "[NO]"
        print(f"  {cell_addr}: {filled} ({label}) = {val}")
    
    wb.close()
else:
    print(f"[NO] Fichier non trouve: {dsf_2023}")

print("\nFin du test!")
