import openpyxl
from pathlib import Path

def check_dsf(path):
    print(f"Checking {path}...")
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        print(f"Error loading workbook: {e}")
        return

    print(f"Sheets: {wb.sheetnames}")

    # Fuzzy find sheet
    def get_sheet_fuzzy(name):
        for s in wb.sheetnames:
            if s.strip() == name.strip():
                return wb[s]
        return None

    # Check Note 7 (Clients)
    ws = get_sheet_fuzzy("NOTE 7")
    if ws:
        print(f"NOTE 7 found as '{ws.title}'.")
        # naive check of content
        has_data = False
        for row in ws.iter_rows(min_row=5, max_row=50, values_only=True):
            if any(row):
                # print(f"Sample data in Note 7: {row}")
                has_data = True
                break
        print(f"NOTE 7 has data: {has_data}")
    else:
        print("NOTE 7 missing")

    # Check Note 16 (Emprunts)
    if "NOTE 16" in wb.sheetnames:
        ws = wb["NOTE 16"]
        print("NOTE 16 found.")
        has_data = False
        for row in ws.iter_rows(min_row=5, max_row=50, values_only=True):
            if any(row):
                has_data = True
                break
        print(f"NOTE 16 has data: {has_data}")

    # Check Balance Sheet (BILAN) - usually "R1" or "ACTIF" / "PASSIF"
    # Gulfcam template might use "BILAN ACTIF", "BILAN PASSIF" or just "ACTIF", "PASSIF"
    sheets = wb.sheetnames
    actif_sheet = next((s for s in sheets if "ACTIF" in s and "NOTE" not in s), None)
    if actif_sheet:
        print(f"Actif sheet: {actif_sheet}")
        ws = wb[actif_sheet]
        # Check a few cells
        print(f"Cell D10 value: {ws['D10'].value}") 
    
    passif_sheet = next((s for s in sheets if "PASSIF" in s and "NOTE" not in s), None)
    if passif_sheet:
        print(f"Passif sheet: {passif_sheet}")

if __name__ == "__main__":
    check_dsf("output/DSF_OUTPUT_2024.xlsx")
