import openpyxl
import glob
import os
from pathlib import Path

def verify_r2():
    # Find the latest output file
    list_of_files = glob.glob('output/DSF_OUTPUT_2024*.xlsx') 
    if not list_of_files:
        print("No output file found.")
        return
    
    latest_file = max(list_of_files, key=os.path.getctime)
    print(f"Verifying R2 in: {latest_file}")
    
    try:
        wb = openpyxl.load_workbook(latest_file, data_only=True, read_only=True)
    except Exception as e:
        print(f"Error loading workbook: {e}")
        return

    if "Fiche R2" not in wb.sheetnames:
        print("Sheet 'Fiche R2' not found.")
        return

    ws = wb["Fiche R2"]
    
    checks = {
        "A4": "GULFCAM",
        "A5": "M050900027774W",
        "B9": "08",
        "B11": "1",
        "B13": "2",
        "B29": "TRANSPORT MARITIME",
        "F30": "0 3 1 0 0",
        "M30": "16797523137", # Not formatted with spaces in verification, just raw value
        "O30": "0.94"
    }
    
    all_passed = True
    for coord, expected in checks.items():
        val = str(ws[coord].value)
        # Normalize float strings
        if val.endswith(".0"): val = val[:-2]
        
        if expected in val:
             print(f"[PASS] {coord}: {val}")
        else:
             print(f"[FAIL] {coord}: Expected '{expected}' in '{val}'")
             all_passed = False
             
    if all_passed:
        print("\nSUCCESS: Fiche R2 is correctly filled.")
    else:
        print("\nFAILURE: Some fields are missing or incorrect.")

if __name__ == "__main__":
    verify_r2()
