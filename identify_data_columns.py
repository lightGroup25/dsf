#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Identifier EXACTEMENT quelles colonnes contiennent les données
"""
import openpyxl
from openpyxl.utils import get_column_letter

balance_file = r"c:\Users\Emmaneul Ambadiang\Desktop\DSF\input\BALANCE AU 31 12 2024 VERSION DU 26 05 25.xlsx"

wb = openpyxl.load_workbook(balance_file, data_only=True)
ws = wb.active

print("="*150)
print("IDENTIFICATION EXACTE DES COLONNES DE DONNEES")
print("="*150)

# Pour chaque colonne, afficher:
# 1. Le header (rows 8-10)
# 2. Si elle a des données (row 20)

print("\nColonne | R8 Header | R10 Header | R20 Data (Account 10) | R25 Data (Account 11)")
print("-" * 150)

for col_num in range(1, 30):
    r8 = ws.cell(8, col_num).value
    r10 = ws.cell(10, col_num).value
    r20 = ws.cell(20, col_num).value
    r25 = ws.cell(25, col_num).value
    
    col_letter = get_column_letter(col_num)
    
    r8_str = str(r8)[:20] if r8 else ""
    r10_str = str(r10)[:15] if r10 else ""
    r20_str = str(r20)[:20] if r20 else ""
    r25_str = str(r25)[:20] if r25 else ""
    
    # Afficher seulement si y a du contenu
    if r8 or r10 or r20 or r25:
        print(f"{col_letter:>3s}({col_num:2d}) | {r8_str:20s} | {r10_str:15s} | {r20_str:20s} | {r25_str:20s}")

print("\n" + "="*150)
print("CONCLUSION - QUELLES COLONNES CONTIENNENT LES VRAIES VALEURS?")
print("="*150)

# Identifier les colonnes avec des données numeriques
print("\nColonnes avec donnees numeriques (row 20, account 10):")
for col_num in range(1, 30):
    cell = ws.cell(20, col_num)
    if cell.value and isinstance(cell.value, (int, float)):
        r8 = ws.cell(8, col_num).value or ""
        r10 = ws.cell(10, col_num).value or ""
        print(f"  Col {col_num} ({get_column_letter(col_num)}): Value={cell.value:15.0f} | R8={str(r8)[:40]} | R10={str(r10)[:20]}")

wb.close()
