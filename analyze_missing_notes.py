."""
Analyse pourquoi les notes 3C-12 et 14-34 ne sont pas remplies
"""
import openpyxl
from pathlib import Path
import json

OUTPUT_FILE = Path(r"output/DSF_OUTPUT_UI_20260219_140207.xlsx")

def analyze_sheet_content(ws):
    """Analyse le contenu d'une feuille"""
    filled_cells = 0
    total_cells = 0
    filled_data_cells = 0  # Cellules avec des nombres (pas des labels)
    
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                total_cells += 1
                # Essayer de déterminer si c'est une cellule de données ou un label
                if isinstance(cell.value, (int, float)):
                    filled_data_cells += 1
                    filled_cells += 1
                elif isinstance(cell.value, str) and cell.value.strip():
                    filled_cells += 1
    
    return {
        'filled_cells': filled_cells,
        'filled_data_cells': filled_data_cells,
        'total_cells': total_cells,
        'is_empty': filled_data_cells == 0
    }

def main():
    if not OUTPUT_FILE.exists():
        print(f"❌ Fichier non trouvé: {OUTPUT_FILE}")
        return
    
    print(f"📊 Analyse du fichier: {OUTPUT_FILE.name}\n")
    print("="*80)
    
    wb = openpyxl.load_workbook(OUTPUT_FILE, data_only=True)
    
    # Liste des notes à vérifier
    notes_to_check = []
    
    # Notes 3C à 12
    notes_3c_12 = ["NOTE  3C", "NOTE 4", "NOTE 5", "NOTE 6", "NOTE 7", 
                   "NOTE 8", "NOTE 9", "NOTE 10", "NOTE 11", "NOTE 12"]
    
    # Notes 14 à 34
    notes_14_34 = [f"NOTE {i}" for i in range(14, 35)]
    # Ajouter les variations de noms
    notes_14_34.extend(["NOTE 14  ", "NOTE 15A ", "NOTE 15B", 
                        "NOTE 16A", "NOTE 16B", "NOTE 16B BIS", "NOTE 16C",
                        "NOTE 17 ", "NOTE 18", "NOTE 19", "NOTE 20", "NOTE 21",
                        "NOTE 22", "NOTE 23", "NOTE 24", "NOTE 25", "NOTE 26",
                        "NOTE 27A", "NOTE 27B", "NOTE 28", "NOTE 29", "NOTE 30",
                        "NOTE 31", "NOTE 32", "NOTE 33", "NOTE 34"])
    
    notes_to_check = notes_3c_12 + notes_14_34
    
    # Vérifier uniquement les feuilles qui existent
    existing_notes = []
    for sheet_name in wb.sheetnames:
        sheet_upper = sheet_name.strip().upper()
        for note_check in notes_to_check:
            if note_check.strip().upper() in sheet_upper or sheet_upper in note_check.strip().upper():
                existing_notes.append(sheet_name)
                break
    
    # Analyser chaque note
    empty_notes = []
    filled_notes = []
    
    for sheet_name in existing_notes:
        ws = wb[sheet_name]
        stats = analyze_sheet_content(ws)
        
        if stats['is_empty']:
            empty_notes.append({
                'name': sheet_name,
                'stats': stats
            })
        else:
            filled_notes.append({
                'name': sheet_name,
                'stats': stats
            })
    
    # Afficher les résultats
    print(f"\n📋 NOTES VIDES (0 données numériques):")
    print("-"*80)
    if empty_notes:
        for note in empty_notes:
            print(f"  ❌ {note['name']:30} | Cellules remplies: {note['stats']['filled_cells']:4} | Données: {note['stats']['filled_data_cells']}")
    else:
        print("  ✓ Aucune note vide")
    
    print(f"\n📋 NOTES REMPLIES:")
    print("-"*80)
    if filled_notes:
        for note in filled_notes:
            print(f"  ✓ {note['name']:30} | Cellules remplies: {note['stats']['filled_cells']:4} | Données: {note['stats']['filled_data_cells']}")
    else:
        print("  ❌ Aucune note remplie")
    
    # Résumé
    print("\n" + "="*80)
    print("📊 RÉSUMÉ:")
    print(f"  Total notes analysées: {len(existing_notes)}")
    print(f"  Notes vides: {len(empty_notes)}")
    print(f"  Notes remplies: {len(filled_notes)}")
    
    # Vérifier si ces notes sont dans l'inventaire
    print("\n" + "="*80)
    print("🔍 VÉRIFICATION INVENTAIRE:")
    print("-"*80)
    
    inventory_file = Path("data/dsf_inventory.json")
    if inventory_file.exists():
        with open(inventory_file, 'r', encoding='utf-8') as f:
            inventory = json.load(f)
        
        sheets_in_inventory = list(inventory.get('sheets', {}).keys())
        
        for note in empty_notes:
            sheet_name = note['name']
            if sheet_name in sheets_in_inventory:
                sheet_data = inventory['sheets'][sheet_name]
                input_cols = len(sheet_data.get('input_columns', []))
                fields = len(sheet_data.get('fields', []))
                print(f"  ✓ {sheet_name:30} | Colonnes input: {input_cols:3} | Champs: {fields:4}")
            else:
                print(f"  ❌ {sheet_name:30} | NON TROUVÉ dans l'inventaire")
    
    wb.close()

if __name__ == "__main__":
    main()
