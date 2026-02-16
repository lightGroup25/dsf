"""Verify coverage for NOTE sheets and synthesis pages using inventory + output DSF."""
import re
import sys
from pathlib import Path
from openpyxl import load_workbook

sys.path.insert(0, "src")
from dsf_pipeline import DSFPipelineConfig
from dsf_inventory import DSFInventory


def _pick_output_file() -> Path:
    output_dir = Path("output")
    candidates = [p for p in output_dir.glob("*.xlsx") if not p.name.startswith("~$")]
    if not candidates:
        raise FileNotFoundError("No DSF output file found in output/")
    return max(candidates, key=lambda p: p.stat().st_mtime)


def _is_note_sheet(name: str) -> bool:
    return "NOTE" in name.upper()


def _is_synthese_sheet(name: str) -> bool:
    upper = name.upper()
    return "SYNTHESE" in upper or "SYNTH" in upper or "SYNTHÈSE" in upper


def main() -> None:
    dsf_path = _pick_output_file()
    inventory = DSFInventory.from_json(Path("data/dsf_inventory.json"))

    wb = load_workbook(dsf_path, data_only=True)

    note_sheets = [s for s in inventory.list_sheet_names() if _is_note_sheet(s)]
    synthese_sheets = [s for s in inventory.list_sheet_names() if _is_synthese_sheet(s)]

    def summarize(sheet_names):
        summary = []
        for sheet in sheet_names:
            if sheet not in wb.sheetnames:
                summary.append((sheet, 0, 0, 0.0, "MISSING"))
                continue
            ws = wb[sheet]
            fields = [f for f in inventory.iter_fields(sheet) if not f.locked and not f.has_formula]
            if not fields:
                summary.append((sheet, 0, 0, 0.0, "NO_FIELDS"))
                continue
            filled = 0
            for f in fields:
                cell = ws[f.cell]
                if cell.value not in (None, ""):
                    filled += 1
            total = len(fields)
            rate = (filled / total) * 100 if total else 0.0
            summary.append((sheet, filled, total, rate, "OK"))
        return summary

    note_summary = summarize(note_sheets)
    synth_summary = summarize(synthese_sheets)

    print(f"DSF output: {dsf_path}")
    print("\nNOTES:")
    print("-" * 80)
    for sheet, filled, total, rate, status in note_summary:
        print(f"{sheet:20} | {filled:4}/{total:4} | {rate:6.1f}% | {status}")

    print("\nSYNTHESE:")
    print("-" * 80)
    for sheet, filled, total, rate, status in synth_summary:
        print(f"{sheet:20} | {filled:4}/{total:4} | {rate:6.1f}% | {status}")

    wb.close()


if __name__ == "__main__":
    main()
