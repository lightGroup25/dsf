#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Test rapide de l'intégration du système de calculs dans le pipeline DSF.
"""

import logging
import sys
from pathlib import Path

# Add src to path
sys.path.insert(0, str(Path(__file__).parent / "src"))

from dsf_pipeline import DSFPipeline, DSFPipelineConfig

def test_pipeline_with_calculations():
    """Test pipeline DSF avec système de calculs intégré."""
    
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s %(levelname)s %(name)s %(message)s"
    )
    
    print("="*80)
    print("TEST PIPELINE DSF AVEC SYSTÈME DE CALCULS")
    print("="*80)
    
    # Configuration du pipeline
    config = DSFPipelineConfig(
        template_dsf=Path("templates/DSF Normal standard.xlsx"),
        balance_input=Path("input/BALANCE AU 31 12 2024 VERSION DU 26 05 25.xlsx"),
        inventory_path=Path("data/dsf_inventory.json"),
        prefill_mapping_path=Path("config/dsf_prefill_mapping.json"),
        dsf_output=Path("output/DSF_TEST_WITH_CALCULATIONS.xlsx"),
        report_json_path=Path("output/reports/test_calculations_report.json"),
        report_html_path=Path("output/reports/test_calculations_report.html"),
        filling_method="semantic",
        fuzzy_threshold=0.6,
        apply_calculations=True,  # ← ACTIVER CALCULS
        use_calculation_formulas=True,  # ← UTILISER FORMULES EXCEL
    )
    
    print(f"\n📋 Configuration:")
    print(f"   Template: {config.template_dsf}")
    print(f"   Balance: {config.balance_input}")
    print(f"   Output: {config.dsf_output}")
    print(f"   Mode remplissage: {config.filling_method}")
    print(f"   Calculs automatiques: {config.apply_calculations}")
    print(f"   Formules Excel: {config.use_calculation_formulas}")
    
    # Vérifier fichiers
    missing = []
    if not config.template_dsf.exists():
        missing.append(f"Template: {config.template_dsf}")
    if not config.balance_input.exists():
        missing.append(f"Balance: {config.balance_input}")
    if not config.inventory_path.exists():
        missing.append(f"Inventory: {config.inventory_path}")
    if not config.prefill_mapping_path.exists():
        missing.append(f"Prefill mapping: {config.prefill_mapping_path}")
    
    if missing:
        print(f"\n❌ Fichiers manquants:")
        for m in missing:
            print(f"   - {m}")
        return False
    
    print(f"\n✓ Tous les fichiers requis présents")
    
    # Exécuter pipeline
    print(f"\n🚀 Exécution du pipeline...")
    print(f"-" * 80)
    
    try:
        pipeline = DSFPipeline(config)
        artifacts = pipeline.run()
        
        print(f"\n" + "="*80)
        print(f"✅ PIPELINE COMPLÉTÉ AVEC SUCCÈS")
        print(f"="*80)
        print(f"\n📄 Fichiers générés:")
        print(f"   DSF Output: {artifacts.dsf_output}")
        print(f"   Template pré-rempli: {artifacts.prefilled_template}")
        if artifacts.report_json:
            print(f"   Rapport JSON: {artifacts.report_json}")
        if artifacts.report_html:
            print(f"   Rapport HTML: {artifacts.report_html}")
        
        # Vérifier que les calculs ont été appliqués
        print(f"\n🔍 Vérification des calculs...")
        if artifacts.dsf_output.exists():
            # Importer openpyxl pour vérifier
            try:
                import openpyxl
                wb = openpyxl.load_workbook(artifacts.dsf_output)
                
                formulas_found = 0
                for sheet_name in wb.sheetnames:
                    ws = wb[sheet_name]
                    for row in ws.iter_rows():
                        for cell in row:
                            if cell.data_type == 'f':  # Formula
                                formulas_found += 1
                
                wb.close()
                
                print(f"   ✓ {formulas_found} formules Excel détectées dans le DSF")
                
                if formulas_found > 0:
                    print(f"\n✨ SUCCÈS: Le système de calculs est intégré!")
                else:
                    print(f"\n⚠️ Aucune formule détectée (peut-être pas de lignes TOTAL avec données)")
                
            except ImportError:
                print(f"   Note: openpyxl non disponible pour vérification")
        
        return True
        
    except Exception as e:
        print(f"\n❌ ERREUR lors de l'exécution du pipeline:")
        print(f"   {type(e).__name__}: {e}")
        import traceback
        traceback.print_exc()
        return False


if __name__ == "__main__":
    success = test_pipeline_with_calculations()
    sys.exit(0 if success else 1)
