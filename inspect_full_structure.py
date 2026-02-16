"""
Check if balance file has detailed accounts that could fill NOTE sections
Also check all worksheets in DSF template
"""

import openpyxl
from pathlib import Path

print("="*70)
print("CHECKING BALANCE FILE STRUCTURE")
print("="*70)

balance_file = r'input/BALANCE AU 31 12 2024 VERSION DU 26 05 25.xlsx'
wb = openpyxl.load_workbook(balance_file)

print(f"\nBalance file sheets: {wb.sheetnames}")

# Check the first sheet
ws = wb.active
print(f"\nActive sheet: {ws.title}")
print(f"Max row: {ws.max_row}, Max col: {ws.max_column}")

# Count how many rows have data
data_rows = 0
for row_idx in range(1, ws.max_row + 1):
    account_num = ws[f'A{row_idx}'].value
    account_label = ws[f'D{row_idx}'].value
    if account_num and account_label:
        data_rows += 1

print(f"Rows with account data: {data_rows}")

print("\n" + "="*70)
print("CHECKING DSF TEMPLATE STRUCTURE")
print("="*70)

dsf_file = r'templates/DSF Normal standard.xlsx'
wb_dsf = openpyxl.load_workbook(dsf_file)

print(f"\nDSF template sheets: {wb_dsf.sheetnames}")

# Check each sheet
for sheet_name in wb_dsf.sheetnames:
    ws = wb_dsf[sheet_name]
    print(f"\nSheet: {sheet_name}")
    print(f"  Max row: {ws.max_row}, Max col: {ws.max_column}")
    
    # Look for NOTE markers
    note_count = 0
    for row_idx in range(1, min(500, ws.max_row + 1)):
        cell_value = ws[f'A{row_idx}'].value
        if cell_value and 'NOTE' in str(cell_value):
            note_count += 1
            if note_count <= 5:
                print(f"  Found: Row {row_idx} - {cell_value}")
    
    if note_count > 0:
        print(f"  Total NOTEs in this sheet: {note_count}")

print("\n" + "="*70)
