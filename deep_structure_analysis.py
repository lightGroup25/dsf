"""
Deep analysis: For each major NOTE, understand the REAL structure
What columns exist, what they mean, which should be filled
"""

import openpyxl

dsf_file = r'templates/DSF Normal standard.xlsx'
wb = openpyxl.load_workbook(dsf_file)

# Check a few key NOTEs in detail
sheets_to_deep_analyze = ['NOTE 3A', 'NOTE 3B', 'NOTE 4', 'NOTE 6', 'NOTE 13']

for sheet_name in sheets_to_deep_analyze:
    if sheet_name not in wb.sheetnames:
        continue
    
    ws = wb[sheet_name]
    print(f"\n{'='*70}")
    print(f"DEEP ANALYSIS: {sheet_name}")
    print(f"{'='*70}")
    
    # Show ALL rows from 1-20 to understand structure
    print(f"\nFirst 20 rows - all columns A through M:\n")
    
    for row_idx in range(1, 21):
        row_content = {}
        has_content = False
        
        for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M']:
            val = ws[f'{col}{row_idx}'].value
            if val:
                has_content = True
                row_content[col] = str(val)[:40]
        
        if has_content:
            print(f"Row {row_idx:2d}: ", end="")
            for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M']:
                if col in row_content:
                    print(f"[{col}:{row_content[col]}] ", end="")
            print()
    
    print(f"\nConclusion for {sheet_name}:")
    print("Which columns should be filled? (Based on headers you see)")
