
import openpyxl
import sys
import os

def inspect_sheet(filename, partial_sheet_name):
    print(f"Loading {filename}...")
    try:
        wb = openpyxl.load_workbook(filename, data_only=True)
    except Exception as e:
        print(f"Error loading workbook: {e}")
        return

    sheet_name = None
    for name in wb.sheetnames:
        if partial_sheet_name.lower() in name.lower():
            sheet_name = name
            break
    
    if not sheet_name:
        print(f"Sheet containing '{partial_sheet_name}' not found.")
        print(f"Available sheets: {wb.sheetnames}")
        return

    print(f"Inspecting sheet: {sheet_name}")
    ws = wb[sheet_name]

    # Search for key headers to locate tables
    for r in range(1, 60):
        for c in range(1, 10):
            val = ws.cell(row=r, column=c).value
            if val and isinstance(val, str):
                if "CONSEIL" in val.upper() or "ADMINISTRATION" in val.upper() or "MEMBRE" in val.upper():
                    print(f"Header Found at {ws.cell(row=r, column=c).coordinate}: {val}")

    # Inspect Rows 20-35 specifically for Board Members alignment
    print("\n--- Rows 20-35 Dump (Col A-F) ---")
    for r in range(20, 36):
        row_vals = []
        for c in range(1, 7): # A-F
            val = ws.cell(row=r, column=c).value
            row_vals.append(f"{get_column_letter(c)}{r}: {val or ''}")
        print(" | ".join(row_vals))

from openpyxl.utils import get_column_letter

if __name__ == "__main__":
    # Using the output file from previous run
    filename = r"c:\Users\Emmaneul Ambadiang\Desktop\DSF\output\DSF_OUTPUT_2024_20260213_182524.xlsx"
    if not os.path.exists(filename):
         filename = r"c:\Users\Emmaneul Ambadiang\Desktop\DSF\data\MODELE_O_01_DSF_S_N_SY_2024.xlsx"
    inspect_sheet(filename, "R3")
