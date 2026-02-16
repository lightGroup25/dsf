# -*- coding: utf-8 -*-
"""
Test du système de détection de formules dans les en-têtes de colonnes
"""
import logging
from pathlib import Path

import sys
sys.path.insert(0, str(Path(__file__).parent / "src"))

from column_formula_detector import ColumnFormulaDetector
from openpyxl import load_workbook

logging.basicConfig(
    level=logging.INFO,
    format="%(levelname)s - %(name)s - %(message)s"
)

logger = logging.getLogger(__name__)


def main():
    logger.info("="*70)
    logger.info("TEST COLUMN FORMULA DETECTOR")
    logger.info("="*70)
    
    template_path = Path("templates/DSF Normal standard.xlsx")
    
    if not template_path.exists():
        logger.error(f"Template not found: {template_path}")
        return
    
    logger.info(f"\nLoading template: {template_path}")
    wb = load_workbook(template_path, data_only=False)
    
    total_formulas = 0
    
    # Test sur chaque feuille
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        
        logger.info(f"\n{'='*70}")
        logger.info(f"Sheet: {sheet_name}")
        logger.info(f"{'='*70}")
        
        detector = ColumnFormulaDetector(ws)
        formulas = detector.detect()
        
        if not formulas:
            logger.info("  ✗ No column formulas detected")
            continue
        
        logger.info(f"  ✓ {len(formulas)} formula column(s) detected:")
        
        for col_letter, formula in formulas.items():
            logger.info(f"\n  Column {col_letter}:")
            logger.info(f"    Header: '{formula.header_text}'")
            logger.info(f"    Type: {formula.operation_type}")
            logger.info(f"    Pattern: {formula.formula_pattern}")
            logger.info(f"    References: {', '.join(formula.referenced_columns)}")
            
            # Exemples de formules générées
            logger.info(f"    Examples:")
            for row in [15, 20, 25]:
                excel_formula = formula.generate_formula(row)
                logger.info(f"      Row {row}: {excel_formula}")
            
            total_formulas += 1
    
    wb.close()
    
    logger.info(f"\n{'='*70}")
    logger.info(f"SUMMARY")
    logger.info(f"{'='*70}")
    logger.info(f"Total sheets scanned: {len(wb.sheetnames)}")
    logger.info(f"Total formula columns: {total_formulas}")
    
    if total_formulas > 0:
        logger.info(f"\n✅ SUCCESS: Column formula detection is working!")
    else:
        logger.warning(f"\n⚠️  WARNING: No formulas detected in template")


if __name__ == "__main__":
    main()
