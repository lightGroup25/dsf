#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import openpyxl

wb = openpyxl.load_workbook(r'templates\DSF Normal standard.xlsx', data_only=False)
ws = wb['NOTE 3A']

# Chercher toutes les formules dans NOTE 3A
print("RECHERCHE DE FORMULES DANS NOTE 3A:")
found_formulas = False
for row_num in range(1, 100):
    for col_num in range(1, 20):
        cell = ws.cell(row_num, col_num)
        if cell.value and isinstance(cell.value, str) and cell.value.startswith('='):
            found_formulas = True
            print(f"Row {row_num}, Col {col_num}: {cell.value}")

if not found_formulas:
    print("Aucune formule trouvee dans NOTE 3A")

# Chercher des colonnes avec TAUX ou %
print("\n\nRECHERCHE DE TAUX/POURCENTAGES DANS NOTE 3A:")
found_taux = False
for row_num in range(1, 100):
    for col_num in range(1, 20):
        cell = ws.cell(row_num, col_num)
        if cell.value:
            val_str = str(cell.value).upper()
            if 'TAUX' in val_str or '%' in val_str or 'POURCENTAGE' in val_str:
                found_taux = True
                print(f"Row {row_num}, Col {col_num}: {cell.value}")

if not found_taux:
    print("Aucun taux/pourcentage trouve dans NOTE 3A")

# Analyser completement les headers
print("\n\nHEADERS COMPLETS NOTE 3A - Rows 8-10:")
for row_num in [8, 9, 10]:
    print(f"\nRow {row_num}:")
    for col_num in range(1, 20):
        val = ws.cell(row_num, col_num).value
        if val:
            print(f"  Col {col_num}: {val}")

wb.close()

# Maintenant analyser aussi autres NOTES
wb2 = openpyxl.load_workbook(r'templates\DSF Normal standard.xlsx', data_only=False)

print("\n\n" + "="*80)
print("ANALYSE DES AUTRES NOTES - Chercher formules/pourcentages")
print("="*80)

for sheet_name in wb2.sheetnames:
    if 'NOTE' in sheet_name:
        ws = wb2[sheet_name]
        
        # Chercher formules
        has_formulas = False
        for row_num in range(1, 50):
            for col_num in range(1, 16):
                cell = ws.cell(row_num, col_num)
                if cell.value and isinstance(cell.value, str) and cell.value.startswith('='):
                    if not has_formulas:
                        print(f"\n{sheet_name} - FORMULES TROUVEES:")
                        has_formulas = True
                    print(f"  Row {row_num}, Col {col_num}: {cell.value[:60]}")
        
        # Chercher taux/pourcentages  
        has_taux = False
        for row_num in range(1, 50):
            for col_num in range(1, 16):
                cell = ws.cell(row_num, col_num)
                if cell.value:
                    val_str = str(cell.value).upper()
                    if 'TAUX' in val_str or '%' in val_str:
                        if not has_taux:
                            print(f"\n{sheet_name} - TAUX/% TROUVES:")
                            has_taux = True
                        print(f"  Row {row_num}, Col {col_num}: {cell.value}")

wb2.close()

print("\n\nDone!")
