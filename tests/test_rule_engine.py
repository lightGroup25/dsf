# -*- coding: utf-8 -*-
from __future__ import annotations

import json
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook

from balance_normalizer import NormalizedBalanceRow
from dsf_inventory import DSFInventory
from dsf_rule_config import load_rule_set
from dsf_rule_engine import Assignment, RuleEngine
from dsf_secure_writer import ProtectedDSFWriter

TEST_ROOT = Path(__file__).parent
DATA_ROOT = TEST_ROOT / "data"


def _normalized_row(
    compte: str,
    label: str,
    classe: str,
    debit: str = "0",
    credit: str = "0",
    buckets: list[str] | None = None,
) -> NormalizedBalanceRow:
    debit_decimal = Decimal(debit)
    credit_decimal = Decimal(credit)
    solde = debit_decimal - credit_decimal
    return NormalizedBalanceRow(
        compte=compte,
        label=label,
        classe=classe,
        natural_side="debit" if debit_decimal >= credit_decimal else "credit",
        solde_final=solde,
        debit_balance=debit_decimal,
        credit_balance=credit_decimal,
        dsf_buckets=buckets or [],
        metadata={"balance_side": "debit" if debit_decimal else "credit", "account_path": [f"Classe {classe}"]},
    )


class RuleEngineTests(unittest.TestCase):
    def setUp(self) -> None:
        self.inventory = DSFInventory.from_json(DATA_ROOT / "inventory_sample.json")
        self.rule_set = load_rule_set(DATA_ROOT / "rules_sample.yaml")
        self.engine = RuleEngine(self.rule_set, self.inventory)

    def test_rule_engine_assigns_accounts(self) -> None:
        rows = [
            _normalized_row("201", "FRAIS DÉVELOPPEMENT", "2", debit="100000", buckets=["BILAN_ACTIF"]),
            _normalized_row("101", "CAPITAL", "1", credit="50000", buckets=["BILAN_PASSIF"]),
            _normalized_row("601", "ACHATS", "6", debit="12000", buckets=["CR_CHARGES"]),
            _normalized_row("701", "VENTES", "7", credit="28000", buckets=["CR_PRODUITS"]),
        ]
        result = self.engine.apply(rows)
        self.assertEqual(len(result.assignments), 4)
        assignment_map = {f"{a.sheet}!{a.cell}": a for a in result.assignments}
        self.assertIn("BILAN PAYSAGE!D14", assignment_map)
        self.assertEqual(assignment_map["BILAN PAYSAGE!D14"].amount, Decimal("100000"))
        self.assertEqual(result.stats["resultat_exercice"], str(Decimal("16000")))
        self.assertIn("BILAN_ACTIF", result.stats["bucket_totals"])
        self.assertEqual(result.stats["rule_usage"]["actif_immos"]["hits"], 1)
        self.assertFalse(result.unmatched)

    def test_rule_engine_reports_unmatched(self) -> None:
        rows = [_normalized_row("801", "COUTS", "8", debit="1000", buckets=["ANALYTIQUE"])]
        result = self.engine.apply(rows)
        self.assertEqual(len(result.assignments), 0)
        self.assertEqual(len(result.unmatched), 1)
        self.assertIn("Aucune règle", result.unmatched[0].reason)

    def test_protected_writer_enforces_inventory(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_path = Path(tmp_dir)
            template = tmp_path / "template.xlsx"
            output = tmp_path / "output.xlsx"
            wb = Workbook()
            wb.remove(wb.active)
            ws_actif = wb.create_sheet("BILAN PAYSAGE")
            ws_actif["D14"] = None
            ws_passif = wb.create_sheet("COMPTE DE RESULTAT")
            ws_passif["D5"] = None
            wb.save(template)

            writer = ProtectedDSFWriter(template, output, self.inventory)
            assignment = Assignment(
                sheet="BILAN PAYSAGE",
                cell="D14",
                rule_id="actif_immos",
                label="IMMOBILISATIONS",
                amount=Decimal("123.45"),
            )
            writer.write_assignment(assignment)
            writer.save()
            writer.close()
            self.assertTrue(output.exists())
            self.assertEqual(len(writer.logs), 1)

            with self.assertRaises(ValueError):
                invalid_assignment = Assignment(
                    sheet="BILAN PAYSAGE",
                    cell="Z99",
                    rule_id="invalid",
                    label="",
                    amount=Decimal("10"),
                )
                writer.write_assignment(invalid_assignment)


if __name__ == "__main__":
    unittest.main()
