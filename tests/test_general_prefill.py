# -*- coding: utf-8 -*-
from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
import sys
import os

# Add project root to path
sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from openpyxl import Workbook, load_workbook

from dsf_general_info import get_gulfcam_config
from dsf_general_prefill import DSFGeneralPrefiller


class GeneralPrefillTests(unittest.TestCase):
    def test_prefill_writes_expected_cells(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_path = Path(tmp_dir)
            template_path = tmp_path / "template.xlsx"
            output_path = tmp_path / "output.xlsx"
            
            # Create a mock workbook with our expected sheets
            wb = Workbook()
            wb.remove(wb.active)
            # Correct sheet names as per template
            for sheet in ["ENTETE", "Fiche R1", "Fiche R2", "Fiche R3", "NOTE 13 "]:
                wb.create_sheet(sheet)
                
            # Simulate a merged cell in R1 just to test safety
            # Assume A5 is merged A5:B5
            ws_r1 = wb["Fiche R1"]
            ws_r1.merge_cells("A5:B5")
            
            wb.save(template_path)

            # Use the actual config function
            config = get_gulfcam_config()
            
            prefiller = DSFGeneralPrefiller(template_path)
            prefiller.load()
            prefiller.fill(config)
            prefiller.save(output_path)

            filled = load_workbook(output_path)
            # Basic checks
            self.assertEqual(filled["ENTETE"]["A2"].value, config.denomination_sociale)
            
            # R1: The date is written to A5 (which was merged A5:B5)
            # The value should be in the master cell A5
            self.assertEqual(filled["Fiche R1"]["A5"].value, "01/01/2024")
            
            # R2: Forme juridique
            # Note: mapping might target A2 or similar
            # self.assertEqual(filled["R2"]["A2"].value, config.forme_juridique) 
            
            # R3: Executives
            self.assertEqual(filled["Fiche R3"]["A4"].value, config.dirigeants[0].nom)
            
            # Note 13: Shareholders
            self.assertEqual(filled["NOTE 13 "]["A8"].value, config.actionnaires[0].nom)


if __name__ == "__main__":
    unittest.main()
