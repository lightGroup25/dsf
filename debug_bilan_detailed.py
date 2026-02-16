"""Debug DÉTAILLÉ - BILAN PAYSAGE uniquement"""

import openpyxl
from difflib import SequenceMatcher

# Charger la balance
balance_file = r'input/BALANCE AU 31 12 2024 VERSION DU 26 05 25.xlsx'
wb_bal = openpyxl.load_workbook(balance_file)
ws_bal = wb_bal.active

accounts = {}
for row_idx in range(1, ws_bal.max_row + 1):
    account_num = ws_bal[f'A{row_idx}'].value
    account_label = ws_bal[f'D{row_idx}'].value
    
    if not account_num or not account_label:
        continue
    
    try:
        int(str(account_num).split('.')[0])
        num_str = str(account_num).strip()
        label_str = str(account_label).strip()
        
        accounts[num_str] = {
            'label': label_str,
            'col14_opening': ws_bal[f'N{row_idx}'].value,
            'col21_movements': ws_bal[f'U{row_idx}'].value,
            'col27_closing': ws_bal[f'AA{row_idx}'].value,
        }
    except:
        pass

print("="*80)
print("DEBUG BILAN PAYSAGE - DÉTAILLÉ")
print("="*80)

# Charger DSF BILAN PAYSAGE
dsf_file = r'templates/DSF Normal standard.xlsx'
wb_dsf = openpyxl.load_workbook(dsf_file)
ws = wb_dsf['BILAN PAYSAGE']

# ÉTAPE 1: Détecter colonnes
print("\n[1] COLONNE DETECTION")
print("-" * 80)

column_info = {}
for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M', 'N']:
    # Get headers from rows 1-20
    headers = []
    for r in range(1, 21):
        val = ws[f'{col}{r}'].value
        if val:
            headers.append(str(val).upper())
    
    header_text = ' '.join(headers)
    
    # Classifier
    col_type = 'UNKNOWN'
    if any(x in header_text for x in ['AMORT', 'DEPREC', 'AMORTISSEMENTS']):
        col_type = 'AMORT'
    elif 'N-1' in header_text or 'EXERCICE AU 31/12/N-1' in header_text:
        col_type = 'COMPARATIF'
    elif ('NET' in header_text and 'CLOTURE' not in header_text  
          and 'BRUT' not in header_text and 'MONTANT' not in header_text):
        col_type = 'NET'
    elif 'CLOTURE' in header_text or 'A LA CLOTURE' in header_text:
        col_type = 'CLOSING'
    elif 'OUVERTURE' in header_text or ('A L' in header_text and 'BRUTE' in header_text):
        col_type = 'OPENING'
    elif 'BRUT' in header_text and 'CLOTURE' not in header_text and 'AMORT' not in header_text:
        col_type = 'OPENING'
    elif any(x in header_text for x in ['ACQUISITION', 'APPORT', 'CREATION', 'MOUVEMENT']):
        col_type = 'MOVEMENTS'
    
    print(f"Col {col}: {col_type:<12} ({header_text[:50]})")
    
    # Track untuk remplissage
    if col_type in ['OPENING', 'CLOSING', 'MOVEMENTS']:
        column_info[col] = {'type': col_type, 'header': header_text[:60]}

print(f"\nFillable columns: {list(column_info.keys())}")
print(f"Column info: {column_info}")

# ÉTAPE 2: Extraire libellés des lignes
print("\n[2] LIGNE LABELS (rows 12-25)")
print("-" * 80)

for row_idx in range(12, 26):
    col_a = ws[f'A{row_idx}'].value
    col_b = ws[f'B{row_idx}'].value
    
    line_label = str(col_b if col_b else col_a).strip() if (col_a or col_b) else None
    
    if not line_label or len(line_label) < 2:
        print(f"Row {row_idx}: SKIPPED (empty)")
        continue
    
    print(f"\nRow {row_idx}: {line_label}")
    
    # ÉTAPE 3: Fuzzy match avec balance
    best_score = 0
    best_account = None
    
    for acc_num, acc_data in accounts.items():
        score = SequenceMatcher(None, line_label.lower(), acc_data['label'].lower()).ratio()
        if score > best_score:
            best_score = score
            best_account = (acc_num, acc_data)
    
    if best_score >= 0.25:
        acc_num, acc_data = best_account
        print(f"  [MATCH] {acc_num}: {acc_data['label'][:50]}")
        print(f"    Score: {best_score:.3f}")
        print(f"    Data: Col14={acc_data['col14_opening']}, Col21={acc_data['col21_movements']}, Col27={acc_data['col27_closing']}")
        
        # Vérifier les valeurs
        if acc_data['col14_opening']:
            print(f"    -> Would fill OPENING column with {acc_data['col14_opening']}")
        else:
            print(f"    -> Would fill OPENING column but NO COL14 DATA")
    else:
        print(f"  [NO MATCH] Best: {best_account[1]['label'][:40]} @ {best_score:.3f}")
