"""
Enhanced Option 3 - Map balance accounts to ALL sheets including NOTE sections
"""

import openpyxl
from openpyxl.utils import get_column_letter
from difflib import SequenceMatcher
import json
from pathlib import Path
from datetime import datetime

# Configuration
BALANCE_FILE = r"input/BALANCE AU 31 12 2024 VERSION DU 26 05 25.xlsx"
DSF_TEMPLATE = r"templates/DSF Normal standard.xlsx"
OUTPUT_FILE = r"output/DSF_OPTION3_COMPLETE_WITH_NOTES.xlsx"

def load_balance_data(filepath):
    """Load balance file and extract ALL account data"""
    print(f"Loading balance file: {filepath}")
    
    try:
        wb = openpyxl.load_workbook(filepath)
        ws = wb.active
        
        accounts = {}
        
        # Find data range
        for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row, 
                                                    min_col=1, max_col=30, 
                                                    values_only=False), 1):
            if row_idx < 9:
                continue
                
            account_num_cell = row[0]
            account_num = account_num_cell.value
            
            if account_num is None:
                continue
            
            try:
                int(str(account_num).split('.')[0])
            except:
                continue
            
            account_label = row[3].value if row[3] else None
            opening_balance = row[13].value if row[13] else None
            movements = row[20].value if row[20] else None
            closing_balance = row[26].value if row[26] else None
            
            if account_label and (opening_balance or closing_balance):
                accounts[str(account_num).strip()] = {
                    'row': row_idx,
                    'label': str(account_label).strip(),
                    'opening': opening_balance,
                    'movements': movements,
                    'closing': closing_balance
                }
        
        print(f"✓ Loaded {len(accounts)} accounts from balance file")
        return accounts
        
    except Exception as e:
        print(f"✗ Error loading balance file: {e}")
        raise

def find_fillable_rows_in_sheet(ws):
    """Find all rows that should be filled with data (not headers/totals)"""
    fillable_rows = []
    
    for row_idx in range(1, min(ws.max_row + 1, 500)):
        col_a = ws[f'A{row_idx}'].value
        col_b = ws[f'B{row_idx}'].value
        
        if col_a is None and col_b is None:
            continue
        
        col_a_str = str(col_a).strip() if col_a else ""
        col_b_str = str(col_b).strip() if col_b else ""
        
        # Skip header/footer rows
        is_header = any(x in col_a_str.upper() for x in ['BILAN', 'DESSIGNATION', 'NATURE', 'LIBELLE', 'COMPTE', 'TOTAL', 'NOTE'])
        is_header = is_header or any(x in col_b_str.upper() for x in ['TOTAL', 'LIBELLE', 'DESIGNATION'])
        
        if not is_header and (col_a_str or col_b_str):
            fillable_rows.append({
                'row': row_idx,
                'col_a': col_a_str,
                'col_b': col_b_str,
                'label': col_b_str if col_b_str else col_a_str
            })
    
    return fillable_rows

def fuzzy_match_single(balance_label, target_label, threshold=0.4):
    """Match a single account with fuzzy similarity"""
    if not target_label:
        return 0
    
    balance_clean = balance_label.lower()
    target_clean = target_label.lower()
    
    return SequenceMatcher(None, balance_clean, target_clean).ratio()

def process_sheet(wb_template, sheet_name, balance_accounts, stats):
    """Process a single sheet in the template, filling with balance data"""
    try:
        ws = wb_template[sheet_name]
    except:
        return 0
    
    fillable_rows = find_fillable_rows_in_sheet(ws)
    cells_filled = 0
    
    for row_info in fillable_rows:
        row_idx = row_info['row']
        row_label = row_info['label']
        
        best_score = 0
        best_account = None
        
        # Find best matching account
        for account_num, account_data in balance_accounts.items():
            score = fuzzy_match_single(account_data['label'], row_label, threshold=0.3)
            if score > best_score:
                best_score = score
                best_account = account_data
        
        # If we found a good match, fill the cells
        if best_score >= 0.4 and best_account:
            try:
                # Try to fill columns with data
                col_d = ws[f'D{row_idx}']
                col_e = ws[f'E{row_idx}']
                col_j = ws[f'J{row_idx}']
                
                # Fill opening balance
                if best_account['opening'] is not None and not col_d.value:
                    col_d.value = best_account['opening']
                    cells_filled += 1
                
                # Fill movements
                if best_account['movements'] is not None and not col_e.value:
                    col_e.value = best_account['movements']
                    cells_filled += 1
                
                # Fill closing with formula
                if not col_j.value or (isinstance(col_j.value, str) and not col_j.value.startswith('=')):
                    col_j.value = f"=D{row_idx}+E{row_idx}"
                    cells_filled += 1
                
            except:
                pass
    
    return cells_filled

def main():
    print("="*70)
    print("OPTION 3 ENHANCED - FILL ALL SHEETS INCLUDING NOTES")
    print("="*70)
    
    try:
        # Load data
        balance_accounts = load_balance_data(BALANCE_FILE)
        
        # Load template
        print(f"Loading DSF template: {DSF_TEMPLATE}")
        wb_template = openpyxl.load_workbook(DSF_TEMPLATE)
        print(f"✓ Template has {len(wb_template.sheetnames)} sheets")
        
        # Process each sheet
        total_cells_filled = 0
        sheets_processed = 0
        stats = {}
        
        print(f"\nProcessing sheets...")
        
        for sheet_name in wb_template.sheetnames:
            cells_filled = process_sheet(wb_template, sheet_name, balance_accounts, stats)
            
            if cells_filled > 0:
                total_cells_filled += cells_filled
                sheets_processed += 1
                if sheets_processed <= 10 or 'NOTE' in sheet_name:
                    print(f"  ✓ {sheet_name}: {cells_filled} cells")
        
        # Save output
        print(f"\nSaving complete DSF to {OUTPUT_FILE}")
        Path(OUTPUT_FILE).parent.mkdir(parents=True, exist_ok=True)
        wb_template.save(OUTPUT_FILE)
        
        print("\n" + "="*70)
        print("SUMMARY - OPTION 3 COMPLETE")
        print("="*70)
        print(f"Sheets processed: {sheets_processed}/{len(wb_template.sheetnames)}")
        print(f"Total cells filled: {total_cells_filled}")
        print(f"Balance accounts available: {len(balance_accounts)}")
        print(f"\n✓ Output: {OUTPUT_FILE}")
        
    except Exception as e:
        print(f"\n✗ ERROR: {e}")
        import traceback
        traceback.print_exc()

if __name__ == "__main__":
    main()
