#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Test du remplissage des Notes 2, 3A, 3B, 3C existantes dans le DSF 2023"""

import sys
from pathlib import Path
from openpyxl import load_workbook

# Ajouter src au path
sys.path.insert(0, str(Path(__file__).parent / "src"))

print("=" * 100)
print("TEST: REMPLISSAGE DES NOTES 2, 3A, 3B, 3C")
print("=" * 100)

# Charger le DSF 2023 pour voir la structure des notes
dsf_2023 = Path(r"C:\Users\Emmaneul Ambadiang\Desktop\dsf\input\DSF GULFCAM 2023 V3.xlsx")

if not dsf_2023.exists():
    print(f"[NO] Fichier non trouve: {dsf_2023}")
    sys.exit(1)

wb = load_workbook(dsf_2023, data_only=True)

print(f"\n[OK] Fichier trouve: {dsf_2023.name}")

# Vérifier la présence des notes
notes_to_check = ["NOTE 2", "NOTE 3A", "NOTE 3B", "NOTE  3C"]

print("\n### PRESENCE DES NOTES ###")
for note_name in notes_to_check:
    found = False
    for sheet in wb.sheetnames:
        if sheet.upper().replace(" ", "") == note_name.upper().replace(" ", ""):
            found = True
            print(f"  [OK] {note_name:15} -> trouvee comme '{sheet}'")
            
            # Vérifier la structure interne
            ws = wb[sheet]
            print(f"       max_row={ws.max_row}, max_col={ws.max_column}")
            
            # Chercher headers (Brut, Amort, Net, N, N-1)
            headers = []
            for row in range(1, min(20, ws.max_row + 1)):
                for col in range(1, ws.max_column + 1):
                    cell = ws.cell(row, col)
                    if cell.value and isinstance(cell.value, str):
                        val_upper = str(cell.value).upper().strip()
                        if val_upper in ["BRUT", "AMORT", "NET", "N", "N-1", "DESIGNATION", "SOLDE", "CAPTION"]:
                            headers.append((row, col, val_upper))
            
            if headers:
                print(f"       Headers trouves: {len(headers)}")
                for row, col, val in headers[:5]:
                    print(f"         Row {row}, Col {col}: '{val}'")
            break
    
    if not found:
        print(f"  [NO] {note_name:15} -> NON TROUVEE")

# Vérifier le mapping des comptes
print("\n### MAPPING DES COMPTES ###")
from dsf_notes_filler import NOTE_ACCOUNT_MAPPING

for note_key in notes_to_check:
    note_short = note_key.replace(" ", "").replace("NOTE", "").strip()  # "2", "3A", "3B", "3C"
    # Chercher dans le mapping
    mapping_key = None
    for key in NOTE_ACCOUNT_MAPPING.keys():
        if key.upper().replace(" ", "") == note_key.upper().replace(" ", ""):
            mapping_key = key
            break
    
    if mapping_key:
        count = len(NOTE_ACCOUNT_MAPPING[mapping_key])
        accounts = [acc[0] for acc in NOTE_ACCOUNT_MAPPING[mapping_key]]
        print(f"  [OK] {note_key:15} -> {count} comptes: {', '.join(accounts[:3])}...")
    else:
        print(f"  [NO] {note_key:15} -> PAS DANS LE MAPPING")

print("\nFin du test!")
wb.close()
