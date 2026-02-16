from openpyxl import load_workbook
import sys

path = "input/BALANCE AU 31 12 2024 VERSION DU 26 05 25.xlsx"
print(f"Loading {path}...")
wb = load_workbook(path, read_only=True, data_only=True)
ws = wb.active

print("--- Rows 6-15 ---")
for i, row in enumerate(ws.iter_rows(min_row=6, max_row=15, values_only=True)):
    print(f"Row {i+6}: {row}")
