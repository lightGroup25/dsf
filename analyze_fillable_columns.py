"""
Identify which columns should be filled vs which are protected/calculated
Check the first few data rows to see the pattern
"""

import openpyxl

dsf_file = r'templates/DSF Normal standard.xlsx'
wb = openpyxl.load_workbook(dsf_file)
ws = wb['BILAN PAYSAGE']

print("BILAN PAYSAGE - Column analysis\n")

# Check what's in key rows and columns
print("Row 8 (Headers):")
for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K']:
    cell = ws[f'{col}8']
    print(f"  {col}8: {cell.value}")

print("\nRow 10 (Sub-headers):")
for col in ['D', 'E', 'F', 'G']:
    cell = ws[f'{col}10']
    print(f"  {col}10: {cell.value}")

print("\nRow 12 (First data row - IMMOBILISATIONS INCORPORELLES):")
for col in ['A', 'B', 'D', 'E', 'F', 'G']:
    cell = ws[f'{col}12']
    val = cell.value
    is_formula = isinstance(val, str) and val.startswith('=')
    print(f"  {col}12: Value={val}, IsFormula={is_formula}")

print("\nRow 30 (Totals row):")
for col in ['A', 'B', 'D', 'E', 'F', 'G']:
    cell = ws[f'{col}30']
    val = cell.value
    is_formula = isinstance(val, str) and val.startswith('=')
    print(f"  {col}30: Value={val}, IsFormula={is_formula}")

# Check all formulas in column E and F (likey amort and net columns)
print("\n\nFormulas found in worksheet:")
formula_count = {'D': 0, 'E': 0, 'F': 0, 'G': 0}
for row in ws.iter_rows(min_row=12, max_row=50):
    for cell in row:
        if cell.value and isinstance(cell.value, str) and cell.value.startswith('='):
            col = cell.column_letter
            if col in formula_count:
                formula_count[col] += 1
                if formula_count[col] <= 3:
                    print(f"  {cell.coordinate}: {cell.value}")

print("\nFormula count by column:")
for col, count in formula_count.items():
    print(f"  Col {col}: {count} formulas")
