"""See the ACTUAL header row and structure of balance file"""

import openpyxl

balance_file = r'input/BALANCE AU 31 12 2024 VERSION DU 26 05 25.xlsx'
wb_bal = openpyxl.load_workbook(balance_file)
ws_bal = wb_bal.active

print("="*100)
print("BALANCE FILE - HEADER ROW (Row 1)")
print("="*100)

for col_idx in range(1, 30):
    col_letter = openpyxl.utils.get_column_letter(col_idx)
    header = ws_bal[f'{col_letter}1'].value
    print(f"  Col {col_letter} ({col_idx:2d}): {header}")

print("\n" + "="*100)
print("ROW 1 (headers) + Row 2 (all columns for reference)")
print("="*100)
for row_idx in [1, 2]:
    print(f"\nRow {row_idx}:")
    for col_idx in range(1, 30):
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        val = ws_bal[f'{col_letter}{row_idx}'].value
        if val:
            print(f"  {col_letter} ({col_idx:2d}): {val}")

print("\n" + "="*100)
print("LAST FEW DATA ROWS (see how many columns have data)")
print("="*100)

last_row = ws_bal.max_row
for row_idx in [last_row - 2, last_row - 1, last_row]:
    print(f"\nRow {row_idx}:")
    print(f"  A: {ws_bal[f'A{row_idx}'].value}")
    print(f"  D: {ws_bal[f'D{row_idx}'].value}")
    
    # Find last non-empty column
    last_col = 1
    for col_idx in range(1, 50):
        val = ws_bal[f'{openpyxl.utils.get_column_letter(col_idx)}{row_idx}'].value
        if val:
            last_col = col_idx
    
    print(f"  Last non-empty column: {openpyxl.utils.get_column_letter(last_col)} (Col {last_col})")
    
    # Show all non-empty columns
    for col_idx in range(1, last_col + 1):
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        val = ws_bal[f'{col_letter}{row_idx}'].value
        if val:
            print(f"    {col_letter} ({col_idx:2d}): {val}")
