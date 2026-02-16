#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import openpyxl

wb = openpyxl.load_workbook(r'templates\DSF Normal standard.xlsx', data_only=False)
ws = wb['NOTE 3A']

print("="*120)
print("ANALYSE DES LIBELLES - DSF NOTE 3A")
print("="*120)

print("\nLibelles des colonnes (Row 8):")
for col_num in range(1, 12):
    val = ws.cell(8, col_num).value
    print(f"  Col {col_num:2d}: {str(val)[:80] if val else '[VIDE]'}")

print("\n" + "="*120)
print("Structure complete des en-tetes:")
print("="*120)

for col_num in [1, 2, 3, 4, 5, 6, 7, 8, 9, 10]:
    r8 = str(ws.cell(8, col_num).value)[:50] if ws.cell(8, col_num).value else ""
    r9 = str(ws.cell(9, col_num).value)[:50] if ws.cell(9, col_num).value else ""
    r10 = str(ws.cell(10, col_num).value)[:50] if ws.cell(10, col_num).value else ""
    
    print(f"\nCol {col_num}:")
    print(f"  R8: {r8}")
    print(f"  R9: {r9}")
    print(f"  R10: {r10}")

wb.close()

print("\n" + "="*120)
print("REMARQUES IMPORTANTES")
print("="*120)

print("""
1. Les libelles mentionnent EXPLICITEMENT:
   - "MONTANT BRUTE A L'OUVERTURE" = Col 4 (ouverture)
   - "MONTANT BRUT A LA CLOTURE" = Col 10 (cloture)
   - "ACQUISITIONS", "VIREMENTS", "REEVALUATION", "CESSIONS" = Col 5-9 (mouvements)

2. Cette structure est IDEALE pour le mapping:
   - On peut AVEC CERTITUDE remplir col 10 avec balance cloture
   - On DOIT LAISSER col 4-9 vides (pas de source adequate)

3. Les libeles explicites permettent d'affirmer que:
   - Le DSF demande VRAIMENT les details des mouvements
   - Ce n'est pas optional, c'est structure
   - Mais on n'a pas les donnees pour cela

CONCLUSION POUR OPTION 1:
   - Remplir col 10 (MONTANT BRUT A LA CLOTURE)
   - Laisser col 4-9 vides
   C'est honnete et transparent
""")
