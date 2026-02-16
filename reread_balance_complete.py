#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
RELECTURE COMPLETE du fichier BALANCE
Chercher TOUTES les colonnes et TOUTES les données
"""
import openpyxl

balance_file = r"c:\Users\Emmaneul Ambadiang\Desktop\DSF\input\BALANCE AU 31 12 2024 VERSION DU 26 05 25.xlsx"

wb = openpyxl.load_workbook(balance_file, data_only=True)
ws = wb.active

print("="*150)
print("RELECTURE COMPLETE - FICHIER BALANCE")
print("="*150)

# Vérifie max columns
print(f"\nMax row: {ws.max_row}")
print(f"Max column: {ws.max_column}")

# Afficher TOUS les headers (rows 8-10) pour TOUTES les colonnes
print("\n" + "="*150)
print("HEADERS COMPLETS (Rows 8-10) - TOUTES LES COLONNES")
print("="*150)

for row_num in [8, 9, 10]:
    print(f"\nRow {row_num}:")
    for col_num in range(1, ws.max_column + 1):
        cell = ws.cell(row_num, col_num)
        val = cell.value
        if val:
            print(f"  Col {col_num:2d}: {str(val)[:70]}")

# Afficher les données d'un compte exemple
print("\n" + "="*150)
print("DONNEES D'UN COMPTE EXEMPLE (Row 20 - Account 10)")
print("="*150)

for col_num in range(1, ws.max_column + 1):
    cell = ws.cell(20, col_num)
    val = cell.value
    if val is not None:
        print(f"  Col {col_num:2d}: {val}")

# Afficher structure des données (rows 8-10) avec meilleure formatage
print("\n" + "="*150)
print("STRUCTURE DETAILLEE - HIERARCHIE DES COLONNES")
print("="*150)

for col_num in range(1, min(35, ws.max_column + 1)):
    r8 = str(ws.cell(8, col_num).value)[:40] if ws.cell(8, col_num).value else ""
    r9 = str(ws.cell(9, col_num).value)[:40] if ws.cell(9, col_num).value else ""
    r10 = str(ws.cell(10, col_num).value)[:40] if ws.cell(10, col_num).value else ""
    
    if r8 or r9 or r10:
        print(f"\nCol {col_num:2d}:")
        if r8:
            print(f"  R8: {r8}")
        if r9:
            print(f"  R9: {r9}")
        if r10:
            print(f"  R10: {r10}")

wb.close()
