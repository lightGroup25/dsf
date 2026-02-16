"""
VALIDATION REPORT - V7 MAPPER
Vérifier les résultats, les formules appliquées, et la couverture
"""

import openpyxl
from pathlib import Path

output_file = r'output/DSF_FINAL_V7_OPTIMIZED.xlsx'
balance_file = r'input/BALANCE AU 31 12 2024 VERSION DU 26 05 25.xlsx'

# Load balance for reference
wb_bal = openpyxl.load_workbook(balance_file)
ws_bal = wb_bal.active

# Build account reference
accounts_ref = {}
for row_idx in range(1, ws_bal.max_row + 1):
    acc_num = ws_bal[f'A{row_idx}'].value
    acc_label = ws_bal[f'D{row_idx}'].value
    col_n = ws_bal[f'N{row_idx}'].value
    col_u = ws_bal[f'U{row_idx}'].value
    col_aa = ws_bal[f'AA{row_idx}'].value
    
    if acc_num and acc_label:
        try:
            int(str(acc_num).split('.')[0])
            accounts_ref[str(acc_label).strip()] = {
                'number': str(acc_num).strip(),
                'col_n': col_n,
                'col_u': col_u,
                'col_aa': col_aa,
            }
        except:
            pass

# Load DSF output
if not Path(output_file).exists():
    print(f"ERROR: {output_file} not found!")
    exit(1)

wb = openpyxl.load_workbook(output_file)

print("="*100)
print("VALIDATION REPORT - DSF_FINAL_V7_OPTIMIZED.xlsx")
print("="*100)

# Check key sheets
test_sheets = [
    ('BILAN PAYSAGE', 5),
    ('NOTE 3A', 131),
    ('NOTE 3B', 57),
    ('Note 1', 37),
]

for sheet_name, expected_count in test_sheets:
    if sheet_name not in wb.sheetnames:
        print(f"\n[MISSING] {sheet_name}")
        continue
    
    ws = wb[sheet_name]
    print(f"\n{'='*100}")
    print(f"Sheet: {sheet_name} (expected ~{expected_count} cells)")
    print(f"{'='*100}")
    
    filled_count = 0
    samples = []
    
    for row in ws.iter_rows(min_row=12, max_row=50):
        for col_idx, cell in enumerate(row, 1):
            if cell.value and isinstance(cell.value, (int, float)) and cell.value > 1000:
                filled_count += 1
                if len(samples) < 5:
                    col_letter = openpyxl.utils.get_column_letter(col_idx)
                    line_label = ws[f'B{cell.row}'].value or ws[f'A{cell.row}'].value
                    samples.append({
                        'cell': f'{col_letter}{cell.row}',
                        'value': cell.value,
                        'line': str(line_label)[:40] if line_label else 'N/A'
                    })
    
    print(f"Total cells with data: {filled_count}")
    print(f"\nSample cells filled:")
    for s in samples:
        print(f"  {s['cell']:6s}: {s['value']:15,.0f} (line: {s['line']})")

# Summary statistics
print(f"\n{'='*100}")
print("SUMMARY - ALL SHEETS")
print(f"{'='*100}")

all_cells = 0
sheets_with_data = 0
sheet_counts = []

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    sheet_cells = 0
    
    for row in ws.iter_rows():
        for cell in row:
            if cell.value and isinstance(cell.value, (int, float)) and cell.value > 1000:
                sheet_cells += 1
    
    if sheet_cells > 0:
        all_cells += sheet_cells
        sheets_with_data += 1
        sheet_counts.append((sheet_name, sheet_cells))

print(f"Total sheets: {len(wb.sheetnames)}")
print(f"Sheets with data: {sheets_with_data}")
print(f"Total cells with numeric data > 1000: {all_cells}")

print(f"\nTop 10 sheets by cell count:")
for sheet, count in sorted(sheet_counts, key=lambda x: -x[1])[:10]:
    print(f"  {sheet:40s}: {count:4d} cells")

# Verify formulas were applied (CLOSING = OPENING + MOVEMENTS)
print(f"\n{'='*100}")
print("VERIFY FORMULA APPLICATION (CLOSING = OPENING + MOVEMENTS)")
print(f"{'='*100}")

# Check NOTE 3A specifically (should have all 3 columns)
if 'NOTE 3A' in wb.sheetnames:
    ws = wb['NOTE 3A']
    print(f"\nNOTE 3A - Check formula application:")
    
    # Expected columns: D (OPENING), E (MOVEMENTS), J (CLOSING)
    col_d_count = 0
    col_e_count = 0
    col_j_count = 0
    formula_verified = 0
    
    for row_idx in range(12, 40):
        d_val = ws[f'D{row_idx}'].value
        e_val = ws[f'E{row_idx}'].value
        j_val = ws[f'J{row_idx}'].value
        
        if d_val and isinstance(d_val, (int, float)):
            col_d_count += 1
        if e_val and isinstance(e_val, (int, float)):
            col_e_count += 1
        if j_val and isinstance(j_val, (int, float)):
            col_j_count += 1
        
        # Check if formula applies
        if d_val and e_val and j_val:
            if isinstance(d_val, (int, float)) and isinstance(e_val, (int, float)) and isinstance(j_val, (int, float)):
                expected = d_val + e_val
                if abs(j_val - expected) < 1:  # Allow for rounding
                    formula_verified += 1
    
    print(f"  Column D (OPENING): {col_d_count} cells")
    print(f"  Column E (MOVEMENTS): {col_e_count} cells")
    print(f"  Column J (CLOSING): {col_j_count} cells")
    print(f"  Formula verified (CLOSING = OPENING + MOVEMENTS): {formula_verified} rows")

print(f"\n{'='*100}")
print("[OK] Validation Report Complete")
print("="*100)
