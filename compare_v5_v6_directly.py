#!/usr/bin/env python3
"""
Direct comparison: V5 vs V6 by reading Excel files
Counts filled cells directly
"""

from openpyxl import load_workbook
from pathlib import Path

def count_filled_cells(xlsx_path):
    """Count non-empty cells (excluding headers)"""
    wb = load_workbook(xlsx_path, data_only=True)
    
    total_filled = 0
    filled_by_sheet = {}
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        
        sheet_filled = 0
        for row in ws.iter_rows(min_row=11, values_only=False):  # Skip headers (rows 1-10)
            for cell in row:
                # Count cells with numeric values (account amounts)
                if cell.value is not None and isinstance(cell.value, (int, float)):
                    if cell.value != 0:  # Only count non-zero
                        sheet_filled += 1
                        total_filled += 1
        
        if sheet_filled > 0:
            filled_by_sheet[sheet_name] = sheet_filled
    
    return total_filled, filled_by_sheet

print("=" * 80)
print("DIRECT EXCEL FILE COMPARISON")
print("=" * 80)

v5_path = Path("output/DSF_OUTPUT_2024_20260213_190157.xlsx")  # Last known V5
v6_path = Path("output/test_semantic_v6_output.xlsx")

if not v5_path.exists():
    # Try the last semantic test output
    v5_path = Path("output/test_semantic_output.xlsx")

print(f"\nV5: {v5_path}")
if v5_path.exists():
    v5_total, v5_by_sheet = count_filled_cells(v5_path)
    print(f"✓ Filled cells: {v5_total}")
    print(f"✓ Sheets with hits: {len(v5_by_sheet)}")
    top_v5 = sorted(v5_by_sheet.items(), key=lambda x: -x[1])[:3]
    print(f"  Top: {', '.join([f'{k}({v})' for k,v in top_v5])}")
else:
    print("✗ File not found")
    v5_total = 0

print(f"\nV6: {v6_path}")
if v6_path.exists():
    v6_total, v6_by_sheet = count_filled_cells(v6_path)
    print(f"✓ Filled cells: {v6_total}")
    print(f"✓ Sheets with hits: {len(v6_by_sheet)}")
    top_v6 = sorted(v6_by_sheet.items(), key=lambda x: -x[1])[:3]
    print(f"  Top: {', '.join([f'{k}({v})' for k,v in top_v6])}")
else:
    print("✗ File not found")
    v6_total = 0

print(f"\n{'=' * 80}")
print("COMPARISON")
print(f"{'=' * 80}")
print(f"V5: {v5_total} cells")
print(f"V6: {v6_total} cells")

if v6_total > v5_total:
    improvement = v6_total - v5_total
    improvement_pct = (improvement / v5_total * 100) if v5_total > 0 else 0
    print(f"\n✅ IMPROVEMENT: +{improvement} cells (+{improvement_pct:.1f}%)")
elif v6_total == v5_total:
    print(f"\n⚠ NO CHANGE: Hybrid strategies did not help")
else:
    print(f"\n❌ DEGRADATION: {v6_total - v5_total} cells")
