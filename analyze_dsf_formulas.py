"""Analyze formula patterns in DSF template - headers, lines, and cells"""

import openpyxl
import re
from collections import defaultdict

dsf_file = r'templates/DSF Normal standard.xlsx'
wb = openpyxl.load_workbook(dsf_file)

print("="*100)
print("ANALYZE FORMULAS IN DSF TEMPLATE")
print("="*100)

formula_patterns = defaultdict(list)
cell_formulas = []

# Check a few important sheets
test_sheets = ['BILAN PAYSAGE', 'NOTE 3A', 'NOTE 3B', 'NOTE 28']

for sheet_name in test_sheets:
    if sheet_name not in wb.sheetnames:
        continue
    
    ws = wb[sheet_name]
    print(f"\n[{sheet_name}]")
    print("-" * 100)
    
    # Look for formulas in cells
    formula_count = 0
    for row in ws.iter_rows(min_row=1, max_row=50):
        for cell in row:
            if cell.value and isinstance(cell.value, str) and cell.value.startswith('='):
                formula_count += 1
                cell_formulas.append({
                    'sheet': sheet_name,
                    'cell': cell.coordinate,
                    'formula': cell.value
                })
                print(f"  Cell {cell.coordinate}: {cell.value}")
    
    if formula_count == 0:
        print("  (no formulas found in cells)")
    
    # Look for patterns in headers that suggest calculations
    print(f"\n  Header analysis:")
    for row_idx in range(1, 20):
        for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K']:
            cell = ws[f'{col}{row_idx}']
            if cell.value:
                val_str = str(cell.value).upper()
                # Look for keywords suggesting calculations
                if any(x in val_str for x in ['TOTAL', '+', '-', 'MOINS', 'PLUS', 'RESULTAT', 'DIFF', 'ECART', 'NET']):
                    # Check for formulaic hints
                    if 'TOTAL' in val_str or 'MOINS' in val_str or 'PLUS' in val_str:
                        print(f"    {col}{row_idx}: {str(cell.value)[:60]}")

print("\n" + "="*100)
print("FORMULA PATTERNS DISCOVERED")
print("="*100)

if cell_formulas:
    print(f"\nCell formulas found: {len(cell_formulas)}")
    for f in cell_formulas:
        print(f"  {f['sheet']} - {f['cell']}: {f['formula']}")
else:
    print("No cell formulas found!")

# Now analyze the balance file to identify formula patterns there
print("\n" + "="*100)
print("BALANCE STRUCTURE - For formula inference")
print("="*100)

balance_file = r'input/BALANCE AU 31 12 2024 VERSION DU 26 05 25.xlsx'
wb_bal = openpyxl.load_workbook(balance_file)
ws_bal = wb_bal.active

print("\nBalance Column Headers (Row 8-10):")
for row_idx in range(8, 11):
    print(f"\nRow {row_idx}:")
    for col_idx in range(1, 26):
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        val = ws_bal[f'{col_letter}{row_idx}'].value
        if val:
            print(f"  {col_letter} ({col_idx:2d}): {val}")

# Verify the known formula: Col27 = Col14 + Col21
print("\n" + "="*100)
print("VERIFY FORMULA: Col27 = Col14 + Col21")
print("="*100)

print("\nSample accounts:")
for row_idx in range(13, 23):
    acc_num = ws_bal[f'A{row_idx}'].value
    col14 = ws_bal[f'N{row_idx}'].value
    col21 = ws_bal[f'U{row_idx}'].value
    col27 = ws_bal[f'AA{row_idx}'].value
    
    if col14 or col21 or col27:
        print(f"\nRow {row_idx}: Account {acc_num}")
        print(f"  Col14 (N):  {col14}")
        print(f"  Col21 (U):  {col21}")
        print(f"  Col27 (AA): {col27}")
        
        if col14 and col21 and col27:
            expected = col14 + col21
            match = "✓ FORMULA OK" if expected == col27 else f"✗ MISMATCH: {col14} + {col21} = {expected}, but got {col27}"
            print(f"  {match}")
