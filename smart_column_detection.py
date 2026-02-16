"""
Smart column detection - analyze sheet structure properly
Determine which columns in each sheet should be filled with balance file data
"""

import openpyxl
from pathlib import Path
import json

dsf_file = r'templates/DSF Normal standard.xlsx'
wb = openpyxl.load_workbook(dsf_file)

# Define rules for each column type
COLUMN_DETECTION_RULES = {
    'fillable': {
        'keywords': [
            'MONTANT BRUTE',  # Opening balance (ACTIF)
            'ACQUISITIONS',   # Purchases/movements
            'VIREMENTS',      # Transfers
            'REEVALUATION',   # Revaluations  
            'CESSIONS',       # Sales/disposals
            'VARIATIONS',     # Stock variations
            'CREATIONS',      # Creations
            'MOUVEMENTS'      # Movements
        ]
    },
    'protected': {
        'keywords': [
            'NET',            # Calculated net
            'AMORT',          # Depreciation (calculated)
            'CLÔTURE',        # Closing (calculated)
            'CLOTURE',        # Alternative spelling
            'N-1',            # Prior year
            'COMPARATIF',     # Comparison
            'EXERCICE PRECEDENT',  # Prior period
            'TOTAL'           # Totals
        ]
    }
}

def detect_column_type(header_text):
    """Determine if column should be filled or is protected"""
    if not header_text:
        return 'unknown'
    
    text_upper = str(header_text).upper()
    
    # Check protected first (higher priority)
    for keyword in COLUMN_DETECTION_RULES['protected']['keywords']:
        if keyword in text_upper:
            return 'protected'
    
    # Check fillable
    for keyword in COLUMN_DETECTION_RULES['fillable']['keywords']:
        if keyword in text_upper:
            return 'fillable'
    
    # Special case: "BRUT" without other keywords = fillable (opening balance)
    if 'BRUT' in text_upper and not any(x in text_upper for x in ['NET', 'AMORT']):
        return 'fillable'
    
    return 'unknown'

def analyze_sheet_structure(ws, sheet_name):
    """Analyze a sheet and return its fillable columns"""
    
    # Find header row
    header_row = None
    for row_idx in range(1, 30):
        row_text = []
        for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M']:
            val = ws[f'{col}{row_idx}'].value
            if val:
                row_text.append(str(val)[:20])
        
        row_str = ' '.join(row_text).upper()
        # Look for row that has data column headers
        if any(x in row_str for x in ['MONTANT', 'BRUT', 'ACQUISITIONS', 'VIREMENTS', 'OUVERTURE']):
            header_row = row_idx
            break
    
    if not header_row:
        return {'error': 'No header row found'}
    
    # Analyze each column
    fillable_cols = []
    protected_cols = []
    unknown_cols = []
    
    for col in ['D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M']:
        header = ws[f'{col}{header_row}'].value
        if not header:
            continue
        
        col_type = detect_column_type(header)
        
        if col_type == 'fillable':
            fillable_cols.append((col, str(header)[:50]))
        elif col_type == 'protected':
            protected_cols.append((col, str(header)[:50]))
        else:
            unknown_cols.append((col, str(header)[:50]))
    
    return {
        'sheet': sheet_name,
        'header_row': header_row,
        'fillable': fillable_cols,
        'protected': protected_cols,
        'unknown': unknown_cols
    }

# Analyze multiple sheets
sheets_to_analyze = ['BILAN PAYSAGE', 'NOTE 3B', 'NOTE 4', 'NOTE 6', 'NOTE 7', 'NOTE 13']

results = []
for sheet_name in sheets_to_analyze:
    if sheet_name not in wb.sheetnames:
        continue
    
    ws = wb[sheet_name]
    result = analyze_sheet_structure(ws, sheet_name)
    results.append(result)
    
    print(f"\n{'='*70}")
    print(f"Sheet: {sheet_name}")
    print(f"{'='*70}")
    
    if 'error' in result:
        print(f"  {result['error']}")
        continue
    
    print(f"\nHeader row: {result['header_row']}")
    
    if result['fillable']:
        print(f"\n✓ FILLABLE columns ({len(result['fillable'])}):")
        for col, header in result['fillable']:
            print(f"  {col}: {header}")
    
    if result['protected']:
        print(f"\n✗ PROTECTED columns ({len(result['protected'])}):")
        for col, header in result['protected']:
            print(f"  {col}: {header}")
    
    if result['unknown']:
        print(f"\n? UNKNOWN columns ({len(result['unknown'])}):")
        for col, header in result['unknown']:
            print(f"  {col}: {header}")

# Save results
output_path = r'output/column_detection_results.json'
Path(output_path).parent.mkdir(parents=True, exist_ok=True)
with open(output_path, 'w', encoding='utf-8') as f:
    json.dump(results, f, indent=2, ensure_ascii=False, default=str)

print(f"\n\nResults saved to {output_path}")
