import openpyxl

wb = openpyxl.load_workbook(r'templates/DSF Normal standard.xlsx')
ws = wb['NOTE 3A']

print("NOTE 3A - HEADERS (rows 1-10):")
for r in range(1, 11):
    for c in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K']:
        val = ws[f'{c}{r}'].value
        if val:
            print(f"  {c}{r}: {str(val)[:50]}")

print("\n\nNOTE 3A - FIRST DATA ROWS (rows 12-20):")
for r in range(12, 21):
    row_vals = {}
    for c in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K']:
        val = ws[f'{c}{r}'].value
        if val:
            row_vals[c] = str(val)[:40]
    
    if row_vals:
        print(f"Row {r}: {row_vals}")
