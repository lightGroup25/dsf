
import openpyxl

def find_text_in_sheet(filename, sheet_name, search_text):
    try:
        wb = openpyxl.load_workbook(filename, data_only=True)
        if sheet_name not in wb.sheetnames:
            print(f"Sheet {sheet_name} not found.")
            return

        ws = wb[sheet_name]
        print(f"Searching '{search_text}' in {sheet_name}...")
        
        for row in ws.iter_rows():
            for cell in row:
                if cell.value and search_text.lower() in str(cell.value).lower():
                    print(f"Found '{search_text}' at {cell.coordinate}: {cell.value}")
                    return cell.coordinate
        print("Not found.")

    except Exception as e:
        print(f"Error: {e}")

filename = "DSF Normal standard.xlsx"
find_text_in_sheet(filename, "Fiche R2", "ACTIVITÉ")
find_text_in_sheet(filename, "Fiche R3", "Nom")
find_text_in_sheet(filename, "Fiche R3", "Qualité")
