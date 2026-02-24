#!/usr/bin/env python3
"""
Get COMPLETE header structure with better formatting
Show rows 8-10 with ACTUAL column boundaries
"""
import openpyxl

balance_file = r"c:\Users\Emmaneul Ambadiang\Desktop\DSF\input\BALANCE AU 31 12 2024 VERSION DU 26 05 25.xlsx"

wb = openpyxl.load_workbook(balance_file, data_only=True)
ws = wb.active

print("=" * 150)
print("COMPLETE HEADER STRUCTURE (Rows 8-10)")
print("=" * 150)

print("\nROW 8 (Main section headers):")
print("-" * 150)
for col in range(1, 35):
    val = ws.cell(8, col).value
    if val:
        print(f"  Col {col:2d}: {val}")

print("\nROW 9 (Sub-headers row 1):")
print("-" * 150)
for col in range(1, 35):
    val = ws.cell(9, col).value
    if val:
        print(f"  Col {col:2d}: {val}")

print("\nROW 10 (Sub-headers row 2 - Débit/Crédit):")
print("-" * 150)
for col in range(1, 35):
    val = ws.cell(10, col).value
    if val:
        print(f"  Col {col:2d}: {val}")

# Now let's show the grid more visually
print("\n" + "=" * 150)
print("VISUAL GRID - Rows 8-10 All Columns")
print("=" * 150)
for col in range(1, 35):
    r8 = ws.cell(8, col).value or ""
    r9 = ws.cell(9, col).value or ""
    r10 = ws.cell(10, col).value or ""
    
    # Only print if there's content
    if r8 or r9 or r10:
        r8_str = str(r8)[:15] if r8 else ""
        r9_str = str(r9)[:15] if r9 else ""
        r10_str = str(r10)[:15] if r10 else ""
        print(f"Col {col:2d} | R8: {r8_str:15s} | R9: {r9_str:15s} | R10: {r10_str:15s}")

# Show what's in row 11 (first data row)
print("\n" + "=" * 150)
print("FIRST DATA ROW (Row 11)")
print("=" * 150)
for col in range(1, 10):
    val = ws.cell(11, col).value
    print(f"  Col {col}: {val}")

print("\n" + "=" * 150)
print("KEY COLUMNS IDENTIFIED")
print("=" * 150)
print("""
From Row 10 analysis:
1. Col 10: Débit (under "Mouvements au 31/12/")
2. Col 13: Crédit (under "Mouvements au 31/12/")
3. Col 17: Débit (under "Mouvements")
4. Col 20: Crédit (under "Mouvements")
5. Col 23: Débit (under "Soldes cumulés")
6. Col 26: Crédit (under "Soldes cumulés")

Question: Are cols 14, 21, 27 actual data columns or merged/hidden?
""")

# Check if there are merged cells
print("\nMerged cells in balance file:")
for merged_range in ws.merged_cells.ranges:
    print(f"  {merged_range}")

wb.close()
