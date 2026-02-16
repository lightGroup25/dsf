import openpyxl
import os

output_path = r"c:\Users\Emmaneul Ambadiang\Desktop\DSF\output\DSF_OUTPUT_2024.xlsx"

if not os.path.exists(output_path):
    print(f"File not found: {output_path}")
    exit(1)

wb = openpyxl.load_workbook(output_path, data_only=True)

def inspect_sheet(sheet_name, row_range, col_idx):
    if sheet_name not in wb.sheetnames:
        print(f"Sheet {sheet_name} not found")
        return

    ws = wb[sheet_name]
    print(f"\n--- {sheet_name} (Col {col_idx}) ---")
    
    found_data = False
    for row in row_range:
        cell_val = ws[f"{col_idx}{row}"].value
        if cell_val is not None and cell_val != 0 and cell_val != "":
             print(f"Row {row}: {cell_val}")
             found_data = True
    
    if not found_data:
        print("No data found in specified range.")

# Inspect Note 3A (Fixed Assets)
# Assuming Forced write to lines 15-30, col J
inspect_sheet("NOTE 3A", range(15, 35), "J")

# Inspect Note 3C (Amortization)
# Assuming Forced write to lines 15-30, col J
inspect_sheet("NOTE 3C", range(15, 35), "J")

print("\nFinished inspection.")
