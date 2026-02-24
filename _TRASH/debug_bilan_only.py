"""Debug BILAN PAYSAGE spécifiquement - voir structure exacte"""

import openpyxl
from difflib import SequenceMatcher

# Charger balance
balance_file = r'input/BALANCE AU 31 12 2024 VERSION DU 26 05 25.xlsx'
wb_bal = openpyxl.load_workbook(balance_file)
ws_bal = wb_bal.active

balance_accounts = {}
for r in range(1, ws_bal.max_row + 1):
    acc_num = ws_bal[f'A{r}'].value
    acc_label = ws_bal[f'D{r}'].value
    if acc_num and acc_label:
        try:
            int(str(acc_num).split('.')[0])
            balance_accounts[str(acc_label).strip()] = str(acc_num).strip()
        except:
            pass

print("="*70)
print("BILAN PAYSAGE - DEBUG COMPLET")
print("="*70)

# Charger DSF BILAN PAYSAGE
dsf_file = r'templates/DSF Normal standard.xlsx'
wb_dsf = openpyxl.load_workbook(dsf_file)
ws_dsf = wb_dsf['BILAN PAYSAGE']

print("\nColonnes détectées (Row 10 - en-têtes principal):")
print("-" * 70)
for col in ['D', 'E', 'F', 'G', 'H', 'I', 'J']:
    header = ws_dsf[f'{col}10'].value
    print(f"  Col {col}: {header}")

print("\nLibellés des lignes (colonnes B et A, rows 12-25):")
print("-" * 70)
for r in range(12, 26):
    col_b = ws_dsf[f'B{r}'].value
    col_a = ws_dsf[f'A{r}'].value
    print(f"Row {r}:")
    if col_a:
        print(f"  Col A: {col_a}")
    if col_b:
        print(f"  Col B: {col_b}")
    
    # Regarder les valeurs en colonnes D, E, J
    val_d = ws_dsf[f'D{r}'].value
    val_e = ws_dsf[f'E{r}'].value
    val_j = ws_dsf[f'J{r}'].value
    
    if val_d or val_e or val_j:
        print(f"  DONNÉES EXISTANTES: D={val_d}, E={val_e}, J={val_j}")
    
    # Tester fuzzy match avec chaque label
    label_to_test = col_b or col_a
    if label_to_test:
        label_clean = str(label_to_test).strip()
        if len(label_clean) > 5:  # Ignore very short labels
            best_score = 0
            best_match = None
            for bal_label in balance_accounts.keys():
                score = SequenceMatcher(None, label_clean.lower(), bal_label.lower()).ratio()
                if score > best_score:
                    best_score = score
                    best_match = bal_label
            
            print(f"  Fuzzy match: {label_clean[:40]}")
            print(f"    → {best_match[:50]} (score: {best_score:.3f})")

print("\n" + "="*70)
print("Comptes balance avec 'IMMOBILISATION' dans le libellé:")
print("="*70)
for label, num in balance_accounts.items():
    if 'IMMOBILISATION' in label.upper():
        print(f"  {num}: {label}")
