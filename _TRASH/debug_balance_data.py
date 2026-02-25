"""Check if balance accounts have actual data in Col14, Col21, Col27"""

import openpyxl

balance_file = r'input/BALANCE AU 31 12 2024 VERSION DU 26 05 25.xlsx'
wb_bal = openpyxl.load_workbook(balance_file)
ws_bal = wb_bal.active

print("="*80)
print("BALANCE - Check data in Col14 (N), Col21 (U), Col27 (AA)")
print("="*80)

# Chercher "IMMOBILISATIONS INCORPORELLES"
print("\n[1] Find IMMOBILISATIONS INCORPORELLES:")
for row_idx in range(1, ws_bal.max_row + 1):
    acc_num = ws_bal[f'A{row_idx}'].value
    acc_label = ws_bal[f'D{row_idx}'].value
    
    if acc_label and 'IMMOBILISATIONS INCORPORELLES' in str(acc_label).upper():
        print(f"\nRow {row_idx}: {acc_num} - {acc_label}")
        
        col14 = ws_bal[f'N{row_idx}'].value
        col21 = ws_bal[f'U{row_idx}'].value
        col27 = ws_bal[f'AA{row_idx}'].value
        
        print(f"  Col14 (N): {col14}")
        print(f"  Col21 (U): {col21}")
        print(f"  Col27 (AA): {col27}")
        
        if not col14:
            print(f"  *** WARNING: Col14 is EMPTY/None!")

# Chercher "IMMOBILISATIONS CORPORELLES"  
print("\n[2] Find IMMOBILISATIONS CORPORELLES:")
for row_idx in range(1, ws_bal.max_row + 1):
    acc_num = ws_bal[f'A{row_idx}'].value
    acc_label = ws_bal[f'D{row_idx}'].value
    
    if acc_label and 'IMMOBILISATIONS CORPORELLES' in str(acc_label).upper():
        print(f"\nRow {row_idx}: {acc_num} - {acc_label}")
        
        col14 = ws_bal[f'N{row_idx}'].value
        col21 = ws_bal[f'U{row_idx}'].value
        col27 = ws_bal[f'AA{row_idx}'].value
        
        print(f"  Col14 (N): {col14}")
        print(f"  Col21 (U): {col21}")
        print(f"  Col27 (AA): {col27}")

# Check how many accounts have Col14 data
print("\n[3] Count accounts with Col14 data:")
col14_yes = 0
col14_no = 0
for row_idx in range(1, ws_bal.max_row + 1):
    acc_num = ws_bal[f'A{row_idx}'].value
    col14 = ws_bal[f'N{row_idx}'].value
    
    if not acc_num:
        continue
    
    if col14:
        col14_yes += 1
    else:
        col14_no += 1

print(f"  With Col14 data: {col14_yes}")
print(f"  Without Col14 data: {col14_no}")
