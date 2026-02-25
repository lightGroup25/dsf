"""Analyser complètement la fiche R3 pour trouver tous les tableaux"""
from pathlib import Path
from openpyxl import load_workbook

template = Path("templates/DSF Normal standard.xlsx")
wb = load_workbook(template)
ws = wb["Fiche R3"]

print("FICHE R3 - Analyse complete")
print("=" * 100)

# Afficher toutes les lignes 1-60
print("\nContenu complet (lignes 1-60):")
print("-" * 100)

for row_num in range(1, 61):
    row_data = []
    for col_num in range(1, 7):  # A-F
        cell = ws.cell(row=row_num, column=col_num)
        if cell.value not in [None, ""]:
            col_letter = chr(64 + col_num)
            val = str(cell.value).replace("\n", " ")[:50]
            row_data.append(f"{col_letter}:{val}")
    
    if row_data:
        print(f"L{row_num:2d}: {' | '.join(row_data)}")

wb.close()
