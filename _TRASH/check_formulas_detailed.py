#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.utils import get_column_letter

wb = openpyxl.load_workbook(r'templates\DSF Normal standard.xlsx', data_only=False)

# Analyser NOTE 3B complètement
ws3b = wb['NOTE 3B']

print("="*150)
print("ANALYSE COMPLETE - NOTE 3B (DEPRECIATIONS/AMORTISSEMENTS)")
print("="*150)

print("\nHEADERS NOTE 3B - TOUTES LES COLONNES (Row 8-10):")
for row_num in [8, 9, 10]:
    print(f"\nRow {row_num}:")
    for col_num in range(1, 25):  # Check 24 columns
        cell = ws3b.cell(row_num, col_num)
        val = cell.value
        if val or col_num <= 15:  # Show at least first 15
            col_letter = get_column_letter(col_num)
            val_str = str(val)[:50] if val else ""
            print(f"  Col {col_num:2d} ({col_letter}): {val_str}")

print("\n" + "="*150)
print("PREMIERE LIGNE DE DONNEES (Row 11) - Chercher les formules/percentages")
print("="*150)

for col_num in range(1, 25):
    cell = ws3b.cell(11, col_num)
    if cell.value:
        val_str = str(cell.value)[:80]
        if isinstance(cell.value, str) and cell.value.startswith('='):
            print(f"Col {col_num}: FORMULE = {val_str}")
        else:
            print(f"Col {col_num}: {cell.data_type} = {val_str}")

print("\n" + "="*150)
print("VERIFIER D'AUTRES ENDROITS POUR LES TAUX - Chercher 'TAUX', '%', 'AMORT'")
print("="*150)

# Chercher "TAUX" quelque part dans la feuille
for row_num in range(1, 50):
    for col_num in range(1, 30):
        cell = ws3b.cell(row_num, col_num)
        if cell.value:
            val_str = str(cell.value).upper()
            if 'TAUX' in val_str or '%' in val_str or 'AMORT' in val_str:
                print(f"Row {row_num}, Col {col_num}: {str(cell.value)[:80]}")

wb.close()

print("\n\nVérification aussi des autres NOTE pour voir si elles ont des structures calculées...")

wb2 = openpyxl.load_workbook(r'templates\DSF Normal standard.xlsx', data_only=False)
ws_note4 = wb2['NOTE 4 ']

print("\n" + "="*150)
print("NOTE 4 (STOCKS) - Structure")
print("="*150)

print("\nHEADERS NOTE 4 (Row 8):")
for col_num in range(1, 16):
    val = ws_note4.cell(8, col_num).value
    if val:
        print(f"  Col {col_num}: {str(val)[:70]}")

print("\nY a-t-il des colonnes pour les taux, pourcentages ou formules?")

wb2.close()
