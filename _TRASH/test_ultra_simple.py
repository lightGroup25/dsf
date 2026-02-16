"""
Ultra-simple test - write directly and verify immediately
"""
from openpyxl import load_workbook
import json

print("\n" + "="*70)
print("ULTRA-SIMPLE WRITE TEST")
print("="*70)

# Load template
print("\n1. Loading template...")
wb = load_workbook("DSF Normal standard.xlsx")
ws = wb["ENTETE"]
print(f"   ✓ Template loaded, ENTETE sheet found")

# Check what's currently there
print(f"\n2. Before writing:")
print(f"   B2 = {ws['B2'].value}")
print(f"   B3 = {ws['B3'].value}")

# Write directly
print(f"\n3. Writing data directly...")
ws['B2'].value = "TEST GULFCAM"
ws['B3'].value = "TEST SIGLE"
ws['B4'].value = "TEST ADRESSE"

print(f"   ✓ Values written to cells")

# Check in memory
print(f"\n4. After writing (in memory):")
print(f"   B2 = {ws['B2'].value}")
print(f"   B3 = {ws['B3'].value}")
print(f"   B4 = {ws['B4'].value}")

# Save
print(f"\n5. Saving file...")
wb.save("DSF_TEST_ULTRASMPLE.xlsx")
print(f"   ✓ File saved")

# Reload and verify
print(f"\n6. Reloading file to verify...")
wb2 = load_workbook("DSF_TEST_ULTRASMPLE.xlsx")
ws2 = wb2["ENTETE"]

print(f"\n7. After reloading:")
print(f"   B2 = {ws2['B2'].value}")
print(f"   B3 = {ws2['B3'].value}")
print(f"   B4 = {ws2['B4'].value}")

if ws2['B2'].value == "TEST GULFCAM":
    print(f"\n✅ SUCCESS - Data persisted!")
else:
    print(f"\n❌ FAILED - Data not saved!")

print()
