import openpyxl

dsf_file = r'templates/DSF Normal standard.xlsx'
wb = openpyxl.load_workbook(dsf_file)

# Check BILAN PAYSAGE sheet first
ws = wb['BILAN PAYSAGE']

print("Checking BILAN PAYSAGE for colored cells...\n")

for row_idx in range(1, min(50, ws.max_row + 1)):
    for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K']:
        cell = ws[f'{col}{row_idx}']
        
        has_color = False
        color_info = None
        
        # Check fill
        if cell.fill:
            try:
                if hasattr(cell.fill, 'fgColor') and cell.fill.fgColor:
                    fgc = cell.fill.fgColor
                    if hasattr(fgc, 'rgb') and fgc.rgb and fgc.rgb != '00000000':
                        has_color = True
                        color_info = fgc.rgb
                    elif hasattr(fgc, 'theme') and fgc.theme is not None:
                        has_color = True
                        color_info = f"theme:{fgc.theme}"
                    elif hasattr(fgc, 'index') and fgc.index:
                        has_color = True
                        color_info = f"index:{fgc.index}"
            except:
                pass
        
        if has_color and cell.value:
            print(f"Row {row_idx}, {col}: Color={color_info}, Value={str(cell.value)[:40]}")
