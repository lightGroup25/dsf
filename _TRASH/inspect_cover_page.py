import openpyxl

def inspect_cover_page(path):
    print(f"Loading {path}...")
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:
        print(f"Error: {e}")
        return

    sheet_name = "PAGE DE GARDE"
    if sheet_name not in wb.sheetnames:
        print(f"Sheet '{sheet_name}' not found.")
        return

    ws = wb[sheet_name]
    print(f"Inspecting '{sheet_name}'...")
    
    output_file = Path("output/cover_page.txt")
    print(f"Writing to {output_file.absolute()}")
    with open(output_file, "w", encoding="utf-8") as f:
        for row in ws.iter_rows(max_row=50):
            for cell in row:
                if cell.value:
                    val = str(cell.value).strip()
                    if len(val) > 2:
                        line = f"{cell.coordinate}: {val}\n"
                        print(line.strip())
                        f.write(line)



if __name__ == "__main__":
    from pathlib import Path
    import glob
    import os

    # Find the latest output file
    list_of_files = glob.glob('output/DSF_OUTPUT_2024*.xlsx') 
    if not list_of_files:
        print("No output file found.")
        exit(1)
    
    latest_file = max(list_of_files, key=os.path.getctime)
    print(f"Inspecting latest file: {latest_file}")
    
    wb = openpyxl.load_workbook(latest_file, data_only=True)
    if "PAGE DE GARDE" not in wb.sheetnames:
        print("PAGE DE GARDE not found!")
        exit(1)
        
    ws = wb["PAGE DE GARDE"]
    cells_to_check = ["A27", "A29", "A31", "A33", "B10", "B18", "C36"]
    
    for coord in cells_to_check:
        val = ws[coord].value
        print(f"{coord}: {val}")


