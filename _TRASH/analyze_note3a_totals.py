"""
Analyser NOTE 3A en détail pour comprendre la structure des TOTAUX
"""

import openpyxl
from pathlib import Path

TEMPLATE_FILE = Path(r"c:\Users\Emmaneul Ambadiang\Desktop\DSF\templates\DSF Normal standard.xlsx")

def main():
    wb = openpyxl.load_workbook(TEMPLATE_FILE)
    ws = wb['NOTE 3A']
    
    print("="*80)
    print("ANALYSE NOTE 3A - Structure des TOTAUX")
    print("="*80)
    
    # Chercher les lignes TOTAL
    for row_idx in range(10, 60):
        label = None
        for col in ['A', 'B']:
            val = ws[f'{col}{row_idx}'].value
            if val and isinstance(val, str):
                label = val
                break
        
        if label and any(kw in label.upper() for kw in ['TOTAL', 'SOUS-TOTAL', 'SOMME']):
            print(f"\nRow {row_idx}: {label}")
            print("-" * 60)
            
            # Vérifier les valeurs dans les colonnes de données
            for col_letter in ['D', 'E', 'F', 'G', 'H', 'I', 'J', 'K']:
                cell = ws[f'{col_letter}{row_idx}']
                val = cell.value
                
                # Vérifier le header de la colonne
                header = ""
                for h_row in range(1, 15):
                    h_val = ws[f'{col_letter}{h_row}'].value
                    if h_val:
                        header += str(h_val).upper() + " "
                
                val_display = str(val)[:30] if val else "(vide)"
                if isinstance(val, str) and val.startswith('='):
                    val_display = f"FORMULE: {val[:50]}"
                
                print(f"  {col_letter}{row_idx}: {val_display:40s} | Header: {header[:40]}")
            
            # Vérifier les lignes au-dessus (données à sommer)
            print(f"\n  Lignes au-dessus (échantillon):")
            for prev_row in range(max(row_idx-5, 15), row_idx):
                label_above = ws[f'A{prev_row}'].value or ws[f'B{prev_row}'].value
                if label_above:
                    d_val = ws[f'D{prev_row}'].value
                    e_val = ws[f'E{prev_row}'].value
                    print(f"    Row {prev_row}: {str(label_above)[:40]:40s} | D={d_val} E={e_val}")

if __name__ == "__main__":
    main()
