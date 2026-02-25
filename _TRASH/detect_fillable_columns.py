"""
Detect which columns should be filled by analyzing ROW HEADERS
Each sheet has different structure - we need to identify filling columns by header text
"""

import openpyxl

dsf_file = r'templates/DSF Normal standard.xlsx'
wb = openpyxl.load_workbook(dsf_file)

sheets_to_check = ['BILAN PAYSAGE', 'NOTE 3A', 'NOTE 3B', 'NOTE 4', 'NOTE 6', 'NOTE 7']

for sheet_name in sheets_to_check:
    if sheet_name not in wb.sheetnames:
        continue
    
    ws = wb[sheet_name]
    print(f"\n{'='*70}")
    print(f"Sheet: {sheet_name}")
    print(f"{'='*70}")
    
    # Find header row by looking for typical header keywords
    header_row = None
    for row_idx in range(1, 20):
        row_vals = []
        for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L']:
            val = ws[f'{col}{row_idx}'].value
            if val:
                row_vals.append(str(val)[:30])
        
        # Check if this looks like a header row
        row_str = ' '.join(row_vals).upper()
        if any(x in row_str for x in ['MONTANT', 'BRUT', 'ACQUISITIONS', 'CLOTURE', 'OUVERTURE', 'DESIGNAT']):
            header_row = row_idx
            print(f"\nHeader row: {row_idx}")
            
            # Show each column header
            for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L']:
                val = ws[f'{col}{row_idx}'].value
                if val:
                    header_text = str(val)[:50]
                    print(f"  Col {col}: {header_text}")
            
            break
    
    if header_row is None:
        print("Could not detect header row structure")
    else:
        # Identify which columns should be FILLED based on headers
        fillable_columns = []
        
        for col in ['D', 'E', 'F', 'G', 'H', 'I', 'J', 'K']:
            header_val = ws[f'{col}{header_row}'].value
            if not header_val:
                continue
            
            header_text = str(header_val).upper()
            
            # Identify fillable columns based on keywords
            # FILLABLE: MONTANT (OUVERTURE/BRUT), ACQUISITIONS, VARIATIONS, MOUVEMENTS, etc.
            # NOT FILLABLE: NET, AMORT, CLÔTURE (usually calculated), COMPARATIF, Exercice N-1
            
            is_fillable = False
            is_calculated = False
            
            # Definitely fillable
            if any(x in header_text for x in ['MONTANT BRUTE A L\'OU', 'ACQUISITIONS', 'VARIATIONS', 
                                              'VIREMENTS', 'REEVALUATION', 'CESSIONS']):
                is_fillable = True
            
            # Definitely NOT fillable (calculated/comparatif)
            if any(x in header_text for x in ['NET', 'AMORT', 'CLOTURE', 'N-1', 'DEPREC', 'COMPARATIF', 'EXERCICE']):
                is_calculated = True
                is_fillable = False  # Override
            
            status = "FILL" if is_fillable else ("CALCULATED" if is_calculated else "UNKNOWN")
            fillable_columns.append((col, status))
            
            print(f"\n  Decision: Col {col} = {status}")
            print(f"    Header: {header_text}")
