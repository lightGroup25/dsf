import openpyxl
import glob
import os
from pathlib import Path

def verify_cover_page():
    # Find the latest output file
    list_of_files = glob.glob('output/DSF_OUTPUT_2024*.xlsx') 
    if not list_of_files:
        print("No output file found.")
        return
    
    latest_file = max(list_of_files, key=os.path.getctime)
    print(f"Verifying file: {latest_file}")
    
    try:
        print("Loading workbook (read_only)...")
        wb = openpyxl.load_workbook(latest_file, data_only=True, read_only=True)
    except Exception as e:
        print(f"Error loading workbook: {e}")
        return

    if "PAGE DE GARDE" not in wb.sheetnames:
        print("Sheet 'PAGE DE GARDE' not found.")
        return

    ws = wb["PAGE DE GARDE"]
    
    checks = {
        "A27": "GULFCAM",
        "A29": "GULFCAM",
        "B10": "GRANDES ENTREPRISES",
        "A33": "M050900027774W",
        "A31": "AFFAIRES MARITIMES",
        "B18": "31/12/2024"
    }
    
    all_passed = True
    for coord, expected in checks.items():
        val = str(ws[coord].value)
        if expected in val:
             print(f"[PASS] {coord}: {val}")
        else:
             print(f"[FAIL] {coord}: Expected '{expected}' in '{val}'")
             all_passed = False
             
    if all_passed:
        print("\nSUCCESS: Cover page is correctly filled.")
    else:
        print("\nFAILURE: Some fields are missing or incorrect.")

if __name__ == "__main__":
    verify_cover_page()
