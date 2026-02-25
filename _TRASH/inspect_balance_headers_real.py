#!/usr/bin/env python3
"""
Inspect ACTUAL balance file headers to understand column structure
"""
import openpyxl
import os

balance_file = r"c:\Users\Emmaneul Ambadiang\Desktop\DSF\input\BALANCE AU 31 12 2024 VERSION DU 26 05 25.xlsx"

if not os.path.exists(balance_file):
    print(f"ERROR: Balance file not found: {balance_file}")
    exit(1)

print("=" * 100)
print("BALANCE FILE HEADER INSPECTION")
print("=" * 100)

try:
    wb = openpyxl.load_workbook(balance_file, data_only=True)
    print(f"\nSheets in balance file: {wb.sheetnames}")
    
    ws = wb.active
    print(f"Active sheet: {ws.title}")
    
    # Read rows 1-12 (header section)
    print("\n" + "=" * 100)
    print("COLUMN HEADERS (Rows 1-12)")
    print("=" * 100)
    
    for row_num in range(1, 13):
        row_values = []
        for col_num in range(1, 35):  # First 34 columns
            cell = ws.cell(row_num, col_num)
            val = cell.value
            if val is None:
                val = ""
            else:
                val = str(val)[:20]  # Truncate long strings
            row_values.append(f"C{col_num}: {val}")
        
        print(f"\nRow {row_num}:")
        for i in range(0, len(row_values), 4):
            print("  " + " | ".join(row_values[i:i+4]))
    
    # Now check what columns are in row 10 (final header row)
    print("\n" + "=" * 100)
    print("ROW 10 (FINAL HEADER) - ALL COLUMNS")
    print("=" * 100)
    
    for col_num in range(1, 35):
        cell = ws.cell(10, col_num)
        val = cell.value
        if val:
            print(f"Column {col_num}: {val}")
    
    # Check some data rows to understand structure
    print("\n" + "=" * 100)
    print("SAMPLE DATA ROW (Row 20)")
    print("=" * 100)
    
    for col_num in range(1, 35):
        cell = ws.cell(20, col_num)
        val = cell.value
        if val is not None:
            print(f"Column {col_num}: {val}")
    
    wb.close()
    
except Exception as e:
    print(f"ERROR: {e}")
    import traceback
    traceback.print_exc()

print("\n" + "=" * 100)
print("INTERPRETATION NEEDED")
print("=" * 100)
print("""
Questions to answer:
1. Which column contains "Soldes - Débit"? (Should be 14?)
2. Which column contains "Soldes - Crédit"? (Should be 27?)
3. Which columns contain "Mouvements"? (For acquisitions, transfers, etc.)
4. Is there an "Opening balance" or "Soldes initiaux"?
5. What are ALL the column types?
""")
