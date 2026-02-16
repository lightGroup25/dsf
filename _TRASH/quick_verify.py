from openpyxl import load_workbook

print("\n" + "="*60)
print("QUICK VERIFICATION")
print("="*60 + "\n")

try:
    wb = load_workbook('DSF_GULFCAM_RESPECTFUL.xlsx', data_only=False)
    print("✓ File loaded successfully")
    print(f"  Sheets: {wb.sheetnames}")
    
    ws = wb['ENTETE']
    print(f"\nENTETE Sheet Data:")
    print(f"  A2 = {ws['A2'].value}")
    print(f"  A3 = {ws['A3'].value}")  
    print(f"  A4 = {ws['A4'].value}")
    print(f"  A5 = {ws['A5'].value}")
    print(f"  A9 = {ws['A9'].value}")
    
    # Check formatting
    cell_a2 = ws['A2']
    print(f"\nENTETE A2 Formatting:")
    print(f"  Font: {cell_a2.font}")
    print(f"  Fill: {cell_a2.fill}")
    print(f"  Alignment: {cell_a2.alignment}")
    
except Exception as e:
    print(f"❌ Error: {e}")
    import traceback
    traceback.print_exc()
