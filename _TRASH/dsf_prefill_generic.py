"""
Pré-remplisseur DSF générique basé sur mapping JSON
Utilise un mapping JSON pour populate les cellules plutôt que du code hard-codé
"""

import json
import logging
from pathlib import Path
from datetime import date
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment
from dsf_general_info import DSF_InfosGenerales, format_currency, format_date

logger = logging.getLogger(__name__)

class DSF_Prefiller_Generic:
    """Remplit les pages DSF en utilisant un mapping JSON"""
    
    def __init__(self, template_path: str, mapping_path: str = "dsf_prefill_mapping.json"):
        """
        Initialise le pré-remplisseur
        
        Args:
            template_path: Chemin du fichier DSF template
            mapping_path: Chemin du fichier de mapping JSON
        """
        self.template_path = Path(template_path)
        self.mapping_path = Path(mapping_path)
        self.wb = None
        self.mapping = None
        self.loaded = False
        
    def load_mapping(self):
        """Charge le mapping JSON"""
        if not self.mapping_path.exists():
            raise FileNotFoundError(f"Mapping non trouvé: {self.mapping_path}")
        
        with open(self.mapping_path, 'r', encoding='utf-8') as f:
            self.mapping = json.load(f)
        
        logger.info(f"Mapping chargé: {self.mapping_path}")
        return self.mapping
    
    def load_template(self):
        """Charge le template Excel"""
        if not self.template_path.exists():
            raise FileNotFoundError(f"Template non trouvé: {self.template_path}")
        
        try:
            self.wb = load_workbook(str(self.template_path))
            logger.info(f"Template chargé: {self.template_path}")
        except Exception as e:
            logger.warning(f"Impossible de charger template: {e}. Création nouveau workbook.")
            self.wb = Workbook()
        
        self.loaded = True
    
    def get_or_create_sheet(self, sheet_name: str):
        """Obtient ou crée une feuille"""
        if sheet_name not in self.wb.sheetnames:
            ws = self.wb.create_sheet(sheet_name)
        else:
            ws = self.wb[sheet_name]
        return ws
    
    def _get_attribute_value(self, infos: DSF_InfosGenerales, attribute_path: str):
        """
        Récupère la valeur d'un attribut dans l'objet infos
        Supporte les chemins imbriqués: "dirigeant.nom"
        """
        parts = attribute_path.split('.')
        value = infos
        for part in parts:
            if hasattr(value, part):
                value = getattr(value, part)
            else:
                return None
        return value
    
    def _format_value(self, value, format_type: str = None):
        """Formate une valeur selon son type"""
        if value is None:
            return ""
        
        if format_type == "date":
            if isinstance(value, date):
                return format_date(value)
            return str(value)
        elif format_type == "currency":
            if isinstance(value, (int, float)):
                return format_currency(value)
            return str(value)
        elif format_type == "percentage":
            if isinstance(value, (int, float)):
                return f"{value:.2f}%"
            return str(value)
        
        return str(value) if value is not None else ""
    
    def _write_cell(self, ws, cell_ref: str, value, bold: bool = False, currency: bool = False):
        """Écrit une cellule avec formatage, gère les cellules fusionnées"""
        try:
            cell = ws[cell_ref]
            
            # Si c'est une cellule fusionnée, trouver la cellule de départ
            if isinstance(cell, type(None)) or str(type(cell).__name__) == 'MergedCell':
                for merged_range in ws.merged_cells.ranges:
                    if cell_ref in merged_range:
                        cell_ref = merged_range.start_cell.coordinate
                        cell = ws[cell_ref]
                        break
            
            # Formater la valeur
            if currency and isinstance(value, (int, float)):
                cell.value = format_currency(value)
                cell.number_format = '#,##0'
            else:
                cell.value = value
            
            if bold:
                cell.font = Font(bold=True)
            
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            
        except Exception as e:
            logger.debug(f"Cellule {cell_ref} skipped: {e}")
    
    def fill_sheet_from_mapping(self, sheet_name: str, infos: DSF_InfosGenerales):
        """
        Remplit une feuille en utilisant le mapping
        """
        if sheet_name not in self.mapping:
            logger.warning(f"Aucun mapping pour {sheet_name}")
            return
        
        sheet_mapping = self.mapping[sheet_name]
        ws = self.get_or_create_sheet(sheet_name)
        
        # Écrire le titre si défini
        if "title" in sheet_mapping:
            self._write_cell(ws, "A1", sheet_mapping["title"], bold=True)
        
        # Remplir les cellules simples
        if "cells" in sheet_mapping:
            for cell_ref, cell_config in sheet_mapping["cells"].items():
                if isinstance(cell_config, dict) and "attribute" in cell_config:
                    value = self._get_attribute_value(infos, cell_config["attribute"])
                    
                    # Appliquer le formatage
                    format_type = cell_config.get("format")
                    if format_type == "date" and isinstance(value, date):
                        value = format_date(value)
                    
                    bold = cell_config.get("bold", False)
                    self._write_cell(ws, cell_ref, value, bold=bold)
        
        # Remplir les tableaux
        if "tables" in sheet_mapping:
            current_row_offset = 1  # Démarrer après le titre
            
            for table_name, table_config in sheet_mapping["tables"].items():
                # Obtenir les données du tableau
                table_data = self._get_attribute_value(infos, table_name)
                
                if not table_data:
                    logger.warning(f"Pas de données pour table {table_name}")
                    continue
                
                # Écrire le titre du tableau
                if "title_row" in table_config:
                    title_row = int(table_config["title_row"].replace("A", ""))
                    self._write_cell(ws, f'A{title_row}', table_config.get("title", ""), bold=True)
                
                # Écrire les en-têtes
                if "header_row" in table_config:
                    header_row = int(table_config["header_row"].replace("A", ""))
                    for header in table_config["headers"]:
                        col = header["column"]
                        label = header["label"]
                        self._write_cell(ws, f'{col}{header_row}', label, bold=True)
                
                # Écrire les données
                if "data_start_row" in table_config:
                    data_row = int(table_config["data_start_row"].replace("A", ""))
                    fields = table_config["fields"]
                    
                    for idx, item in enumerate(table_data, 1):
                        for col, field_name in fields.items():
                            if field_name == "index":
                                value = str(idx)
                            elif field_name == "nom_prenoms":
                                # Combiner nom et prénoms
                                nom = getattr(item, 'nom', '')
                                prenoms = getattr(item, 'prenoms', '')
                                value = f"{nom} {prenoms}".strip()
                            else:
                                value = getattr(item, field_name, '')
                            
                            # Appliquer le formatage
                            if field_name in ['chiffre_affaire_ht', 'montant_total', 'montant']:
                                self._write_cell(ws, f'{col}{data_row}', value, currency=True)
                            elif field_name in ['pourcentage_ca', 'pourcentage']:
                                if isinstance(value, (int, float)):
                                    self._write_cell(ws, f'{col}{data_row}', f"{value:.2f}%")
                                else:
                                    self._write_cell(ws, f'{col}{data_row}', value)
                            else:
                                self._write_cell(ws, f'{col}{data_row}', value)
                        
                        data_row += 1
                    
                    # Écrire les totaux si configurés
                    if "totals_row" in table_config:
                        totals_config = table_config["totals_row"]
                        label_col = totals_config["label_column"]
                        label = totals_config["label"]
                        
                        self._write_cell(ws, f'{label_col}{data_row}', label, bold=True)
                        
                        # Calculer les sommes
                        for sum_col in totals_config.get("sum_columns", []):
                            if sum_col == "E":  # Colonne montant/CA
                                total = sum(getattr(item, 'montant_total', 0) or getattr(item, 'chiffre_affaire_ht', 0) 
                                          for item in table_data)
                                if "%" not in label:
                                    self._write_cell(ws, f'{sum_col}{data_row}', total, currency=True)
                            elif sum_col in ["D", "E"] and "%" in str(table_config["headers"]):
                                # Somme des pourcentages
                                total = sum(getattr(item, 'pourcentage_ca', 0) for item in table_data)
                                self._write_cell(ws, f'{sum_col}{data_row}', f"{total:.2f}%", bold=True)
        
        logger.info(f"{sheet_name} remplie")
    
    def fill_all(self, infos: DSF_InfosGenerales):
        """Remplit toutes les pages selon le mapping"""
        if not self.loaded:
            self.load_template()
        
        if not self.mapping:
            self.load_mapping()
        
        pages_remplies = []
        for sheet_name in self.mapping.keys():
            self.fill_sheet_from_mapping(sheet_name, infos)
            pages_remplies.append(sheet_name)
        
        logger.info(f"✓ Toutes les pages remplies avec succès: {', '.join(pages_remplies)}")
    
    def save(self, output_path: str = None):
        """Sauvegarde le fichier DSF"""
        if not self.loaded:
            raise RuntimeError("Aucun fichier DSF chargé")
        
        if output_path is None:
            output_path = str(self.template_path)
        
        self.wb.save(output_path)
        logger.info(f"✓ Fichier sauvegardé: {output_path}")


def prefill_dsf_generic(
    template_path: str, 
    infos: DSF_InfosGenerales, 
    output_file: str = None,
    mapping_path: str = "dsf_prefill_mapping.json"
):
    """
    Remplit un fichier DSF en utilisant le mapping JSON
    
    Args:
        template_path: Chemin du fichier DSF
        infos: Objet DSF_InfosGenerales
        output_file: Fichier de sortie (optionnel)
        mapping_path: Chemin du mapping JSON
    
    Returns:
        Chemin du fichier généré
    """
    prefiller = DSF_Prefiller_Generic(template_path, mapping_path)
    prefiller.load_mapping()
    prefiller.load_template()
    prefiller.fill_all(infos)
    
    output = output_file or template_path
    prefiller.save(output)
    
    return output


if __name__ == "__main__":
    from dsf_general_info import GULFCAM_SAS
    import logging
    
    logging.basicConfig(level=logging.INFO)
    
    # Test avec GULFCAM
    output = prefill_dsf_generic(
        "DSF Normal standard.xlsx",
        GULFCAM_SAS,
        "DSF_GULFCAM_MAPPING.xlsx"
    )
    
    print(f"✓ DSF généré avec mapping: {output}")
