"""See the actual data structure - first few rows with accounts"""

import openpyxl

balance_file = r'input/BALANCE AU 31 12 2024 VERSION DU 26 05 25.xlsx'
wb_bal = openpyxl.load_workbook(balance_file)
ws_bal = wb_bal.active

print("="*100)
print("BALANCE FILE - FIRST 20 ROWS (find where headers are)")
print("="*100)

for row_idx in range(1, 21):
    print(f"\nRow {row_idx}:")
    has_data = False
    for col_idx in range(1, 25):
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        val = ws_bal[f'{col_letter}{row_idx}'].value
        if val:
            print(f"  {col_letter} ({col_idx:2d}): {val}")
            has_data = True
    
    if not has_data:
        print("  (empty row)")

print("\n" + "="*100)
print("LET'S LOOK AT A REAL ACCOUNT ROW (e.g., row 10)")
print("="*100)

row = 10
print(f"\nRow {row} (full row, all columns 1-30):")
for col_idx in range(1, 31):
    col_letter = openpyxl.utils.get_column_letter(col_idx)
    val = ws_bal[f'{col_letter}{row}'].value
    if val is not None:
        print(f"  {col_letter} ({col_idx:2d}): {val}")
