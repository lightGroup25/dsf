#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Analyser UN COMPTE EN DETAIL pour comprendre la structure
"""
import openpyxl

balance_file = r"c:\Users\Emmaneul Ambadiang\Desktop\DSF\input\BALANCE AU 31 12 2024 VERSION DU 26 05 25.xlsx"

wb = openpyxl.load_workbook(balance_file, data_only=True)
ws = wb.active

print("="*150)
print("ANALYSE DETAILLEE D'UN COMPTE - Account 10 (CAPITAL)")
print("="*150)

# Account 10 est en row 20
account_row = 20

print("\nRow 20 - Account 10 (CAPITAL):")
print("-" * 150)

for col_num in range(1, 30):
    cell = ws.cell(account_row, col_num)
    val = cell.value
    
    if val is not None:
        # Afficher la valeur ET chercher le header
        r8 = ws.cell(8, col_num).value
        r10 = ws.cell(10, col_num).value
        
        header = ""
        if r8:
            header = str(r8)[:30]
        if r10:
            header += f" ({str(r10)[:10]})"
        
        print(f"  Col {col_num:2d} | Value: {str(val):20s} | Header: {header}")

print("\n" + "="*150)
print("INTERPRETATION DES DONNEES")
print("="*150)

# Col 14, 21, 27 contiennent les seules valeurs non-zéro
# Essayons de comprendre ce qu'elles signifient

col_14 = ws.cell(account_row, 14).value
col_21 = ws.cell(account_row, 21).value
col_27 = ws.cell(account_row, 27).value

print(f"""
Account 10 (CAPITAL):
  Col 14: {col_14:>15} (Mouvement 31/12/23 Crédit ou Solde 31/12/23 Débit?)
  Col 21: {col_21:>15} (Mouvement intermédiaire Crédit?)
  Col 27: {col_27:>15} (Solde final Crédit?)

Logique possible:
  Si Col 14 est "Solde Débit 31/12/23 Mouvement"
  Et Col 27 est "Solde Crédit final"
  
  Alors: Col 27 = Col 14 + Col 21 (mouvements)?
  Vérification: {col_14} + {col_21} = {col_14 + col_21 if col_21 else 'impossible'}
  Resultat: {col_14 + col_21 == col_27 if col_21 else 'N/A'}
""")

# Regarder un autre compte
print("\n" + "="*150)
print("VERIFIER AVEC ACCOUNT 11 (RESERVES)")
print("="*150)

account_row_2 = 25  # RESERVES est probablement en row 25

account_11_num = ws.cell(account_row_2, 1).value
print(f"Row {account_row_2} contains account: {account_11_num}")

col_14_2 = ws.cell(account_row_2, 14).value
col_21_2 = ws.cell(account_row_2, 21).value
col_27_2 = ws.cell(account_row_2, 27).value

print(f"""
Account {account_11_num}:
  Col 14: {col_14_2}
  Col 21: {col_21_2}
  Col 27: {col_27_2}
""")

# Aussi vérifier les colonnes 10, 11, 12, 13 qui correspondraient à mouvement 31/12
print("\n" + "="*150)
print("VERIFIER TOUTES LES COLONNES POUR ACCOUNT 10")
print("="*150)

account_row = 20
print(f"Row {account_row} - ALL columns 1-27:")

for col_num in range(1, 28):
    val = ws.cell(account_row, col_num).value
    print(f"  Col {col_num:2d} ({chr(64+col_num):s}): {str(val) if val is not None else '[Empty]':>20s}")

wb.close()
