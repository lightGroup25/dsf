
import openpyxl
import os
import sys

def verify_note13(filename):
    print(f"Verifying {filename}...")
    try:
        wb = openpyxl.load_workbook(filename, data_only=True)
    except Exception as e:
        print(f"Error: {e}")
        return

    sheet_name = None
    for name in wb.sheetnames:
        if "NOTE 13" in name.upper() or "NOTE13" in name.upper():
            sheet_name = name
            break
    
    if not sheet_name:
        print("FAIL: Sheet NOTE 13 not found")
        return

    ws = wb[sheet_name]
    print(f"Sheet: {sheet_name}")

    checks = [
        ("A3", "GULFCAM SAS"),
        ("A4", "M050900027774W"),
        ("A11", "SOFIMAR SICAV"),
        ("B11", "CAMEROUNAISE"),
        ("C11", "ORDINAIRES"),
        ("D11", 276954),
        ("A20", "HESNAULT"),
        ("A21", "Apporteurs"),
        ("A26", "fusion-absorption"),
    ]

    all_passed = True
    for coord, expected in checks:
        val = ws[coord].value
        # Check based on type
        if isinstance(expected, str):
            passed = val and expected.upper() in str(val).upper()
        else:
            # Numeric loose equality
            try:
                passed = abs(float(val) - expected) < 1.0
            except:
                passed = False
                
        status = "PASS" if passed else "FAIL"
        print(f"[{status}] {coord}: {val} (Expected: {expected})")
        if not passed:
            all_passed = False

    if all_passed:
        print("SUCCESS: NOTE 13 is correctly filled.")
    else:
        print("FAILURE: Some checks failed.")
        sys.exit(1)

if __name__ == "__main__":
    filename = r"c:\Users\Emmaneul Ambadiang\Desktop\DSF\output\DSF_OUTPUT_2024.xlsx"
    verify_note13(filename)
