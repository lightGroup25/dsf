"""
ANALYSE APPROFONDIE - Pourquoi NOTE 3A/3B N'ONT PAS DE FORMULES?
==================================================================
"""

import openpyxl
from pathlib import Path
from openpyxl.utils import get_column_letter

def analyze_note_3_structure():
    """Analyser la structure des NOTE 3A et 3B"""
    
    file_path = Path(r"c:\Users\Emmaneul Ambadiang\Desktop\DSF\DSF_FINAL_V7_FIXED.xlsx")
    
    if not file_path.exists():
        print(f"❌ Fichier introuvable: {file_path}")
        return
    
    wb = openpyxl.load_workbook(file_path)
    
    for sheet_name in ['NOTE 3A', 'NOTE 3B']:
        if sheet_name not in wb.sheetnames:
            print(f"⚠️ {sheet_name} not in workbook")
            continue
        
        ws = wb[sheet_name]
        
        print(f"\n{'='*100}")
        print(f"ANALYSE: {sheet_name}")
        print('='*100)
        
        # Trouver les lignes TOTAL ou SOUS-TOTAL
        print(f"\n1️⃣ RECHERCHE DES LIGNES TOTAL:")
        print("-" * 100)
        
        total_rows = []
        for row_idx in range(1, min(100, ws.max_row + 1)):
            cell_a = ws[f'A{row_idx}']
            if cell_a.value:
                text = str(cell_a.value).upper()
                if any(kw in text for kw in ['TOTAL', 'SOUS', 'SOMME', 'GENERAL', 'GLOBAL']):
                    total_rows.append((row_idx, cell_a.value))
                    print(f"   ✅ Row {row_idx}: {cell_a.value}")
        
        if not total_rows:
            print(f"   ⚠️ Aucune ligne TOTAL détectée!")
        
        # Analyser les données autour des lignes TOTAL
        print(f"\n2️⃣ ANALYSE DES DONNÉES:")
        print("-" * 100)
        
        for total_row, label in total_rows:
            print(f"\n   TOTAL Row {total_row}: {label}")
            
            # Trouver les ranges de données
            start_search = max(1, total_row - 15)
            end_search = min(total_row - 1, ws.max_row)
            
            data_rows = []
            for row_idx in range(start_search, end_search + 1):
                cell_a = ws[f'A{row_idx}']
                has_data = False
                
                # Checker si cette ligne a des données
                for col_idx in range(2, 15):
                    col_letter = get_column_letter(col_idx)
                    cell = ws[f'{col_letter}{row_idx}']
                    if cell.value is not None:
                        has_data = True
                        break
                
                if has_data or (cell_a.value and str(cell_a.value).strip()):
                    data_rows.append(row_idx)
                    
                    # Afficher les 3 premières et dernières lignes
                    if len(data_rows) <= 3 or row_idx >= end_search - 2:
                        label_text = cell_a.value if cell_a.value else "[VIDE]"
                        values = []
                        for col_idx in range(2, 6):  # Colonnes B-E
                            col_letter = get_column_letter(col_idx)
                            cell = ws[f'{col_letter}{row_idx}']
                            val = cell.value if cell.value is not None else "-"
                            values.append(f"{col_letter}:{val}")
                        
                        print(f"      Row {row_idx}: {label_text:<40} | {', '.join(values)}")
            
            if not data_rows:
                print(f"      ⚠️ AUCUNE DONNÉE TROUVÉE (!) dans les lignes précédentes!")
            else:
                print(f"      ✅ {len(data_rows)} lignes avec données")
                
                # Checker la ligne TOTAL elle-même
                print(f"\n      CONTENU de la ligne TOTAL {total_row}:")
                has_total_data = False
                for col_idx in range(2, 15):
                    col_letter = get_column_letter(col_idx)
                    cell = ws[f'{col_letter}{total_row}']
                    if cell.value is not None:
                        has_total_data = True
                        print(f"        {col_letter}: {cell.value} (type: {type(cell.value).__name__})")
                
                if not has_total_data:
                    print(f"        ⚠️ LA LIGNE TOTAL EST VIDE! (pas de données à additionner)")
        
        # Vérifier les formules existantes
        print(f"\n3️⃣ FORMULES EXISTANTES:")
        print("-" * 100)
        
        formula_count = 0
        for row_idx in range(1, ws.max_row + 1):
            for col_idx in range(1, 15):
                col_letter = get_column_letter(col_idx)
                cell = ws[f'{col_letter}{row_idx}']
                if cell.data_type == 'f':  # Formula
                    formula_count += 1
                    print(f"   {col_letter}{row_idx}: {cell.value}")
        
        if formula_count == 0:
            print(f"   ⚠️ Aucune formule trouvée dans {sheet_name}")
        else:
            print(f"   ✅ {formula_count} formules trouvées")

def explain_empty_data_issue():
    """Expliquer l'issue des données vides"""
    
    print(f"\n\n{'='*100}")
    print("EXPLICATION: POURQUOI NOTE 3A/3B N'ONT PAS DE FORMULES?")
    print('='*100)
    
    explanation = """
📊 PROBLÈME IDENTIFIÉ:

1. NOTE 3A et NOTE 3B ont des lignes TOTAL/SOUS-TOTAL (structurellement)
2. MAIS ces lignes n'ont PAS DE DONNÉES À SOMMER
3. Donc le mapper ne crée pas de formules (car il n'y a rien à additionner)

🔍 CAS 1: DONNÉES PRÉSENTES
   Template:
   Row 13: Item 1          | 100
   Row 14: Item 2          | 200
   Row 15: TOTAL          | [VIDE]
   
   Mapper génère: =SUM(D13:D14) ✅
   
🔍 CAS 2: DONNÉES ABSENTES (cas actuel)
   Template:  
   Row 13: Item 1          | [VIDE]
   Row 14: Item 2          | [VIDE]
   Row 15: TOTAL          | [VIDE]
   
   Mapper ne génère rien ⚠️ (puisqu'il n'y a rien à sommer)

💡 SOLUTION POSSIBLE:

Option 1: Attendre que les données arrivent
   - Si les données de NOTE 3A/3B sont importées ultérieurement
   - Les formules seront créées à ce moment

Option 2: Créer les formules "à l'avance" (même sans données)
   - Modifier le mapper pour générer des formules
   - Même si les lignes de données sont vides
   - Les formules se calculeront automatiquement quand les données arriveront

Option 3: Remplir avec des zéros par défaut
   - Les notes ne contenant que des "0" génèreraient des formules = 0
   - Plus logique pour un bilan financier

📌 DÉCISION:
   Actuellement: Option 1 (attendre les données)
   Recommandé: Option 2 ou 3 (créer les formules à l'avance)
    """
    
    print(explanation)

if __name__ == "__main__":
    analyze_note_3_structure()
    explain_empty_data_issue()
