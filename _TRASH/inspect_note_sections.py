"""
Inspect all NOTE sections in DSF template to understand their structure
and identify which rows should be filled
"""

import openpyxl

dsf_file = r'templates/DSF Normal standard.xlsx'
wb = openpyxl.load_workbook(dsf_file)
ws = wb.active

print("="*70)
print("DSF TEMPLATE - ALL NOTE SECTIONS")
print("="*70)

notes_data = {}
current_note = None

for row_idx in range(1, min(300, ws.max_row + 1)):
    col_a = ws[f'A{row_idx}'].value
    col_b = ws[f'B{row_idx}'].value
    
    if col_a is None:
        continue
    
    col_a_str = str(col_a).strip()
    
    # Check if this is a NOTE section
    if col_a_str.startswith('NOTE'):
        current_note = col_a_str
        notes_data[current_note] = []
        print(f"\n{current_note}")
        print("-" * 70)
    
    # If we're in a NOTE section, collect data rows
    if current_note and col_a_str and not col_a_str.startswith('NOTE'):
        # Look for rows with account structure
        col_d = ws[f'D{row_idx}'].value
        col_e = ws[f'E{row_idx}'].value
        col_j = ws[f'J{row_idx}'].value
        
        if col_b or col_d or col_e or col_j:
            print(f"  Row {row_idx}: A={col_a_str}, B={col_b}, "
                  f"D={col_d}, E={col_e}, J={col_j}")
            
            notes_data[current_note].append({
                'row': row_idx,
                'col_a': col_a_str,
                'col_b': col_b,
                'col_d': col_d,
                'col_e': col_e,
                'col_j': col_j
            })

print(f"\n\nSUMMARY OF NOTES:")
print("="*70)
for note_name, rows in notes_data.items():
    print(f"{note_name}: {len(rows)} rows")
    
total_note_rows = sum(len(rows) for rows in notes_data.values())
print(f"\nTotal Note rows to potentially fill: {total_note_rows}")
