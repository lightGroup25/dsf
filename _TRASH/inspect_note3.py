
import openpyxl

def inspect_sheet_range(wb, sheet_name, start_row, end_row, cols=10):
    if sheet_name not in wb.sheetnames:
        print(f"Sheet '{sheet_name}' not found")
        # fuzzy match
        for s in wb.sheetnames:
            if sheet_name.replace(" ", "") in s.replace(" ", ""):
                print(f" -> Found similar: '{s}'")
                sheet_name = s
                break
        else:
            return

    ws = wb[sheet_name]
    print(f"\n--- Inspecting '{sheet_name}' Rows {start_row}-{end_row} ---")
    for r in range(start_row, end_row + 1):
        row_data = []
        for c in range(1, cols + 1):
            cell = ws.cell(row=r, column=c)
            var = cell.value
            val = str(var).strip() if var is not None else ""
            if len(val) > 15: val = val[:12] + "..."
            row_data.append(f"{val}")
        print(f"Row {r:2}: " + " | ".join(row_data))

filename = "temp_template.xlsx"
try:
    print(f"Loading {filename}...")
    wb = openpyxl.load_workbook(filename, data_only=True)
    inspect_sheet_range(wb, "NOTE 3A", 1, 30)
except Exception as e:
    print(f"Error: {e}")
