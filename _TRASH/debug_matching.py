"""Debug: Voir pourquoi BILAN PAYSAGE et autres grandes sheets n'ont pas de match"""

import openpyxl
from difflib import SequenceMatcher

# Charger balance
balance_file = r'input/BALANCE AU 31 12 2024 VERSION DU 26 05 25.xlsx'
wb_bal = openpyxl.load_workbook(balance_file)
ws_bal = wb_bal.active

balance_accounts = {}
for r in range(1, ws_bal.max_row + 1):
    acc_num = ws_bal[f'A{r}'].value
    acc_label = ws_bal[f'D{r}'].value
    if acc_num and acc_label:
        try:
            int(str(acc_num).split('.')[0])
            balance_accounts[str(acc_label).strip()] = str(acc_num).strip()
        except:
            pass

print("Sample balance accounts:")
for i, (label, num) in enumerate(list(balance_accounts.items())[:20]):
    print(f"  {num}: {label}")

# Charger DSF BILAN PAYSAGE
dsf_file = r'templates/DSF Normal standard.xlsx'
wb_dsf = openpyxl.load_workbook(dsf_file)
ws_dsf = wb_dsf['BILAN PAYSAGE']

print("\n" + "="*70)
print("BILAN PAYSAGE - First data rows (B column)")
print("="*70 + "\n")

# Afficher les libellés des lignes dans BILAN PAYSAGE
for r in range(12, 40):
    col_b = ws_dsf[f'B{r}'].value
    if col_b:
        label = str(col_b).strip()
        print(f"Row {r}: {label[:60]}")
        
        # Essayer fuzzy match avec balance
        best_score = 0
        best_match = None
        for bal_label in balance_accounts.keys():
            score = SequenceMatcher(None, label.lower(), bal_label.lower()).ratio()
            if score > best_score:
                best_score = score
                best_match = bal_label
        
        if best_score > 0.3:
            print(f"  MATCH FOUND: {best_match} (score: {best_score:.2f})")
        else:
            print(f"  NO MATCH (best: {best_match[:40]} @ {best_score:.2f})")
