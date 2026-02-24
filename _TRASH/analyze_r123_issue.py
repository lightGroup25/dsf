"""Analyse le problème de remplissage des fiches R1, R2, R3"""
from pathlib import Path
from openpyxl import load_workbook

# Charger le dernier DSF généré
output_dir = Path("output")
dsf_files = [f for f in output_dir.glob("*.xlsx") if not f.name.startswith("~$")]

if not dsf_files:
    print("❌ Aucun fichier DSF trouvé dans output/")
    exit(1)

latest_dsf = max(dsf_files, key=lambda p: p.stat().st_mtime)
print(f"📄 Analyse de: {latest_dsf.name}\n")

wb = load_workbook(latest_dsf, data_only=True)

# Analyser chaque fiche R1, R2, R3
for sheet_name in ["Fiche R1", "Fiche R2", "Fiche R3"]:
    if sheet_name not in wb.sheetnames:
        print(f"⚠️  Feuille '{sheet_name}' non trouvée")
        continue
    
    ws = wb[sheet_name]
    print(f"\n{'='*60}")
    print(f"📋 FICHE: {sheet_name}")
    print(f"{'='*60}")
    
    # Afficher les cellules remplies (non vides et non formules)
    filled_cells = []
    for row in ws.iter_rows(min_row=1, max_row=50, min_col=1, max_col=15):
        for cell in row:
            if cell.value is not None and cell.value != "":
                # Skip headers and structure
                if isinstance(cell.value, str) and (
                    ":" in str(cell.value) or 
                    len(str(cell.value)) < 3 or
                    str(cell.value).strip() in ["", "-"]
                ):
                    continue
                
                filled_cells.append({
                    "cell": cell.coordinate,
                    "value": cell.value,
                    "row": cell.row,
                    "col": cell.column
                })
    
    if filled_cells:
        print(f"\n✅ Cellules remplies: {len(filled_cells)}")
        print("\n🔍 Échantillon des données (10 premières):")
        for item in filled_cells[:10]:
            print(f"   {item['cell']:6} = {item['value']}")
    else:
        print("\n✅ Aucune cellule remplie (bon si c'est intentionnel)")
    
    # Analyser les lignes 5-20 (zone typique de saisie)
    print(f"\n📊 Analyse zone de saisie (lignes 5-20, colonnes A-F):")
    for row_num in range(5, 21):
        row_data = []
        for col_num in range(1, 7):
            cell = ws.cell(row=row_num, column=col_num)
            if cell.value not in [None, "", "-"]:
                row_data.append(f"Col{chr(64+col_num)}={cell.value}")
        
        if row_data:
            print(f"   Ligne {row_num:2d}: {' | '.join(row_data[:3])}")  # Max 3 colonnes

wb.close()

print(f"\n{'='*60}")
print("✅ Analyse terminée")
print(f"{'='*60}\n")
