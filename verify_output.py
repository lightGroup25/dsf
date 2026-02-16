from openpyxl import load_workbook
import glob
import os

# Find latest output
list_of_files = glob.glob('output/DSF_OUTPUT_2024_*.xlsx') 
latest_file = max(list_of_files, key=os.path.getctime)
print(f"Checking {latest_file}...")

wb = load_workbook(latest_file, data_only=True)

checks = [
    ("NOTE 6 ", "D10", "Stocks"),
    ("NOTE 7 ", "D10", "Clients"), # Virtual Field was D10?
    ("NOTE 16A ", "E10", "Emprunts"),
    ("NOTE 27A", "C10", "CA"),
    ("BILAN PAYSAGE", "K12", "Capital"),
    ("BILAN PAYSAGE", "D12", "Immob incorp"),
    ("COMPTE DE RESULTAT", "TA11", "Ventes? No col D"),
    ("COMPTE DE RESULTAT", "D11", "Ventes"),
]

for sheet, cell, label in checks:
    if sheet in wb.sheetnames:
        val = wb[sheet][cell].value
        print(f"{label} ({sheet}:{cell}) = {val}")
    else:
        print(f"Sheet {sheet} MISSING")
