#!/usr/bin/env python3
"""
Examine DSF column headers to understand what data type each needs
Focus on NOTE 3A, 3B, 4, etc to see what columns ask for
"""
import openpyxl
import os

dsf_files = [
    r"c:\Users\Emmaneul Ambadiang\Desktop\DSF\templates\DSF Normal standard.xlsx"
]

dsf_file = None
for f in dsf_files:
    if os.path.exists(f):
        dsf_file = f
        break

if not dsf_file:
    print("ERROR: DSF file not found!")
    print("Searched:")
    for f in dsf_files:
        print(f"  {f}")
    exit(1)

print(f"Using DSF file: {dsf_file}")
print("=" * 150)

wb = openpyxl.load_workbook(dsf_file, data_only=False)

# Check NOTE 3A (inventory fixed assets)
if "NOTE 3A" in wb.sheetnames:
    ws = wb["NOTE 3A"]
    
    print("\nNOTE 3 A - COLUMN HEADERS (Rows 8-10)")
    print("-" * 150)
    
    # Get row 8 headers
    print("\nRow 8 (Main headers):")
    for col in range(1, 15):
        val = ws.cell(8, col).value
        if val:
            print(f"  Col {col}: {val}")
    
    print("\nRow 9 (Sub-headers):")
    for col in range(1, 15):
        val = ws.cell(9, col).value
        if val:
            print(f"  Col {col}: {val}")
    
    print("\nRow 10 (Detail headers):")
    for col in range(1, 15):
        val = ws.cell(10, col).value
        if val:
            print(f"  Col {col}: {val}")
    
    # Show visual grid
    print("\nVisual header grid:")
    print("-" * 150)
    for col in range(1, 15):
        r8 = str(ws.cell(8, col).value or "")[:20]
        r9 = str(ws.cell(9, col).value or "")[:20]
        r10 = str(ws.cell(10, col).value or "")[:20]
        
        if r8 or r9 or r10:
            print(f"Col {col:2d} | R8: {r8:20s} | R9: {r9:20s} | R10: {r10:20s}")
    
    # Check a data row
    print("\n\nFirst data row (Row 11):")
    print("-" * 150)
    for col in range(1, 15):
        val = ws.cell(11, col).value
        if val:
            print(f"  Col {col}: {val}")

else:
    print(f"ERROR: NOTE 3 A not found in sheet names: {wb.sheetnames}")

wb.close()

print("\n" + "=" * 150)
print("WHAT WE NEED TO KNOW:")
print("=" * 150)
print("""
1. Does col 4 in DSF ("MONTANT BRUTE A L'OUVERTURE") map to balance opening balance?
   → If yes, which balance columns contain opening balance?

2. Does col 10 in DSF ("MONTANT BRUT A LA CLOTURE") map to balance closing balance?
   → YES - should be Balance col 14 (debit) or 27 (credit)

3. Are the middle columns (5-9) asking for MOVEMENTS or other data?
   → If yes, which balance columns contain movements?

This would explain WHY our semantic matching is failing:
- We're taking Balance data and putting it in WRONG DSF columns
- We need column-specific mapping, not generic row matching
""")
