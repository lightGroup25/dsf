import openpyxl
import glob
import os

def inspect_r2():
    list_of_files = glob.glob('output/DSF_OUTPUT_2024*.xlsx') 
    if not list_of_files:
        print("No output file found.")
        return
    
    latest_file = max(list_of_files, key=os.path.getctime)
    print(f"Inspecting R2 in: {latest_file}")
    
    try:
        wb = openpyxl.load_workbook(latest_file, data_only=True, read_only=True)
    except Exception as e:
        print(f"Error: {e}")
        return

    if "Fiche R2" not in wb.sheetnames:
        print("Sheet 'Fiche R2' not found.")
        return

    ws = wb["Fiche R2"]
    
    print("--- Scanning rows 1-40 ---")
    for row in ws.iter_rows(min_row=1, max_row=40):
        for cell in row:
            if cell.value:
                val = str(cell.value).strip()
                if len(val) > 2:
                    print(f"{cell.coordinate}: {val}")

if __name__ == "__main__":
    inspect_r2()
