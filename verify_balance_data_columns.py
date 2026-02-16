#!/usr/bin/env python3
"""
Check what's in the actual data columns to verify
which columns contain the Soldes cumulés (final balances)
"""
import openpyxl

balance_file = r"c:\Users\Emmaneul Ambadiang\Desktop\DSF\input\BALANCE AU 31 12 2024 VERSION DU 26 05 25.xlsx"

wb = openpyxl.load_workbook(balance_file, data_only=True)
ws = wb.active

# Check multiple data rows to understand which columns have the balances
print("=" * 120)
print("DATA COLUMN ANALYSIS - Multiple Rows")
print("=" * 120)

for row_num in [11, 15, 20, 25, 30]:
    if row_num > ws.max_row:
        continue
    
    print(f"\nRow {row_num}:")
    print("-" * 120)
    
    acct_num = ws.cell(row_num, 1).value
    acct_label = ws.cell(row_num, 6).value
    
    print(f"  Account: {acct_num} - {acct_label}")
    
    # Show all data columns
    print(f"  Col 10 (Mvt 31/12 D): {ws.cell(row_num, 10).value}")
    print(f"  Col 13 (Mvt 31/12 C): {ws.cell(row_num, 13).value}")
    print(f"  Col 17 (Mvt D):        {ws.cell(row_num, 17).value}")
    print(f"  Col 20 (Mvt C):        {ws.cell(row_num, 20).value}")
    print(f"  Col 23 (Soldes D):     {ws.cell(row_num, 23).value}")
    print(f"  Col 26 (Soldes C):     {ws.cell(row_num, 26).value}")
    
    # Also check columns 14, 21, 27
    print(f"\n  [Other cols - might be empty or merges?]")
    print(f"  Col 14: {ws.cell(row_num, 14).value}")
    print(f"  Col 21: {ws.cell(row_num, 21).value}")
    print(f"  Col 27: {ws.cell(row_num, 27).value}")

print("\n" + "=" * 120)
print("CONCLUSION:")
print("=" * 120)
print("""
The columns that matter for DSF filling are:
- Col 23: Soldes Débit (final balance debit side)
- Col 26: Soldes Crédit (final balance credit side)

These should be the "MONTANT BRUT A LA CLOTURE" for DSF
""")

wb.close()
