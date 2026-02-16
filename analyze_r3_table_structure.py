"""Analyser la structure du tableau des dirigeants dans la fiche R3"""
from pathlib import Path
from openpyxl import load_workbook

template = Path("templates/DSF Normal standard.xlsx")
if not template.exists():
    print(f"❌ Template introuvable: {template}")
    exit(1)

wb = load_workbook(template)

# Chercher la fiche R3
if "Fiche R3" not in wb.sheetnames:
    print("❌ Fiche R3 introuvable")
    exit(1)

ws = wb["Fiche R3"]
print("FICHE R3 - Structure du tableau des dirigeants")
print("=" * 100)

# Analyser les 30 premières lignes
print("\nLignes 1-30 de la fiche R3:")
print("-" * 100)

for row_num in range(1, 31):
    row_data = []
    for col_num in range(1, 12):  # Colonnes A-K
        cell = ws.cell(row=row_num, column=col_num)
        if cell.value not in [None, ""]:
            col_letter = chr(64 + col_num)
            row_data.append(f"{col_letter}:{str(cell.value)[:30]}")
    
    if row_data:
        print(f"Ligne {row_num:2d}: {' | '.join(row_data[:5])}")

# Chercher spécifiquement les headers du tableau des dirigeants
print(f"\n{'=' * 100}")
print("Recherche des en-tetes du tableau des dirigeants:")
print("-" * 100)

found_headers = False
header_row = None

for row_num in range(1, 50):
    for col_num in range(1, 10):
        cell = ws.cell(row=row_num, column=col_num)
        if cell.value and isinstance(cell.value, str):
            val_lower = str(cell.value).lower()
            if "nom" in val_lower and "prénom" in str(ws.cell(row=row_num, column=col_num+1).value or "").lower():
                header_row = row_num
                found_headers = True
                print(f"\nEn-tetes trouves a la ligne {row_num}:")
                
                for c in range(col_num, min(col_num + 6, ws.max_column + 1)):
                    h_val = ws.cell(row=row_num, column=c).value
                    if h_val:
                        print(f"   Colonne {chr(64+c)}: {h_val}")
                break
    if found_headers:
        break

if header_row:
    print(f"\nZone de donnees (lignes {header_row+1} a {header_row+10}):")
    print("-" * 100)
    for row_num in range(header_row + 1, min(header_row + 11, ws.max_row + 1)):
        print(f"Ligne {row_num}: ", end="")
        is_empty = True
        for col_num in range(1, 7):
            cell = ws.cell(row=row_num, column=col_num)
            if cell.value not in [None, ""]:
                is_empty = False
                print(f"{chr(64+col_num)}={cell.value} | ", end="")
        if is_empty:
            print("(vide - zone à remplir)")
        else:
            print()

wb.close()

print(f"\n{'=' * 100}")
print("Analyse terminee")
