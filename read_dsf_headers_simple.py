#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Lire les libellés clés de colonnes DSF - Version simplifiée
"""
import openpyxl
from openpyxl.utils import get_column_letter

dsf_template = r"c:\Users\Emmaneul Ambadiang\Desktop\DSF\templates\DSF Normal standard.xlsx"
wb = openpyxl.load_workbook(dsf_template, data_only=False)

ws = wb["NOTE 3A"]

print("=" * 130)
print("LIBELLES DE COLONNES DSF - NOTE 3A (IMMOBILISATIONS)")
print("=" * 130)

print("\nCOLONNES (Row 8 - Main titles):")
for col_num in range(1, 16):
    val = ws.cell(8, col_num).value
    if val:
        print(f"  Col {col_num:2d}: {str(val)[:80]}")

print("\nCOLONNES (Row 9 - Secondary):")
for col_num in range(1, 16):
    val = ws.cell(9, col_num).value
    if val:
        print(f"  Col {col_num:2d}: {str(val)[:80]}")

print("\nCOLONNES (Row 10 - Details):")
for col_num in range(1, 16):
    val = ws.cell(10, col_num).value
    if val:
        print(f"  Col {col_num:2d}: {str(val)[:80]}")

print("\n" + "=" * 130)
print("OBSERVATIONS IMPORTANTES")
print("=" * 130)

print("""
1. Colonnes 1-3: Typiquement pour labels de comptes (A, B, C)
2. Colonne 4: Premier champ DATA
3. Colonnes 4-10: Contiennent les libellés des données à remplir

QUESTION CLÉ:
   Est-ce que les libellés mentionnent spécifiquement:
   - "OUVERTURE" vs "CLOTURE"?
   - "BILAN D'OUVERTURE 01/01" vs "BILAN DE CLOTURE 31/12"?
   - "BRUT" vs "NET"?
   - Types de mouvements (ACQUI, TRANSFERT, REEVAL, CESSION)?

Si les libellés EUX-MEMES sont explicites, on peut s'y fier pour le mapping!
""")

wb.close()
