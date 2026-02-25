#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Analyse du DSF 2023 pour comprendre les formules et remplissage des notes"""

from openpyxl import load_workbook
from pathlib import Path

dsf_2023 = Path(r"C:\Users\Emmaneul Ambadiang\Desktop\dsf\input\DSF GULFCAM 2023 V3.xlsx")

if not dsf_2023.exists():
    print(f"Fichier non trouvé: {dsf_2023}")
    exit(1)

wb = load_workbook(dsf_2023, data_only=False)

print("=" * 80)
print("ANALYSE DU DSF 2023")
print("=" * 80)

print("\n### FEUILLES DISPONIBLES ###")
for i, sheet_name in enumerate(wb.sheetnames):
    print(f"{i+1:2}. {sheet_name}")

# Examiner BILAN PAYSAGE
if "BILAN PAYSAGE" in wb.sheetnames:
    ws = wb["BILAN PAYSAGE"]
    print(f"\n### BILAN PAYSAGE (max_row={ws.max_row}, max_col={ws.max_column}) ###")
    
    print("\nCellules aux LIGNES 12-16 (cellules non-assignées d'après le rapport):")
    for row in range(12, 17):
        for col in ["G", "H", "I", "J", "K", "L"]:
            cell = ws[f"{col}{row}"]
            if cell.value:
                val_str = str(cell.value)[:50]
                print(f"  {col}{row}: {val_str}")

# Chercher les références vers les notes
print("\n### RECHERCHE DE REFERENCES VERS LES NOTES ###")
for sheet_name in wb.sheetnames[:5]:  # Vérifier les 5 premières feuilles
    ws = wb[sheet_name]
    print(f"\n--- {sheet_name} ---")
    
    # Chercher des cellules avec formules contenant 'NOTE' ou 'BILAN'
    note_refs = []
    for row in ws.iter_rows(min_row=1, max_row=min(50, ws.max_row)):
        for cell in row:
            if cell.value and isinstance(cell.value, str) and ("NOTE" in cell.value or "BILAN" in cell.value):
                note_refs.append((cell.coordinate, cell.value[:60]))
    
    if note_refs:
        for coord, val in note_refs[:5]:  # Afficher max 5
            print(f"  {coord}: {val}")

print("\n### ANALYSE DES formules dans BILAN PAYSAGE (colonnes G, L) ###")
ws = wb["BILAN PAYSAGE"]
for row in range(12, 17):
    for col in ["G", "L"]:
        cell = ws[f"{col}{row}"]
        if cell.data_type == "f" and cell.value:  # Si c'est une formule
            print(f"{col}{row}: FORMULA = {cell.value}")

print("\nDone!")
