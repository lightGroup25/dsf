#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Analyse complète des cellules à remplir dans le BILAN PAYSAGE du DSF 2023"""

from openpyxl import load_workbook
from pathlib import Path

dsf_2023 = Path(r"C:\Users\Emmaneul Ambadiang\Desktop\dsf\input\DSF GULFCAM 2023 V3.xlsx")

if not dsf_2023.exists():
    print(f"Fichier non trouvé: {dsf_2023}")
    exit(1)

wb = load_workbook(dsf_2023, data_only=True)  # data_only pour avoir les valeurs calculées
ws = wb["BILAN PAYSAGE"]

print("=" * 100)
print("ANALYSE COMPLÈTE - CELLULES DU BILAN PAYSAGE À REMPLIR")
print("=" * 100)

# Cellules non-assignées d'après le problème
problem_cells = [
    ("G12", "IMMOBILISATIONS INCORPORELLES"),
    ("L12", "Capital"),
    ("G13", "Frais de développement et de prospection"),
    ("L13", "Apporteurs capital non appelé       (-)"),
    ("G14", "Brevet, licences, logiciels et droits similaires"),
    ("L14", "Primes liées au capital social"),
    ("G15", "Fond commercial et droit au bail"),
    ("L15", "Ecarts de réévaluations"),
    ("G16", "Autres immobilisations incorporelles"),
    ("L16", "Réserves indisponibles"),
]

print("\n### CELLULES PROBLÉMATIQUES DU BILAN ###")
for coord, expected_label in problem_cells:
    cell = ws[coord]
    val_str = str(cell.value)[:50] if cell.value else "(vide)"
    print(f"\n{coord}: {val_str}")
    print(f"  Attendu: {expected_label}")

# Vérifier les NOTES liées
print("\n\n### NOTES CONNEXES DU DSF 2023 ###")
note_names = [name for name in wb.sheetnames if "NOTE" in name.upper() or "CAPITAL" in name.upper()]
for note_name in note_names[:15]:
    print(f"  - {note_name}")

# Chercher où sont les données de Capital dans NOTE 7
if "NOTE 7" in wb.sheetnames:
    ws_note7 = wb["NOTE 7"]
    print("\n### NOTE 7 - Structure (premières lignes) ###")
    for row in range(1, 15):
        for col in range(1, 6):
            cell = ws_note7.cell(row, col)
            if cell.value:
                val_str = str(cell.value)[:40]
                print(f"  {cell.coordinate}: {val_str}")

print("\nFin!")
